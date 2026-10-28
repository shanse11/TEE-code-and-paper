# -*- coding: utf-8 -*-
"""内容一致性 + 模板符合性核验。"""
import json
import os
import re
import sys

import docx
from docx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'
SRC = os.path.join(ROOT, '00_input_backup', '原稿_Hybrid_TEE_Rollup修改版.docx')
OUT = os.path.join(ROOT, '04_deliverable', 'Hybrid_TEE_Rollup_计算机应用_投稿版_待确认.docx')
REPORT = os.path.join(ROOT, '06_reports', 'verify_result.json')


def ptext(p):
    """段落可见文字（含公式的 m:t）。"""
    out = []
    for node in p.iter():
        if node.tag == W + 't':
            out.append(node.text or '')
        elif node.tag == M + 't':
            out.append(node.text or '')
        elif node.tag == W + 'tab':
            out.append('\t')
        elif node.tag == W + 'br':
            out.append(' ')
    return ''.join(out)


def load(path):
    d = docx.Document(path)
    paras, tables = [], []
    for ch in d.element.body:
        if ch.tag == W + 'p':
            pPr = ch.find(W + 'pPr')
            style = None
            num = None
            sect = False
            if pPr is not None:
                s = pPr.find(W + 'pStyle')
                if s is not None:
                    style = s.get(W + 'val')
                n = pPr.find(W + 'numPr')
                if n is not None:
                    num = (n.find(W + 'numId').get(W + 'val'),
                           n.find(W + 'ilvl').get(W + 'val'))
                sect = pPr.find(W + 'sectPr') is not None
            paras.append(dict(text=ptext(ch), style=style, num=num, sect=sect,
                              math=len(ch.findall(M + 'oMath')),
                              img=ch.find('.//' + W + 'drawing') is not None,
                              br=ch.find('.//' + W + 'br') is not None))
        elif ch.tag == W + 'tbl':
            rows = []
            for tr in ch.findall(W + 'tr'):
                rows.append([''.join(t.text or '' for t in tc.iter(W + 't'))
                             for tc in tr.findall(W + 'tc')])
            tables.append(rows)
    return paras, tables


HEAD = re.compile(r'^(\d+(?:\.\d+)*)\s+(.+)$')


def expected(idx, text):
    """源稿第 idx 段在成稿中应当呈现的文字（已授权的版式性差异）。"""
    t = text
    if idx == 3:
        t = t[len('摘  要：'):]
    elif idx == 4:
        t = t[len('关键词：'):]
    elif idx == 5:
        t = 'TP311\u3000\u3000文献标志码：A'
    elif idx == 6:
        t = t.replace('\n', ' ')
    elif idx == 9:
        t = t[len('Abstract: '):]
    else:
        m = HEAD.match(t)
        if m:
            t = m.group(2)
    return t


FORMULA_TXT = {
    48: ['cs = (id, s, F, T)'],
    51: ['M = (S, E, δ)'],
    52: ['δ: S × E → S'],
    54: ["Recover(cs) = cs'"],
    55: ['F(cs) = F(Recover(cs))'],
    94: ['C_total = C_commit + C_DA + p_challenge × C_challenge',
         'C_amortized = C_total / batch_size'],
    98: ['C_total = C_commit + C_DA + p_challenge × C_challenge'],
    99: ['C_amortized = C_total / batch_size'],
}


def main():
    sp, st = load(SRC)
    op, ot = load(OUT)
    res = dict(source_paras=len(sp), source_tables=len(st),
               out_paras=len(op), out_tables=len(ot))

    # ---------- 1. 表格逐格比对 ----------
    src_real = [t for t in st if len(t) > 1 and len(t[0]) > 1]
    tbl_diff = []
    for i, (a, b) in enumerate(zip(src_real, ot)):
        if a != b:
            tbl_diff.append(dict(table=i, src=a, out=b))
    res['tables_src'] = len(src_real)
    res['tables_out'] = len(ot)
    res['table_cell_diff'] = tbl_diff

    # ---------- 2. 段落顺序比对（子序列匹配）----------
    # 公式段：只比对非公式文字（公式另做专项记录）
    def wtext(i):
        p = [c for c in docx.Document(OUT).element.body
             if c.tag == W + 'p'][i]
        return ''.join(t.text or '' for t in p.iter(W + 't'))

    otexts = [p['text'] for p in op]
    otexts_w = [p['text'] for p in op]      # 占位，稍后按公式段替换
    out_paras = [c for c in docx.Document(OUT).element.body if c.tag == W + 'p']
    otexts_w = [''.join(t.text or '' for t in p.iter(W + 't')) for p in out_paras]

    unmatched, formula_rec = [], []
    oi = 0
    for idx, s in enumerate(sp):
        exp = expected(idx, s['text'])
        if idx in FORMULA_TXT:
            for f in FORMULA_TXT[idx]:
                exp = exp.replace(f, '')
            compare = otexts_w
        else:
            compare = otexts
        exp_n = exp.replace(' ', '').replace('\t', '')
        found = False
        j = oi
        while j < len(compare):
            c_n = compare[j].replace(' ', '').replace('\t', '')
            if c_n == exp_n or (exp_n and exp_n in c_n):
                found = True
                break
            j += 1
        if found:
            if idx in FORMULA_TXT:
                formula_rec.append(dict(src_idx=idx, src=s['text'],
                                        out_text_part=compare[j],
                                        out_math=''.join(
                                            t.text or '' for t in
                                            out_paras[j].iter(M + 't'))))
            oi = j + 1
        else:
            unmatched.append(dict(src_idx=idx, expected=exp[:120],
                                  src=sp[idx]['text'][:120]))
    res['unmatched_source_paras'] = unmatched
    res['formula_records'] = formula_rec

    # ---------- 3. 模板符合性 ----------
    styles = {}
    for p in op:
        styles[p['style']] = styles.get(p['style'], 0) + 1
    res['style_usage'] = styles
    res['math_paras'] = sum(1 for p in op if p['math'])
    res['math_total'] = sum(p['math'] for p in op)
    res['image_paras'] = [i for i, p in enumerate(op) if p['img']]
    res['section_breaks'] = [i for i, p in enumerate(op) if p['sect']]

    z = __import__('zipfile').ZipFile(OUT)
    names = z.namelist()
    res['has_comments'] = any('comments' in n for n in names)
    res['has_revision'] = (b'<w:ins ' in z.read('word/document.xml') or
                           b'<w:del ' in z.read('word/document.xml'))
    docxml = z.read('word/document.xml')
    res['header_refs'] = len(re.findall(rb'headerReference', docxml))
    res['footer_refs'] = len(re.findall(rb'footerReference', docxml))
    res['doi_line'] = b'doi:10.11772/j.issn.1001-9081' in docxml
    res['media'] = [n for n in names if n.startswith('word/media/')]

    os.makedirs(os.path.dirname(REPORT), exist_ok=True)
    json.dump(res, open(REPORT, 'w'), ensure_ascii=False, indent=1)

    print('源稿段落', res['source_paras'], '表格', res['source_tables'])
    print('成稿段落', res['out_paras'], '表格', res['out_tables'])
    print('表格差异', len(tbl_diff))
    print('未匹配源段落', len(unmatched))
    for u in unmatched:
        print('   !', u)
    print('公式对象数', res['math_total'], '含公式段落', res['math_paras'])
    print('图片段落', res['image_paras'])
    print('分节符位置', res['section_breaks'])
    print('样式使用', styles)
    print('批注部件', res['has_comments'], '修订标记', res['has_revision'])
    print('页眉引用', res['header_refs'], '页脚引用', res['footer_refs'],
          'DOI 行', res['doi_line'])
    print('媒体', res['media'])


if __name__ == '__main__':
    main()
