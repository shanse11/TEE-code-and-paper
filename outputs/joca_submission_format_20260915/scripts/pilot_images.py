"""小样本验证：SVG / EMF 两种矢量格式在 docx 中的可渲染性（LibreOffice 渲染 + 像素统计）。"""
import os
import shutil
import subprocess
import sys

import docx
from docx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from img_util import make_drawing, svg_size_emu  # noqa: E402

TPL = os.path.join(ROOT, '02_template_docx', '模板_投稿模版.docx')
SRC = os.path.join(ROOT, '00_input_backup', '原稿_Hybrid_TEE_Rollup修改版.docx')
SVGDIR = '/Users/shanse/Documents/ChatGPT/Hybrid TEE-Rollup/outputs/插图矢量版'
EMFDIR = '/tmp/emfout'
OUT = os.path.join(ROOT, '05_qa', 'pilot_images.docx')
SOFFICE = '/Users/shanse/.workbuddy/binaries/libreoffice/LibreOffice.app/Contents/MacOS/soffice'

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def empty_body_keep_sectpr(path):
    d = docx.Document(path)
    body = d.element.body
    final = body.find(qn('w:sectPr'))
    for ch in list(body):
        if ch is not final:
            body.remove(ch)
    # 保留 final sectPr（须为 body 最后一个元素）
    if final is not None and list(body)[-1] is not final:
        body.remove(final)
        body.append(final)
    return d


def main():
    d = empty_body_keep_sectpr(TPL)
    # 取一个可复用的 <w:drawing> 模板（来自原稿图 4）
    sd = docx.Document(SRC)
    tmpl = None
    for ch in sd.element.body.iterchildren():
        if ch.tag == qn('w:tbl'):
            dr = ch.find('.//' + W + 'drawing')
            if dr is not None:
                tmpl = dr
    assert tmpl is not None

    body = d.element.body
    final = body.find(qn('w:sectPr'))

    def para(text):
        p = d.add_paragraph(text)
        if final is not None and list(body)[-1] is not final:
            body.remove(final)
            body.append(final)
        return p

    para('SVG 测试：fig3')
    for kind, path in (('SVG', os.path.join(SVGDIR, 'fig3_byte_size.svg')),
                       ('EMF', os.path.join(EMFDIR, 'fig3_byte_size.emf')),
                       ('SVG', os.path.join(SVGDIR, 'fig4_da_cost.svg')),
                       ('EMF', os.path.join(EMFDIR, 'fig4_da_cost.emf'))):
        p = para('%s %s' % (kind, os.path.basename(path)))
        r = p.add_run()
        el = make_drawing(d, path, 8504 * 12700 // 914400 * 914400 // 1 if False else int(15 / 2.54 * 914400), tmpl, name=os.path.basename(path))
        r._r.append(el)
        if final is not None and list(body)[-1] is not final:
            body.remove(final)
            body.append(final)

    d.save(OUT)
    print('saved', OUT)

    outdir = os.path.join(ROOT, '05_qa', 'pilot_images_pdf')
    shutil.rmtree(outdir, ignore_errors=True)
    os.makedirs(outdir, exist_ok=True)
    subprocess.run([SOFFICE, '--headless', '--norestore',
                    '-env:UserInstallation=file:///tmp/lo_prof_joca',
                    '--convert-to', 'pdf', '--outdir', outdir, OUT],
                   capture_output=True, timeout=300)
    pdf = os.path.join(outdir, 'pilot_images.pdf')
    print('pdf', os.path.exists(pdf), os.path.getsize(pdf) if os.path.exists(pdf) else 0)

    # 逐页渲染并统计墨迹
    import pypdfium2 as pdfium
    from PIL import Image
    pdfdoc = pdfium.PdfDocument(pdf)
    for i in range(len(pdfdoc)):
        page = pdfdoc[i]
        img = page.render(scale=2).to_pil().convert('L')
        w, h = img.size
        px = img.load()
        rows = []
        step = 10
        for y0 in range(0, h, step):
            cnt = 0
            for y in range(y0, min(y0 + step, h), 2):
                for x in range(0, w, 4):
                    if px[x, y] < 200:
                        cnt += 1
            rows.append(cnt)
        nz = [(j * step, c) for j, c in enumerate(rows) if c > 0]
        print('page', i, 'size', w, h, 'ink bands', len(nz), 'total', sum(rows))
        print('   first/last band', nz[0] if nz else None, nz[-1] if nz else None)
        img.save(os.path.join(outdir, 'pilot-%02d.png' % (i + 1)))


if __name__ == '__main__':
    main()
