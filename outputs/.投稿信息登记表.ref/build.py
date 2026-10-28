# -*- coding: utf-8 -*-
"""《计算机应用》投稿信息登记表 —— openpyxl 建表脚本"""
import os

try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = "/Users/shanse/Documents/ChatGPT/Hybrid TEE-Rollup/outputs/投稿信息登记表.xlsx"
if os.path.exists(OUT):
    OUT = OUT.replace(".xlsx", "_20260912_151637.xlsx")


def xl_color(css_hex: str) -> str:
    v = css_hex.removeprefix("#").upper()
    if len(v) != 6:
        raise ValueError(f"Expected #RRGGBB, got: {css_hex}")
    return "FF" + v


XL_PRIMARY = xl_color("#4472C4")     # 表头/标题
XL_SUMMARY = xl_color("#2F5597")     # 汇总
XL_INPUT = xl_color("#D9E2F3")       # 输入区浅底
XL_BORDER = xl_color("#BFBFBF")
XL_GROUPBAR = xl_color("#8EA9DB")

thin = Side(style="thin", color=XL_BORDER)
medium = Side(style="medium", color=XL_GROUPBAR)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

F_TITLE = Font(name="微软雅黑", size=15, bold=True, color=xl_color("#FFFFFF"))
F_SUB = Font(name="微软雅黑", size=10, color=xl_color("#333333"))
F_HEAD = Font(name="微软雅黑", size=10.5, bold=True, color=xl_color("#FFFFFF"))
F_SUM = Font(name="微软雅黑", size=10.5, bold=True, color=xl_color("#FFFFFF"))
F_DATA = Font(name="宋体", size=10.5)
F_CAT = Font(name="微软雅黑", size=10.5, bold=True, color=xl_color("#1F3864"))

FILL_TITLE = PatternFill("solid", fgColor=XL_PRIMARY)
FILL_HEAD = PatternFill("solid", fgColor=XL_PRIMARY)
FILL_SUM = PatternFill("solid", fgColor=XL_SUMMARY)
FILL_INPUT = PatternFill("solid", fgColor=XL_INPUT)

CENTER = Alignment(horizontal="center", vertical="center")
LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)
CENTER_TOP = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()
wb.properties.title = "《计算机应用》投稿信息登记表"

# =========================================================
# Sheet 1 投稿信息登记表
# =========================================================
ws = wb.active
ws.title = "投稿信息登记表"
ws.sheet_view.showGridLines = False

ROWS = [
    # (类别, 登记项, 填写内容, 填写要求与示例, 状态)
    ("A 作者信息", "作者总数", "", "阿拉伯数字；如仅杨帆一人填 1", "待填写"),
    ("A 作者信息", "第 1 作者姓名", "杨帆", "与投稿系统一致；多位作者请在「多作者明细」表逐位登记", "已确认"),
    ("A 作者信息", "姓名拼音 / 英文名", "YANG Fan", "姓全大写、名首字母大写，格式如 HU Dan（胡丹）", "已确认"),
    ("A 作者信息", "出生年", "", "四位年份，如 1998", "待填写"),
    ("A 作者信息", "性别", "", "男 / 女", "待填写"),
    ("A 作者信息", "籍贯（省 + 市）", "", "具体到城市，如：甘肃兰州", "待填写"),
    ("A 作者信息", "职称 / 身份", "", "硕士研究生 / 博士研究生 / 硕士 / 博士 / 讲师 / 副教授 / 教授；在读须标“研究生”，已毕业标“硕士 / 博士”，博士后只标博士", "待填写"),
    ("A 作者信息", "学历", "", "博士研究生 / 硕士研究生 / 博士 / 硕士", "待填写"),
    ("A 作者信息", "是否 CCF 会员", "", "是 / 否；是则必须填写会员号", "待填写"),
    ("A 作者信息", "CCF 会员号", "", "非会员填“无”", "待填写"),
    ("A 作者信息", "主要研究方向", "", "3~5 个，顿号（、）分隔；不要写“计算机应用”这类过宽表述", "待填写"),
    ("A 作者信息", "在本文中的工作量", "", "单作者请列全部环节：选题、方案设计、原型实现、实验、论文撰写", "待填写"),
    ("A 作者信息", "通信作者邮箱", "2991777340@qq.com", "仅一位作者时本人即通信作者", "已确认"),
    ("A 作者信息", "联系人姓名", "杨帆", "用于编辑部沟通，可与通信作者不同", "已确认"),
    ("A 作者信息", "联系人手机号", "13319693336", "必填", "已确认"),
    ("A 作者信息", "联系人邮箱", "2991777340@qq.com", "必填", "已确认"),
    ("A 作者信息", "通信地址", "上海市普陀区中山北路3663号华东师范大学数据科学与工程学院，200062", "精确到门牌号与邮编", "已确认"),

    ("B 单位信息", "中文单位（一级 二级）", "华东师范大学 数据科学与工程学院", "须用正式对外全称，不使用缩略语", "已确认"),
    ("B 单位信息", "单位所在地与邮编", "上海 200062", "直辖市可不写省", "已确认"),
    ("B 单位信息", "英文单位", "School of Data Science and Engineering, East China Normal University, Shanghai 200062, China", "须到学校中英文官网核实", "已确认"),
    ("B 单位信息", "是否有第二单位", "否", "是 / 否", "待填写"),
    ("B 单位信息", "第二单位中英文", "无", "有则填写，须与官网一致", "不适用"),

    ("C 稿件元数据", "中文题名", "面向高频 DApp 的 Hybrid TEE-Rollup 可恢复挑战协议与轻量 DA 成本评估", "一般不超过 20 个汉字；建议不用标点与不常见缩略语、不用副标题", "已确认"),
    ("C 稿件元数据", "备选题名", "", "建议把 DA 展开为“数据可用性”，减少缩略语", "待填写"),
    ("C 稿件元数据", "英文题名", "Recoverable Challenge Protocol and Lightweight DA Cost Evaluation for Hybrid TEE-Rollup in High-Frequency DApps", "与中文对应；须合并为一个段落，让其自动换行", "已确认"),
    ("C 稿件元数据", "中图分类号", "TP311", "请到 clcindex.com 复核；备选 TP309（信息安全）/ TP393（计算机网络）", "待填写"),
    ("C 稿件元数据", "文献标志码", "A", "固定为 A", "已确认"),
    ("C 稿件元数据", "中文关键词（5 个）", "", "必须为 5 个，分号（；）分隔；现为 6 个需删减 1 个", "待填写"),
    ("C 稿件元数据", "英文关键词（5 个）", "", "必须为 5 个；除通用专有名词外一律小写，须写全称，不用复数", "待填写"),
    ("C 稿件元数据", "投稿类型", "新投稿", "新投稿 / 退修稿（退修稿须附逐条修改说明）", "待填写"),
    ("C 稿件元数据", "收稿日期", "", "新投稿留空，由编辑部填写", "不适用"),
    ("C 稿件元数据", "DOI 行", "保留模板原文", "不得修改或删除模板首页页眉与 DOI 行", "已确认"),

    ("D 基金与项目", "是否有基金支持", "", "是 / 否；无也须明确注明", "待填写"),
    ("D 基金与项目", "项目种类", "", "如：国家自然科学基金资助项目、国家重点研发计划项目", "待填写"),
    ("D 基金与项目", "项目编号", "", "必填；因涉密等原因不能提供须另附说明", "待填写"),
    ("D 基金与项目", "项目名称", "", "必填", "待填写"),
    ("D 基金与项目", "项目主持人", "", "必填", "待填写"),
    ("D 基金与项目", "项目起止时间", "", "格式如 2023-01—2026-12", "待填写"),
    ("D 基金与项目", "项目主要研究内容", "", "必填", "待填写"),
    ("D 基金与项目", "本文与项目研究的关系", "", "必填", "待填写"),
    ("D 基金与项目", "本人在项目承担人中的排名", "", "阿拉伯数字，如 3", "待填写"),
    ("D 基金与项目", "本人在项目中承担的工作", "", "必填", "待填写"),
    ("D 基金与项目", "基金项目英文译文", "", "尽量使用全称并核实正确拼写", "待填写"),

    ("E 图表与素材", "图 1 矢量源文件", "", "现为 PNG，须提供 emf/wmf/svg 或绘图源文件", "待填写"),
    ("E 图表与素材", "图 2 矢量源文件", "", "现为 PNG，须提供 emf/wmf/svg 或绘图源文件", "待填写"),
    ("E 图表与素材", "图 3 矢量源文件", "", "现为 PNG，须提供 emf/wmf/svg 或绘图源文件", "待填写"),
    ("E 图表与素材", "图 4 矢量源文件", "", "现为 PNG，须提供 emf/wmf/svg 或绘图源文件", "待填写"),
    ("E 图表与素材", "绘图工具", "", "Matplotlib / Visio / PPT / Excel（需同时提供源文件或数据表）", "待填写"),
    ("E 图表与素材", "是否授权按现有数据重绘", "", "找不到矢量源文件时填“是”", "待填写"),
    ("E 图表与素材", "是否需补中文参考文献", "", "建议补 3~5 篇近三年中文核心期刊文献", "待填写"),

    ("F 格式改造确认", "章节编号改为 0 起", "待办", "0 引言 / 1 相关工作 / 2 系统概览 / 3 核心机制 / 4 实验评估 / 5 结语", "待填写"),
    ("F 格式改造确认", "表编号重编为表 1、表 2", "待办", "现为表 4、表 5，全文不存在表 1~表 3", "待填写"),
    ("F 格式改造确认", "公式改用公式编辑器", "待办", "须为 OMML 对象，不得处理成图片", "待填写"),
    ("F 格式改造确认", "图题移出表格并设嵌入型", "待办", "现图题置于 1×1 表格单元格内，须改为图下方独立段落", "待填写"),
    ("F 格式改造确认", "页眉改为《计算机应用》", "待办", "现页眉为《计算机科学与探索》", "待填写"),
    ("F 格式改造确认", "全文套用模板样式", "待办", "现 147 个段落全部为 Normal 样式", "待填写"),
    ("F 格式改造确认", "表格改三线表并统一精度", "待办", "同列同精度，单位写入项目栏“/单位”", "待填写"),
    ("F 格式改造确认", "随附材料：作者工作量说明", "待办", "投稿必须随附", "待填写"),
    ("F 格式改造确认", "随附材料：研究背景与基金项目介绍", "待办", "投稿必须随附", "待填写"),
    ("F 格式改造确认", "随附材料：变量符号说明表", "待办", "无矩阵 / 矢量变量时填“无”", "待填写"),
    ("F 格式改造确认", "随附材料：修改说明", "不适用", "仅退修稿需要", "不适用"),
]

# 锚点：第1行标题、第2行副标题、第3行进度、第4行空行、第5行表头、第6行起数据
HDR = 5
FIRST = 6
LAST = FIRST + len(ROWS) - 1

ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
c = ws.cell(row=1, column=1, value="《计算机应用》投稿信息登记表")
c.font = F_TITLE
c.fill = FILL_TITLE
c.alignment = CENTER
ws.row_dimensions[1].height = 30

ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=6)
c = ws.cell(row=2, column=2 - 1, value="稿件：Hybrid_TEE_Rollup修改版.docx　｜　模板：投稿模版.doc（joca.cn）　｜　登记日期：2026-09-12")
c.font = F_SUB
c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 20

# 进度汇总区（第 3 行）
for col, val in ((1, "完成率"), (3, "已确认 / 有效必填"), (5, "待填写")):
    cc = ws.cell(row=3, column=col, value=val)
    cc.font = F_SUM
    cc.fill = FILL_SUM
    cc.alignment = CENTER
    cc.border = Border(top=Side(style="medium", color=XL_SUMMARY), bottom=thin, left=thin, right=thin)
N_ALL = f'COUNTA($C${FIRST}:$C${LAST})'
N_NA = f'COUNTIF($F${FIRST}:$F${LAST},"不适用")'
N_OK = f'COUNTIF($F${FIRST}:$F${LAST},"已确认")'
N_TODO = f'COUNTIF($F${FIRST}:$F${LAST},"待填写")'
ws.cell(row=3, column=2, value=f'=IF(OR({N_ALL}=0,{N_ALL}-{N_NA}=0),"",{N_OK}/({N_ALL}-{N_NA}))')
ws.cell(row=3, column=4, value=f'={N_OK}&" / "&({N_ALL}-{N_NA})')
ws.cell(row=3, column=6, value=f'={N_TODO}')
for col in (2, 4, 6):
    cc = ws.cell(row=3, column=col)
    cc.font = F_SUM
    cc.fill = FILL_SUM
    cc.alignment = CENTER
    cc.border = Border(top=Side(style="medium", color=XL_SUMMARY), bottom=thin, left=thin, right=thin)
ws.cell(row=3, column=2).number_format = "0.0%"
ws.row_dimensions[3].height = 22

# 表头（第 5 行）
HEADERS = ["序号", "类别", "登记项", "填写内容", "填写要求与示例", "状态"]
for i, h in enumerate(HEADERS, start=1):
    cc = ws.cell(row=HDR, column=i, value=h)
    cc.font = F_HEAD
    cc.fill = FILL_HEAD
    cc.alignment = CENTER
    cc.border = BOX
ws.row_dimensions[HDR].height = 24

# 数据区
prev_cat = None
for i, (cat, item, val, hint, status) in enumerate(ROWS):
    r = FIRST + i
    is_new_cat = (cat != prev_cat)
    ws.cell(row=r, column=1, value=i + 1)
    ws.cell(row=r, column=2, value=cat)
    ws.cell(row=r, column=3, value=item)
    ws.cell(row=r, column=4, value=val)
    ws.cell(row=r, column=5, value=hint)
    ws.cell(row=r, column=6, value=status)
    prev_cat = cat

    for col in range(1, 7):
        cc = ws.cell(row=r, column=col)
        cc.font = F_CAT if col == 2 else F_DATA
        if col in (1, 6):
            cc.alignment = CENTER
        else:
            cc.alignment = LEFT_TOP
        if col == 4:
            cc.fill = FILL_INPUT
        b = Border(left=thin, right=thin, top=thin, bottom=thin)
        if is_new_cat:
            b = Border(left=thin, right=thin, top=medium, bottom=thin)
        cc.border = b
    ws.cell(row=r, column=4).alignment = LEFT_TOP
    ws.cell(row=r, column=5).alignment = LEFT_TOP

for col, w in zip("ABCDEF", (6, 16, 24, 40, 46, 11)):
    ws.column_dimensions[col].width = w

# 数据验证：状态列
dv_status = DataValidation(type="list", formula1='"待填写,已确认,不适用"', allow_blank=True, showDropDown=False)
dv_status.error = "请从下拉中选择：待填写 / 已确认 / 不适用"
dv_status.errorTitle = "状态取值不合法"
ws.add_data_validation(dv_status)
dv_status.add(f"F{FIRST}:F{LAST}")

# 条件格式：状态列
rng = f"F{FIRST}:F{LAST}"
ws.conditional_formatting.add(rng, CellIsRule(
    operator="equal", formula=['"已确认"'],
    fill=PatternFill("solid", bgColor=xl_color("#C6EFCE")),
    font=Font(name="宋体", size=10.5, color=xl_color("#006100"))))
ws.conditional_formatting.add(rng, CellIsRule(
    operator="equal", formula=['"待填写"'],
    fill=PatternFill("solid", bgColor=xl_color("#FFEB9C")),
    font=Font(name="宋体", size=10.5, color=xl_color("#9C6500"))))
ws.conditional_formatting.add(rng, CellIsRule(
    operator="equal", formula=['"不适用"'],
    fill=PatternFill("solid", bgColor=xl_color("#E7E6E6")),
    font=Font(name="宋体", size=10.5, color=xl_color("#808080"))))

ws.freeze_panes = f"A{FIRST}"
ws.auto_filter.ref = f"A{HDR}:F{LAST}"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_area = f"A1:F{LAST}"
ws.print_title_rows = f"{HDR}:{HDR}"

# =========================================================
# Sheet 2 多作者明细（多于一位作者时使用）
# =========================================================
ws2 = wb.create_sheet("多作者明细")
ws2.sheet_view.showGridLines = False
H2 = ["序号", "姓名", "姓名拼音/英文名", "出生年", "性别", "籍贯（省+市）",
      "职称/身份", "学历", "是否CCF会员", "CCF会员号", "主要研究方向",
      "所属单位序号", "通信作者", "在本文中的工作量"]
W2 = [6, 12, 20, 9, 7, 16, 16, 14, 12, 12, 34, 12, 10, 40]
for i, (h, w) in enumerate(zip(H2, W2), start=1):
    cc = ws2.cell(row=1, column=i, value=h)
    cc.font = F_HEAD
    cc.fill = FILL_HEAD
    cc.alignment = CENTER
    cc.border = BOX
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.row_dimensions[1].height = 30

RESERVED = 5
for i in range(RESERVED):
    r = 2 + i
    ws2.cell(row=r, column=1, value=i + 1)
    for col in range(1, len(H2) + 1):
        cc = ws2.cell(row=r, column=col)
        cc.font = F_DATA
        cc.alignment = CENTER if col in (1, 4, 5, 9, 12, 13) else LEFT_TOP
        if col not in (1,):
            cc.fill = FILL_INPUT
        cc.border = BOX
ws2.freeze_panes = "A2"

DV_SEX = DataValidation(type="list", formula1='"男,女"', allow_blank=True)
DV_YN = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
DV_TITLE = DataValidation(type="list", formula1='"硕士研究生,博士研究生,硕士,博士,讲师,副教授,教授"', allow_blank=True)
DV_DEG = DataValidation(type="list", formula1='"硕士研究生,博士研究生,硕士,博士"', allow_blank=True)
for dv, rng in ((DV_SEX, f"E2:E{1+RESERVED}"), (DV_YN, f"I2:I{1+RESERVED}"),
                (DV_YN, f"M2:M{1+RESERVED}"), (DV_TITLE, f"G2:G{1+RESERVED}"),
                (DV_DEG, f"H2:H{1+RESERVED}")):
    ws2.add_data_validation(dv)
    dv.add(rng)
ws2.page_setup.orientation = "landscape"
ws2.page_setup.fitToWidth = 1
ws2.page_setup.fitToHeight = 0
ws2.sheet_properties.pageSetUpPr.fitToPage = True

# =========================================================
# Sheet 3 图表素材清单
# =========================================================
ws3 = wb.create_sheet("图表素材清单")
ws3.sheet_view.showGridLines = False
H3 = ["序号", "图/表号", "中文题名", "英文题名", "当前格式与尺寸", "矢量源文件路径", "绘图工具", "是否需重出", "备注"]
W3 = [6, 10, 32, 42, 24, 30, 14, 12, 40]
for i, (h, w) in enumerate(zip(H3, W3), start=1):
    cc = ws3.cell(row=1, column=i, value=h)
    cc.font = F_HEAD
    cc.fill = FILL_HEAD
    cc.alignment = CENTER
    cc.border = BOX
    ws3.column_dimensions[get_column_letter(i)].width = w
ws3.row_dimensions[1].height = 30

MEDIA = [
    (1, "图 1", "Hybrid TEE-Rollup 原型系统架构", "Hybrid TEE-Rollup prototype system architecture",
     "PNG，7.73 × 6.07 cm", "", "", "是", "须重出至版心宽度约 15 cm；文字不小于小5号"),
    (2, "图 2", "可恢复挑战协议流程", "Workflow of the recoverable challenge protocol",
     "PNG，8.04 × 6.17 cm", "", "", "是", "流程框架图，不要底色/灰度/阴影"),
    (3, "图 3", "完整负载与紧凑提交的字节规模对比", "Byte-size comparison of full payload and compact commitment",
     "PNG，7.87 × 3.70 cm", "", "", "是", "黑白印刷，用线型与标记区分曲线，不用颜色"),
    (4, "图 4", "不同数据可用性配置下的摊销 gas 用量趋势", "Amortized gas usage trend under different DA configurations",
     "PNG，7.87 × 4.42 cm", "", "", "是", "多曲线尽量用实线；不使用三维图"),
    (5, "表 1", "失败场景分类与当前证据边界", "Failure scenarios and evidence boundaries",
     "Word 表格", "—", "—", "否", "现编号为表 4，须重编为表 1；改三线表"),
    (6, "表 2", "本地 EVM 关键路径平均 gasUsed", "Average gasUsed of local EVM key paths",
     "Word 表格", "—", "—", "否", "现编号为表 5，须重编为表 2；gasUsed 列统一小数位"),
]
for i, row in enumerate(MEDIA):
    r = 2 + i
    for j, v in enumerate(row, start=1):
        cc = ws3.cell(row=r, column=j, value=v)
        cc.font = F_DATA
        cc.alignment = CENTER if j in (1, 2, 5, 7, 8) else LEFT_TOP
        if j in (6, 7):
            cc.fill = FILL_INPUT
        cc.border = BOX
DV_RE = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
ws3.add_data_validation(DV_RE)
DV_RE.add(f"H2:H{1+len(MEDIA)}")
ws3.freeze_panes = "A2"
ws3.page_setup.orientation = "landscape"
ws3.page_setup.fitToWidth = 1
ws3.page_setup.fitToHeight = 0
ws3.sheet_properties.pageSetUpPr.fitToPage = True

# =========================================================
# Sheet 4 填写说明与选项
# =========================================================
ws4 = wb.create_sheet("填写说明与选项")
ws4.sheet_view.showGridLines = False
H4 = ["项目", "说明 / 可选项"]
for i, h in enumerate(H4, start=1):
    cc = ws4.cell(row=1, column=i, value=h)
    cc.font = F_HEAD
    cc.fill = FILL_HEAD
    cc.alignment = CENTER
    cc.border = BOX
ws4.column_dimensions["A"].width = 24
ws4.column_dimensions["B"].width = 96
ws4.row_dimensions[1].height = 26

NOTES = [
    ("状态列取值", "待填写 / 已确认 / 不适用。填完后把状态改为“已确认”，首页完成率会自动更新。"),
    ("职称 / 身份", "在读分别标“硕士研究生 / 博士研究生”；已毕业分别标“硕士 / 博士”；博士后只标明博士。"),
    ("CCF 会员", "是会员必须标明会员号；不是会员直接略过不写，不写“非会员”。"),
    ("主要研究方向", "3~5 个，用顿号（、）隔开；范围不要过大，如“计算机应用”不合格。"),
    ("单位署名", "不使用缩略语，必须是正式对外名称；投稿后作者署名、单位顺序、基金信息不得随意更改。"),
    ("题名", "一般不超过 20 个汉字，建议不用标点及不常见缩略语，不用副标题。"),
    ("关键词", "中英文各 5 个。中文用“；”间隔；英文除 TCP/IP 等通用专有名词外均小写，须写全称，不使用复数。"),
    ("中图分类号", "必须填写，可到 https://www.clcindex.com/ 查询。"),
    ("基金项目", "须提供正确编号；因涉密等原因不能提供编号请另附说明。"),
    ("矢量图", "请用 .emf / .wmf（可编辑文字，不要转曲），或 .svg；同时提供绘图源文件（Matplotlib / Visio / PPT / Excel）。"),
    ("表格", "三线表；同列表示同一参数、精度一致；单位在项目栏以“/单位”形式表示。"),
    ("公式", "只能使用公式编辑器编辑，不能处理成图片；矩阵、矢量用 Times New Roman 加粗斜体，其他变量常规斜体。"),
    ("缩略语", "全文检查：中文摘要“支持向量机(SVM)”，英文摘要“Support Vector Machine (SVM)”，正文首次“支持向量机(Support Vector Machine, SVM)”。"),
    ("百分比与百分点", "75% 对 50%：差值表述为“提升 25 个百分点”；比率表述为“提升 50%”。两者不可混用。"),
    ("随附材料", "① 作者工作量说明；② 研究背景与基金项目介绍；③ 变量符号说明表；④ 修改说明（仅退修稿）。"),
]
for i, (a, b) in enumerate(NOTES):
    r = 2 + i
    ca = ws4.cell(row=r, column=1, value=a)
    cb = ws4.cell(row=r, column=2, value=b)
    ca.font = F_CAT
    cb.font = F_DATA
    ca.alignment = CENTER_TOP
    cb.alignment = LEFT_TOP
    ca.border = BOX
    cb.border = BOX
ws4.freeze_panes = "A2"

wb.save(OUT)
print("SAVED:", OUT)
