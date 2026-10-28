#!/usr/bin/env python3
from copy import deepcopy
from pathlib import Path
import json
import re
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[3]
RUN = ROOT / "outputs/submission_format_20260915_1700"
TEMPLATE = RUN / "work/投稿模版_WPS.docx"
SOURCE = ROOT / "Hybrid_TEE_Rollup修改版.docx"
WORK = RUN / "work/投稿排版_工作稿.docx"
FIG_DIR = RUN / "work/source_figures"
FIG_DIR.mkdir(parents=True, exist_ok=True)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def remove_paragraph(p):
    p._element.getparent().remove(p._element)


def clear_paragraph(p):
    for child in list(p._element):
        if child.tag != qn("w:pPr"):
            p._element.remove(child)


def copy_runs(source_p, target_p, bold_prefix=None):
    clear_paragraph(target_p)
    for sr in source_p.runs:
        if not sr.text:
            continue
        tr = target_p.add_run(sr.text)
        tr.bold = sr.bold
        tr.italic = sr.italic
        tr.underline = sr.underline
        if sr.font.superscript:
            tr.font.superscript = True
        if sr.font.subscript:
            tr.font.subscript = True
    if not target_p.runs and source_p.text:
        target_p.add_run(source_p.text)
    if bold_prefix and target_p.text.startswith(bold_prefix):
        # Rebuild only to format the label; text remains byte-for-byte identical.
        text = target_p.text
        clear_paragraph(target_p)
        target_p.add_run(bold_prefix).bold = True
        target_p.add_run(text[len(bold_prefix):])


def set_style(p, name):
    p.style = name


def suppress_style_numbering(p):
    ppr = p._p.get_or_add_pPr()
    old = ppr.find(qn("w:numPr"))
    if old is not None:
        ppr.remove(old)
    numpr = OxmlElement("w:numPr")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), "0")
    numpr.append(numid)
    ppr.append(numpr)


def add_source_paragraph(doc, sp, style="Normal", heading=False):
    p = doc.add_paragraph(style=style)
    copy_runs(sp, p)
    if heading:
        suppress_style_numbering(p)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True
    return p


def math_el(tag):
    return OxmlElement(f"m:{tag}")


def math_run(text, italic=None):
    r = math_el("r")
    rpr = math_el("rPr")
    sty = math_el("sty")
    sty.set(qn("m:val"), "i" if italic is not False else "p")
    rpr.append(sty)
    wrpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    wrpr.append(sz)
    r.append(rpr)
    r.append(wrpr)
    t = math_el("t")
    if text[:1].isspace() or text[-1:].isspace():
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    return r


def math_sub(base, sub, base_italic=True):
    ssub = math_el("sSub")
    ssub.append(math_el("sSubPr"))
    e = math_el("e")
    e.append(math_run(base, italic=base_italic))
    subel = math_el("sub")
    subel.append(math_run(sub, italic=False))
    ssub.extend([e, subel])
    return ssub


def math_frac(num_nodes, den_nodes):
    frac = math_el("f")
    frac.append(math_el("fPr"))
    num = math_el("num")
    den = math_el("den")
    for n in num_nodes:
        num.append(n)
    for n in den_nodes:
        den.append(n)
    frac.extend([num, den])
    return frac


def omath(nodes):
    om = math_el("oMath")
    for node in nodes:
        om.append(node)
    return om


def add_mixed_formula_paragraph(doc, sp, formula, nodes, style="Normal"):
    text = sp.text
    start = text.index(formula)
    end = start + len(formula)
    p = doc.add_paragraph(style=style)
    if text[:start]:
        p.add_run(text[:start])
    p._p.append(omath(nodes))
    if text[end:]:
        p.add_run(text[end:])
    return p


def c_nodes(subscript):
    return math_sub("C", subscript)


def cost_total_nodes():
    return [
        c_nodes("total"), math_run(" = ", False), c_nodes("commit"), math_run(" + ", False),
        c_nodes("DA"), math_run(" + ", False), math_sub("p", "challenge"),
        math_run(" × ", False), c_nodes("challenge")
    ]


def cost_amortized_nodes():
    return [
        c_nodes("amortized"), math_run(" = ", False),
        math_frac([c_nodes("total")], [math_sub("batch", "size")])
    ]


INLINE_FORMULAS = {
    48: ("cs = (id, s, F, T)", [math_run("cs"), math_run(" = (", False), math_run("id"), math_run(", ", False), math_run("s"), math_run(", ", False), math_run("F"), math_run(", ", False), math_run("T"), math_run(")", False)]),
    51: ("M = (S, E, δ)", [math_run("M"), math_run(" = (", False), math_run("S"), math_run(", ", False), math_run("E"), math_run(", ", False), math_run("δ"), math_run(")", False)]),
    52: ("δ: S × E → S", [math_run("δ"), math_run(": ", False), math_run("S"), math_run(" × ", False), math_run("E"), math_run(" → ", False), math_run("S")]),
    54: ("Recover(cs) = cs'", [math_run("Recover", False), math_run("(", False), math_run("cs"), math_run(") = ", False), math_run("cs'" )]),
    55: ("F(cs) = F(Recover(cs))", [math_run("F"), math_run("(", False), math_run("cs"), math_run(") = ", False), math_run("F"), math_run("(", False), math_run("Recover", False), math_run("(", False), math_run("cs"), math_run("))", False)]),
    77: ("Facts(cs_before) = Facts(cs_after)", [math_run("Facts", False), math_run("(", False), math_sub("cs", "before"), math_run(") = ", False), math_run("Facts", False), math_run("(", False), math_sub("cs", "after"), math_run(")", False)]),
}


def add_numbered_equation(doc, source_text, nodes, number):
    p = doc.add_paragraph(style="MTDisplayEquation")
    suppress_style_numbering(p)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Pt(0)
    p.paragraph_format.right_indent = Pt(0)
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(1.58), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(3.12), WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.keep_together = True
    p.add_run("\t")
    p._p.append(omath(nodes))
    p.add_run(f"\t（{number}）")
    return p


def set_keep(p, with_next=False, together=False):
    p.paragraph_format.keep_with_next = with_next
    p.paragraph_format.keep_together = together


def extract_figures(source):
    out = []
    with zipfile.ZipFile(SOURCE) as z:
        for i, shape in enumerate(source.inline_shapes, 1):
            rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
            part = source.part.related_parts[rid]
            ext = Path(str(part.partname)).suffix
            path = FIG_DIR / f"figure_{i}{ext}"
            path.write_bytes(part.blob)
            out.append((path, shape.width, shape.height))
    return out


def add_figure(doc, source, fig_index, figure_info):
    path, width, height = figure_info
    ip = doc.add_paragraph()
    ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ip.paragraph_format.keep_with_next = True
    ip.add_run().add_picture(str(path), width=width, height=height)
    captions = [p.text for p in source.tables[fig_index].cell(0, 0).paragraphs if p.text]
    for i, text in enumerate(captions):
        cp = doc.add_paragraph(text, style="图表名称")
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.keep_with_next = i < len(captions) - 1


def style_table(table, table_index):
    table.autofit = False
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
    for ri, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        for height in list(trpr.findall(qn("w:trHeight"))):
            trpr.remove(height)
        cant = OxmlElement("w:cantSplit")
        trpr.append(cant)
        if ri == 0:
            header = OxmlElement("w:tblHeader")
            trpr.append(header)
        for cell in row.cells:
            cell.vertical_alignment = 1
            tcpr = cell._tc.get_or_add_tcPr()
            mar = tcpr.find(qn("w:tcMar"))
            if mar is None:
                mar = OxmlElement("w:tcMar")
                tcpr.append(mar)
            for edge, value in (("top", "60"), ("bottom", "60"), ("left", "40"), ("right", "40")):
                el = mar.find(qn(f"w:{edge}"))
                if el is None:
                    el = OxmlElement(f"w:{edge}")
                    mar.append(el)
                el.set(qn("w:w"), value)
                el.set(qn("w:type"), "dxa")
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(9)
    widths = ([0.66, 0.82, 0.86, 0.86] if table_index == 4 else [1.15, 0.75, 1.30])
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def insert_source_table(doc, source_table, table_index):
    body = doc._element.body
    sectpr = body.sectPr
    body.insert(body.index(sectpr), deepcopy(source_table._element))
    table = doc.tables[-1]
    style_table(table, table_index)
    return table


def configure_math_and_fields(doc):
    settings = doc.settings._element
    mathpr = settings.find(qn("m:mathPr"))
    if mathpr is None:
        mathpr = math_el("mathPr")
        settings.append(mathpr)
    font = mathpr.find(qn("m:mathFont"))
    if font is None:
        font = math_el("mathFont")
        mathpr.insert(0, font)
    font.set(qn("m:val"), "Times New Roman")
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    track = settings.find(qn("w:trackRevisions"))
    if track is not None:
        settings.remove(track)


def clear_template_example_footers(doc):
    seen = set()
    for section in doc.sections:
        for footer in (section.first_page_footer, section.footer):
            partname = str(footer.part.partname)
            if partname in seen:
                continue
            seen.add(partname)
            for p in footer.paragraphs:
                clear_paragraph(p)
            for table in list(footer.tables):
                table._element.getparent().remove(table._element)


def remove_front_instruction_shape(doc):
    first = doc.paragraphs[0]._element
    for run in list(first.findall(qn("w:r"))):
        if run.xpath(".//wp:anchor"):
            first.remove(run)


def main():
    template = Document(TEMPLATE)
    source = Document(SOURCE)
    figures = extract_figures(source)

    tp = list(template.paragraphs)
    # Capture the relevant two-column body section before deleting teaching pages.
    body_sect = deepcopy(tp[55]._p.pPr.sectPr)

    front_map = {
        1: (0, "Title", None),
        2: (1, "作者名", None),
        3: (2, "单位", None),
        6: (3, "摘要关键词", "摘  要："),
        12: (4, "摘要关键词", "关键词："),
        13: (5, "固定3", None),
        14: (6, "Title", None),
        15: (7, "作者名", None),
        16: (8, "单位", None),
        18: (9, "摘要关键词", "Abstract:"),
        19: (10, "摘要关键词", "Key words:"),
    }
    for ti, (si, style, prefix) in front_map.items():
        copy_runs(source.paragraphs[si], tp[ti], bold_prefix=prefix)
        set_style(tp[ti], style)
        if ti in (1, 2, 3, 14, 15, 16):
            tp[ti].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Delete unused example slots from the front matter; retain the blank separator and its section break.
    for idx in (4, 5, 7, 8, 9, 10, 11, 17):
        remove_paragraph(tp[idx])

    body = template._element.body
    separator = tp[21]._element
    sep_index = body.index(separator)
    for child in list(body)[sep_index + 1:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    old_final = body.sectPr
    body.replace(old_final, body_sect)

    source_paragraph_index = {p._element: i for i, p in enumerate(source.paragraphs)}
    source_table_index = {t._element: i for i, t in enumerate(source.tables)}
    for child in source._element.body:
        if child.tag == qn("w:p"):
            i = source_paragraph_index[child]
            if i < 11:
                continue
            sp = source.paragraphs[i]
            text = sp.text
            if i in (98, 99):
                add_numbered_equation(template, text, cost_total_nodes() if i == 98 else cost_amortized_nodes(), 1 if i == 98 else 2)
                continue
            if i in INLINE_FORMULAS:
                formula, nodes = INLINE_FORMULAS[i]
                add_mixed_formula_paragraph(template, sp, formula, nodes)
                continue
            if i == 94:
                # This paragraph contains two inline equations; keep prose unchanged around two editable OMML objects.
                f1 = "C_total = C_commit + C_DA + p_challenge × C_challenge"
                f2 = "C_amortized = C_total / batch_size"
                t = text
                a, b = t.index(f1), t.index(f1) + len(f1)
                c, d = t.index(f2), t.index(f2) + len(f2)
                p = template.add_paragraph(style="Normal")
                p.add_run(t[:a])
                p._p.append(omath(cost_total_nodes()))
                p.add_run(t[b:c])
                p._p.append(omath(cost_amortized_nodes()))
                p.add_run(t[d:])
                continue
            if re.match(r"^\d+  ", text):
                add_source_paragraph(template, sp, "一级标题", heading=True)
            elif re.match(r"^\d+\.\d+  ", text):
                add_source_paragraph(template, sp, "二级标题", heading=True)
            elif re.match(r"^\d+\.\d+\.\d+  ", text):
                add_source_paragraph(template, sp, "三级标题", heading=True)
            elif text.startswith("定义 "):
                p = add_source_paragraph(template, sp, "定义定理")
                set_keep(p, with_next=True, together=True)
            elif text == "参考文献：":
                p = add_source_paragraph(template, sp, "固定6")
                set_keep(p, with_next=True, together=True)
            elif 131 <= i <= 142:
                add_source_paragraph(template, sp, "文献列表")
            elif 143 <= i <= 146:
                add_source_paragraph(template, sp, "作者简介")
            elif i in (115, 116, 120, 121):
                p = add_source_paragraph(template, sp, "图表名称")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
            else:
                add_source_paragraph(template, sp, "Normal")
        elif child.tag == qn("w:tbl"):
            ti = source_table_index[child]
            if ti < 4:
                add_figure(template, source, ti, figures[ti])
            else:
                insert_source_table(template, source.tables[ti], ti)

    configure_math_and_fields(template)
    clear_template_example_footers(template)
    remove_front_instruction_shape(template)
    template.save(WORK)
    print(json.dumps({"output": str(WORK), "figures": [str(x[0]) for x in figures]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
