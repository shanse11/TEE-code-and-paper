#!/usr/bin/env python3
from pathlib import Path
import hashlib
import json
import zipfile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree


ROOT = Path(__file__).resolve().parents[3]
RUN = ROOT / "outputs/submission_format_20260915_1700"
SOURCE = ROOT / "Hybrid_TEE_Rollup修改版.docx"
FINAL = RUN / "Hybrid_TEE_Rollup_计算机应用_投稿待确认版.docx"
OUT = RUN / "qa/content_verification.json"
NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
}
FORMULA_PARAGRAPHS = {48, 51, 52, 54, 55, 77, 94, 98, 99}


def sha(data):
    return hashlib.sha256(data).hexdigest()


def p_text(el):
    chunks = []
    for node in el.xpath(".//w:r/w:t | .//w:r/w:tab | .//w:r/w:br | .//w:r/w:cr", namespaces=NS):
        local = etree.QName(node).localname
        if local == "t":
            chunks.append(node.text or "")
        elif local == "tab":
            chunks.append("\t")
        else:
            chunks.append("\n")
    return "".join(chunks)


def nonformula_source_sequence(doc):
    pmap = {p._element: i for i, p in enumerate(doc.paragraphs)}
    tmap = {t._element: i for i, t in enumerate(doc.tables)}
    out = []
    for child in doc._element.body:
        if child.tag == qn("w:p"):
            i = pmap[child]
            if i in FORMULA_PARAGRAPHS:
                continue
            text = doc.paragraphs[i].text
            if text:
                out.append(text)
        elif child.tag == qn("w:tbl"):
            ti = tmap[child]
            if ti < 4:
                out.extend(p.text for p in doc.tables[ti].cell(0, 0).paragraphs if p.text)
    return out


def nonformula_final_sequence(path):
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    out = []
    first = True
    for p in root.xpath("/w:document/w:body/w:p", namespaces=NS):
        if first:
            first = False
            continue  # fixed template DOI/header paragraph
        if p.xpath(".//m:oMath", namespaces=NS):
            continue
        text = p_text(p)
        if text:
            out.append(text)
    return out


def table_cells(table):
    return [["\n".join(p.text for p in cell.paragraphs) for cell in row.cells] for row in table.rows]


def package_math_and_media(path):
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
        equations = []
        for eq in root.xpath(".//m:oMath", namespaces=NS):
            equations.append({
                "text": "".join(eq.xpath(".//m:t/text()", namespaces=NS)),
                "subscripts": len(eq.xpath(".//m:sSub", namespaces=NS)),
                "fractions": len(eq.xpath(".//m:f", namespaces=NS)),
            })
        media = {}
        for n in z.namelist():
            if n.startswith("word/media/"):
                media[n] = sha(z.read(n))
        xml_names = [n for n in z.namelist() if n.startswith("word/") and n.endswith(".xml")]
        revisions = 0
        comments = 0
        template_fragments = []
        prohibited = [
            "摘要用于提示研究对象", "第1作者1,2", "NAME Name1,2",
            "国家863计划项目", "★模板使用说明", "收稿日期: 20yy-mm-dd",
            "张珊珊(1967-)", "keyword; keyword; keyword",
        ]
        for n in xml_names:
            tree = etree.fromstring(z.read(n))
            revisions += sum(len(tree.xpath(f".//w:{tag}", namespaces=NS)) for tag in ("ins", "del", "moveFrom", "moveTo"))
            comments += len(tree.xpath(".//w:commentReference | .//w:commentRangeStart | .//w:commentRangeEnd", namespaces=NS))
            text = "".join(tree.xpath(".//w:t/text()", namespaces=NS))
            for frag in prohibited:
                if frag in text:
                    template_fragments.append({"part": n, "fragment": frag})
        inline = len(root.xpath(".//wp:inline", namespaces=NS))
        anchor = len(root.xpath(".//wp:anchor", namespaces=NS))
        heading_num_ids = root.xpath(
            ".//w:p[w:pPr/w:pStyle[@w:val='64' or @w:val='50' or @w:val='49']]/w:pPr/w:numPr/w:numId/@w:val",
            namespaces=NS,
        )
        doi_text = p_text(root.xpath("/w:document/w:body/w:p[1]", namespaces=NS)[0])
    return equations, media, revisions, comments, template_fragments, inline, anchor, heading_num_ids, doi_text


def source_media_hashes(path):
    with zipfile.ZipFile(path) as z:
        return sorted(sha(z.read(n)) for n in z.namelist() if n.startswith("word/media/"))


def main():
    source = Document(SOURCE)
    final = Document(FINAL)
    source_seq = nonformula_source_sequence(source)
    final_seq = nonformula_final_sequence(FINAL)
    equations, media, revisions, comments, fragments, inline, anchor, num_ids, doi = package_math_and_media(FINAL)
    expected_math_text = [
        "cs = (id, s, F, T)",
        "M = (S, E, δ)",
        "δ: S × E → S",
        "Recover(cs) = cs'",
        "F(cs) = F(Recover(cs))",
        "Facts(csbefore) = Facts(csafter)",
        "Ctotal = Ccommit + CDA + pchallenge × Cchallenge",
        "Camortized = Ctotalbatchsize",
        "Ctotal = Ccommit + CDA + pchallenge × Cchallenge",
        "Camortized = Ctotalbatchsize",
    ]
    expected_subs = [0, 0, 0, 0, 0, 2, 5, 3, 5, 3]
    expected_fracs = [0, 0, 0, 0, 0, 0, 0, 1, 0, 1]
    source_imgs = source_media_hashes(SOURCE)
    final_img_hashes = sorted(media.values())
    checks = {
        "nonformula_text_sequence_exact": source_seq == final_seq,
        "nonformula_source_count": len(source_seq),
        "nonformula_final_count": len(final_seq),
        "tables_exact": [table_cells(t) for t in source.tables[4:]] == [table_cells(t) for t in final.tables],
        "table_count": len(final.tables),
        "source_image_hashes_all_present": all(h in final_img_hashes for h in source_imgs),
        "source_image_hashes": source_imgs,
        "final_media_hashes": final_img_hashes,
        "paper_images_inline_count": inline,
        "template_fixed_anchor_count": anchor,
        "equation_count": len(equations),
        "equation_text_sequence": [e["text"] for e in equations],
        "equation_text_match": [e["text"] for e in equations] == expected_math_text,
        "equation_subscript_structure_match": [e["subscripts"] for e in equations] == expected_subs,
        "equation_fraction_structure_match": [e["fractions"] for e in equations] == expected_fracs,
        "revisions": revisions,
        "comment_markup": comments,
        "template_fragments": fragments,
        "heading_num_ids": num_ids,
        "doi_placeholder_preserved": "doi:10.11772/j.issn.1001-9081.yyyymmnnnn" in doi,
        "reference_count_source": sum(1 for p in source.paragraphs if p.text.startswith("[") and p.text[1:2].isdigit()),
        "reference_count_final": sum(1 for p in final.paragraphs if p.text.startswith("[") and p.text[1:2].isdigit()),
        "author_contact_text_present": all(s in "\n".join(p.text for p in final.paragraphs) for s in (
            "联系人：杨帆", "通信地址：上海市普陀区中山北路3663号华东师范大学数据科学与工程学院，200062",
            "电子信箱：2991777340@qq.com", "电话：13319693336",
        )),
    }
    critical = [
        checks["nonformula_text_sequence_exact"], checks["tables_exact"], checks["source_image_hashes_all_present"],
        checks["equation_text_match"], checks["equation_subscript_structure_match"], checks["equation_fraction_structure_match"],
        checks["revisions"] == 0, checks["comment_markup"] == 0, not checks["template_fragments"],
        set(checks["heading_num_ids"]) == {"0"}, checks["doi_placeholder_preserved"],
        checks["reference_count_source"] == checks["reference_count_final"] == 12,
        checks["author_contact_text_present"], checks["table_count"] == 2, checks["paper_images_inline_count"] == 4,
    ]
    checks["critical_pass"] = all(critical)
    OUT.write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(checks, ensure_ascii=False, indent=2))
    if not checks["critical_pass"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
