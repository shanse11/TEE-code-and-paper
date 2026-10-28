#!/usr/bin/env python3
from copy import deepcopy
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[3]
RUN = ROOT / "outputs/submission_format_20260915_1700"
template = Document(RUN / "work/投稿模版_WPS.docx")
source = Document(ROOT / "Hybrid_TEE_Rollup修改版.docx")

# Replace only one documented body sample slot.
target = template.paragraphs[27]
for child in list(target._element):
    if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
        target._element.remove(child)
target.add_run(source.paragraphs[12].text)

# Copy one real Word table as the complex-object smoke test.
anchor = template.paragraphs[28]._element
anchor.addnext(deepcopy(source.tables[4]._element))

out = RUN / "work/migration_smoke.docx"
template.save(out)
print(out)
