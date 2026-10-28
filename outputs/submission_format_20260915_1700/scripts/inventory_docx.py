#!/usr/bin/env python3
import argparse
import hashlib
import json
import re
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from lxml import etree

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS = {"w": W, "m": M, "wp": WP}


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def xml_text(el):
    return "".join(el.xpath(".//w:t/text() | .//w:delText/text()", namespaces=NS))


def package_parts(path):
    out = []
    with zipfile.ZipFile(path) as z:
        for info in z.infolist():
            data = z.read(info.filename)
            out.append({"path": info.filename, "size": len(data), "sha256": hashlib.sha256(data).hexdigest()})
    return out


def inspect(path):
    doc = Document(path)
    with zipfile.ZipFile(path) as z:
        parts = set(z.namelist())
        xml_parts = {}
        for name in parts:
            if name.startswith("word/") and name.endswith(".xml"):
                try:
                    xml_parts[name] = etree.fromstring(z.read(name))
                except etree.XMLSyntaxError:
                    pass
        root = xml_parts["word/document.xml"]
        revisions = Counter()
        for tag in ("ins", "del", "moveFrom", "moveTo", "pPrChange", "rPrChange", "tblPrChange"):
            revisions[tag] = len(root.xpath(f".//w:{tag}", namespaces=NS))
        comments = None
        if "word/comments.xml" in xml_parts:
            comments = len(xml_parts["word/comments.xml"].xpath(".//w:comment", namespaces=NS))
        media = []
        for name in sorted(p for p in parts if p.startswith("word/media/")):
            data = z.read(name)
            media.append({"path": name, "size": len(data), "sha256": hashlib.sha256(data).hexdigest()})
        math = []
        for name, tree in xml_parts.items():
            for i, eq in enumerate(tree.xpath(".//m:oMath | .//m:oMathPara", namespaces=NS)):
                math.append({"part": name, "index": i, "text": "".join(eq.xpath(".//m:t/text()", namespaces=NS))})
        drawings = {
            "inline": sum(len(t.xpath(".//wp:inline", namespaces=NS)) for t in xml_parts.values()),
            "anchor": sum(len(t.xpath(".//wp:anchor", namespaces=NS)) for t in xml_parts.values()),
        }

    paragraphs = []
    for i, p in enumerate(doc.paragraphs):
        text = p.text
        runs = []
        for r in p.runs:
            vert = None
            if r._element.rPr is not None and r._element.rPr.vertAlign is not None:
                vert = r._element.rPr.vertAlign.val
            runs.append({
                "text": r.text,
                "bold": r.bold,
                "italic": r.italic,
                "underline": bool(r.underline) if r.underline is not None else None,
                "vertAlign": vert,
                "style": r.style.name if r.style else None,
            })
        paragraphs.append({
            "index": i,
            "style": p.style.name if p.style else None,
            "text": text,
            "runs": runs,
        })
    tables = []
    for ti, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append(["\n".join(p.text for p in cell.paragraphs) for cell in row.cells])
        tables.append({"index": ti, "rows": rows})
    sections = []
    for i, s in enumerate(doc.sections):
        sections.append({
            "index": i,
            "page_width": s.page_width,
            "page_height": s.page_height,
            "top_margin": s.top_margin,
            "bottom_margin": s.bottom_margin,
            "left_margin": s.left_margin,
            "right_margin": s.right_margin,
            "header_distance": s.header_distance,
            "footer_distance": s.footer_distance,
            "different_first_page": s.different_first_page_header_footer,
            "header_text": [p.text for p in s.header.paragraphs],
            "footer_text": [p.text for p in s.footer.paragraphs],
        })
    styles = []
    for st in doc.styles:
        if st.type in (1, 2):
            styles.append({"name": st.name, "type": int(st.type), "base": st.base_style.name if st.base_style else None})
    return {
        "path": str(Path(path).resolve()),
        "sha256": sha256(path),
        "paragraphs": paragraphs,
        "tables": tables,
        "sections": sections,
        "styles": styles,
        "media": media,
        "math": math,
        "drawings": drawings,
        "revisions": revisions,
        "comments": comments,
        "package_parts": package_parts(path),
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    result = inspect(args.docx)
    Path(args.out).write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(json.dumps({
        "paragraphs": len(result["paragraphs"]),
        "tables": len(result["tables"]),
        "sections": len(result["sections"]),
        "styles": len(result["styles"]),
        "media": len(result["media"]),
        "math": len(result["math"]),
        "drawings": result["drawings"],
        "revisions": result["revisions"],
        "comments": result["comments"],
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
