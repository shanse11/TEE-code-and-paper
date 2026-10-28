#!/usr/bin/env python3
"""模板小样：一段正文、一个可编辑 OMML 公式和一张真实 SVG。"""
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from img_util import make_drawing

RUN = Path(__file__).resolve().parents[1]
TEMPLATE = RUN / 'work/投稿模版_WPS.docx'
SOURCE = RUN / 'input_backup/Hybrid_TEE_Rollup修改版.docx'
SVG = RUN / 'input_backup/fig3_byte_size.svg'
OUT = RUN / 'work/migration_smoke.docx'

doc = Document(TEMPLATE)
source = Document(SOURCE)
body = doc._element.body
final_sect = deepcopy(body.sectPr)
for child in list(body):
    body.remove(child)
body.append(final_sect)

doc.add_paragraph(source.paragraphs[12].text, style='Normal')

formula = doc.add_paragraph(style='MTDisplayEquation')
om = OxmlElement('m:oMath')
for text in ('C', ' = ', 'C', ' + ', 'C'):
    mr = OxmlElement('m:r')
    mt = OxmlElement('m:t')
    mt.text = text
    mr.append(mt)
    om.append(mr)
formula._p.append(om)

drawing_template = None
for table in source.tables[:4]:
    drawing_template = table._element.find('.//' + qn('w:drawing'))
    if drawing_template is not None:
        break
if drawing_template is None:
    raise RuntimeError('source drawing template missing')
image_p = doc.add_paragraph()
image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = image_p.add_run()
run._r.append(make_drawing(doc, str(SVG), int(15 / 2.54 * 914400),
                           drawing_template, name='smoke_fig3_svg'))

doc.save(OUT)
from build_submission import postprocess
postprocess(str(OUT))
print(OUT)
