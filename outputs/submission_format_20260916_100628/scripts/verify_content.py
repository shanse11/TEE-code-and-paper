#!/usr/bin/env python3
"""对本次排版稿执行内容、对象和 OOXML 结构核验。"""
from pathlib import Path
import hashlib
import json
import re
import zipfile

from docx import Document
from lxml import etree

RUN = Path(__file__).resolve().parents[1]
SOURCE = RUN / "input_backup/Hybrid_TEE_Rollup修改版.docx"
FINAL = RUN / "Hybrid_TEE_Rollup_计算机应用_投稿待确认版.docx"
OUT = RUN / "qa/content_verification.json"
NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "asvg": "http://schemas.microsoft.com/office/drawing/2016/SVG/main",
}
FORMULA_PARAGRAPHS = {48, 51, 52, 54, 55, 77, 94, 98, 99, 100}
FRONT_CHANGED = {1, 3, 4, 5, 6, 7, 9, 10}
SECTION_REF_CHANGED = {108}
HEADING_RE = re.compile(r"^\d+(?:\.\d+)*\s+")


def sha(data):
    return hashlib.sha256(data).hexdigest()


def table_cells(table):
    return [["\n".join(p.text for p in cell.paragraphs) for cell in row.cells]
            for row in table.rows]


def unchanged_source_paragraphs(source):
    out = []
    for i, p in enumerate(source.paragraphs):
        if not p.text or i in FORMULA_PARAGRAPHS | FRONT_CHANGED | SECTION_REF_CHANGED:
            continue
        if HEADING_RE.match(p.text):
            continue
        out.append(p.text)
    return out


def is_subsequence(items, haystack):
    pos, missing = 0, []
    for item in items:
        while pos < len(haystack) and haystack[pos] != item:
            pos += 1
        if pos == len(haystack):
            missing.append(item)
        else:
            pos += 1
    return not missing, missing


def package_checks(path):
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        root = etree.fromstring(z.read("word/document.xml"))
        equations = []
        for eq in root.xpath(".//m:oMath", namespaces=NS):
            equations.append({
                "text": "".join(eq.xpath(".//m:t/text()", namespaces=NS)),
                "subscripts": len(eq.xpath(".//m:sSub", namespaces=NS)),
                "fractions": len(eq.xpath(".//m:f", namespaces=NS)),
            })
        media = {n: sha(z.read(n)) for n in names if n.startswith("word/media/")}
        revisions = sum(len(root.xpath(f".//w:{tag}", namespaces=NS))
                        for tag in ("ins", "del", "moveFrom", "moveTo"))
        comments = sum(len(root.xpath(f".//w:{tag}", namespaces=NS))
                       for tag in ("commentReference", "commentRangeStart", "commentRangeEnd"))
        text_parts = []
        for name in names:
            if name.startswith("word/") and name.endswith(".xml"):
                try:
                    tree = etree.fromstring(z.read(name))
                    text_parts.extend(tree.xpath(".//w:t/text()", namespaces=NS))
                except etree.XMLSyntaxError:
                    pass
        text = "".join(text_parts)
        prohibited = ["摘要用于提示研究对象", "第1作者1,2", "NAME Name1,2",
                      "国家863计划项目", "★模板使用说明", "收稿日期: 20yy-mm-dd",
                      "20yy-mm-dd", "第1作者(出生年—)", "张珊珊(1967-)",
                      "keyword; keyword; keyword", "论文模板"]
        settings = etree.fromstring(z.read("word/settings.xml"))
        core = etree.fromstring(z.read("docProps/core.xml"))
        return {
            "equations": equations,
            "media": media,
            "revisions": revisions,
            "comment_markup": comments,
            "template_fragments": [frag for frag in prohibited if frag in text],
            "inline": len(root.xpath(".//wp:inline", namespaces=NS)),
            "anchor": len(root.xpath(".//wp:anchor", namespaces=NS)),
            "svg_blip": len(root.xpath(".//asvg:svgBlip", namespaces=NS)),
            "heading_num_ids": root.xpath(
                ".//w:p[w:pPr/w:pStyle[@w:val='64' or @w:val='50' or @w:val='49']]"
                "/w:pPr/w:numPr/w:numId/@w:val", namespaces=NS),
            "src_rect_attribute_count": len(root.xpath(".//*[local-name()='srcRect']/@*")),
            "custom_properties_absent": "docProps/custom.xml" not in names,
            "comments_parts_absent": not any("comments" in n for n in names),
            "ole_parts_absent": not any(n.startswith("word/embeddings/") for n in names),
            "track_revisions_absent": not settings.xpath(".//w:trackRevisions", namespaces=NS),
            "update_fields_enabled": bool(settings.xpath(".//w:updateFields[@w:val='true']", namespaces=NS)),
            "core_text": "".join(core.itertext()),
            "package_text": text,
        }


def border_checks(tables):
    results = []
    for table in tables:
        xml = table._element
        vertical = xml.xpath(".//w:tcBorders/w:left | .//w:tcBorders/w:right | .//w:tcBorders/w:insideV")
        repeat_header = bool(xml.xpath(".//w:tr[1]/w:trPr/w:tblHeader"))
        cant_split_rows = len(xml.xpath(".//w:tr/w:trPr/w:cantSplit"))
        results.append({"vertical_border_count": len(vertical), "repeat_header": repeat_header,
                        "cant_split_rows": cant_split_rows, "row_count": len(table.rows)})
    return results


def main():
    source, final = Document(SOURCE), Document(FINAL)
    final_texts = [p.text for p in final.paragraphs if p.text]
    unchanged = unchanged_source_paragraphs(source)
    unchanged_ok, missing = is_subsequence(unchanged, final_texts)
    pkg = package_checks(FINAL)
    expected_math = [
        "cs = (id, s, F, T)", "M = (S, E, δ)", "δ: S × E → S",
        "Recover(cs) = cs'", "F(cs) = F( Recover(cs) )",
        "Facts(csbefore) = Facts(csafter)",
        "Ctotal = Ccommit + CDA + pchallenge × Cchallenge",
        "Ccommit", "CDA", "Cchallenge", "pchallenge", "Camortized = Ctotalbatchsize",
        "Ctotal = Ccommit + CDA + pchallenge × Cchallenge", "Camortized = Ctotalbatchsize",
        "Ccommit", "CDA", "Cchallenge", "pchallenge",
    ]
    expected_subs = [0, 0, 0, 0, 0, 2, 5, 1, 1, 1, 1, 3, 5, 3, 1, 1, 1, 1]
    expected_fracs = [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 1, 0, 1, 0, 0, 0, 0]
    fig_names = ((1, "architecture"), (2, "challenge_flow"),
                 (3, "byte_size"), (4, "da_cost"))
    svg_hashes = [sha((RUN / "input_backup" / f"fig{i}_{name}.svg").read_bytes()) for i, name in fig_names]
    png_hashes = [sha((RUN / "input_backup" / f"fig{i}_{name}.png").read_bytes()) for i, name in fig_names]
    media_hashes = list(pkg["media"].values())
    final_all = "\n".join(final_texts) + "\n" + pkg["package_text"]
    expected_cn_kw = "交互式挑战；数据可用性；高频 DApp；成本摊销；故障恢复"
    expected_en_kw = "interactive challenge; data availability; high-frequency DApp; cost amortization; fault recovery"
    expected_headings = ["引言", "相关工作", "系统概览", "核心机制与成本评估", "实验评估", "结论"]
    heading_texts = [p.text for p in final.paragraphs
                     if p.style and p.style.style_id in {"64", "50", "49"}]
    caption_paras = [p for p in final.paragraphs if p.style and p.style.style_id == "84"]
    captions_direct_black = all(
        p.runs and all(r.font.color.rgb is not None and str(r.font.color.rgb) == "000000" for r in p.runs)
        for p in caption_paras)
    table_borders = border_checks(final.tables)
    checks = {
        "unchanged_text_subsequence_exact": unchanged_ok,
        "unchanged_text_count": len(unchanged),
        "unchanged_text_missing": missing,
        "tables_exact_cell_by_cell": [table_cells(t) for t in source.tables[4:]] == [table_cells(t) for t in final.tables],
        "table_count": len(final.tables),
        "table_border_structure": table_borders,
        "keywords_exact": expected_cn_kw in final_texts and expected_en_kw in final_texts,
        "english_title_joined_only": source.paragraphs[6].text.replace("\n", " ") in final_texts,
        "section_reference_updated": "第 3.1.4 节" in final_all and "第 4.1.4 节" not in final_all,
        "heading_texts": heading_texts,
        "caption_count": len(caption_paras), "captions_direct_black": captions_direct_black,
        "top_level_heading_text_match": [h for h in heading_texts if h in expected_headings] == expected_headings,
        "heading_num_ids": pkg["heading_num_ids"],
        "reference_count_source": sum(bool(re.match(r"^\[\d+\]", p.text)) for p in source.paragraphs),
        "reference_count_final": sum(bool(re.match(r"^\[\d+\]", p.text)) for p in final.paragraphs),
        "formula_text_sequence": [e["text"] for e in pkg["equations"]],
        "formula_text_match": [e["text"] for e in pkg["equations"]] == expected_math,
        "formula_subscript_structure_match": [e["subscripts"] for e in pkg["equations"]] == expected_subs,
        "formula_fraction_structure_match": [e["fractions"] for e in pkg["equations"]] == expected_fracs,
        "specified_svg_hashes_present": all(h in media_hashes for h in svg_hashes),
        "fallback_png_hashes_present": all(h in media_hashes for h in png_hashes),
        "svg_blip_count": pkg["svg_blip"], "inline_image_count": pkg["inline"],
        "anchor_image_count": pkg["anchor"], "src_rect_attribute_count": pkg["src_rect_attribute_count"],
        "author_facts_present": all(x in final_all for x in (
            "杨帆1*", "YANG Fan1*", "2991777340@qq.com", "在读的硕士研究生",
            "湖南邵阳", "区块链系统与跨链互操作")),
        "no_fabricated_funding_statement": "无基金" not in final_all and "无资助" not in final_all,
        "no_english_bio_added": "Biography" not in final_all,
        "contribution_not_inserted_without_template_slot": "作者工作量说明" not in final_all,
        "contact_block_preserved": all(source.paragraphs[i].text in final_texts for i in range(143, 147)),
        "revisions": pkg["revisions"], "comment_markup": pkg["comment_markup"],
        "template_fragments": pkg["template_fragments"],
        "custom_properties_absent": pkg["custom_properties_absent"],
        "comments_parts_absent": pkg["comments_parts_absent"], "ole_parts_absent": pkg["ole_parts_absent"],
        "track_revisions_absent": pkg["track_revisions_absent"], "update_fields_enabled": pkg["update_fields_enabled"],
        "metadata_example_absent": "张珊珊" not in pkg["core_text"] and "论文模板" not in pkg["core_text"],
        "doi_placeholder_preserved": "doi:10.11772/j.issn.1001-9081.yyyymmnnnn" in final.paragraphs[0].text,
    }
    critical = [
        checks["unchanged_text_subsequence_exact"], checks["tables_exact_cell_by_cell"], checks["table_count"] == 2,
        all(t["vertical_border_count"] == 0 and t["repeat_header"] and t["cant_split_rows"] == t["row_count"] for t in table_borders),
        checks["keywords_exact"], checks["english_title_joined_only"], checks["section_reference_updated"],
        checks["top_level_heading_text_match"], checks["captions_direct_black"], set(checks["heading_num_ids"]) == {"2"},
        checks["reference_count_source"] == checks["reference_count_final"] == 12,
        checks["formula_text_match"], checks["formula_subscript_structure_match"], checks["formula_fraction_structure_match"],
        checks["specified_svg_hashes_present"], checks["fallback_png_hashes_present"],
        checks["svg_blip_count"] == checks["inline_image_count"] == 4, checks["anchor_image_count"] == 0,
        checks["src_rect_attribute_count"] == 0, checks["author_facts_present"], checks["contact_block_preserved"],
        checks["contribution_not_inserted_without_template_slot"],
        checks["revisions"] == checks["comment_markup"] == 0, not checks["template_fragments"],
        checks["custom_properties_absent"], checks["comments_parts_absent"], checks["ole_parts_absent"],
        checks["track_revisions_absent"], checks["update_fields_enabled"], checks["metadata_example_absent"],
        checks["doi_placeholder_preserved"],
    ]
    checks["critical_pass"] = all(critical)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(checks, ensure_ascii=False, indent=2))
    if not checks["critical_pass"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
