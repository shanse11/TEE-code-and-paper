from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


BASE = Path(__file__).resolve().parent
SRC = BASE / "Hybrid_TEE_Rollup_参考文献格式修订版.docx"
OUT = BASE / "Hybrid_TEE_Rollup_参考文献格式模板修订版.docx"


def set_east_asian_font(run, name: str) -> None:
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def set_ref_run_style(run) -> None:
    run.font.name = "Times New Roman"
    set_east_asian_font(run, "宋体")
    run.font.size = Pt(9)
    run.bold = False
    run.italic = False
    run.font.superscript = False


def set_ref_paragraph_style(paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.left_indent = Pt(18.4)
    fmt.first_line_indent = Pt(-18.4)
    fmt.line_spacing = 1.2
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    paragraph.alignment = None


def set_heading_style(paragraph) -> None:
    paragraph.clear()
    run = paragraph.add_run("参考文献：")
    run.font.size = Pt(11.5)
    run.bold = True
    set_east_asian_font(run, "黑体")
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.left_indent = None
    paragraph.paragraph_format.line_spacing = None
    paragraph.paragraph_format.space_before = None
    paragraph.paragraph_format.space_after = None
    paragraph.alignment = None


def main() -> None:
    doc = Document(str(SRC))
    paras = doc.paragraphs

    ref_idx = None
    for i, p in enumerate(paras):
        if p.text.strip().replace(" ", "") in {"参考文献", "参考文献："}:
            ref_idx = i
            break
    if ref_idx is None:
        raise RuntimeError("Reference heading not found")

    set_heading_style(paras[ref_idx])

    count = 0
    for p in paras[ref_idx + 1 :]:
        if not re.match(r"^\[\d+\]", p.text.strip()):
            continue
        set_ref_paragraph_style(p)
        for run in p.runs:
            set_ref_run_style(run)
        count += 1

    if count == 0:
        raise RuntimeError("No reference entries found")
    doc.save(str(OUT))
    print(OUT)
    print("references", count)


if __name__ == "__main__":
    main()
