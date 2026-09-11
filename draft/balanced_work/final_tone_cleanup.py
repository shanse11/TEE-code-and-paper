from pathlib import Path

from docx import Document


DOCX = Path("E:/项目/project_code/paper_outputs/draft/final_balanced_low_aigc_academic.docx")


def set_text(paragraph, text):
    style = paragraph.style
    alignment = paragraph.alignment
    runs = list(paragraph.runs)
    for run in runs:
        run.clear()
    if runs:
        runs[0].text = text
    else:
        paragraph.add_run(text)
    paragraph.style = style
    paragraph.alignment = alignment


abstract = (
    "高频去中心化应用（DApp）要求低延迟、低链上开销与公开可验证仲裁并存。"
    "本文研究 Hybrid TEE-Rollup 中轻量链上承诺、链外数据可用性（DA）负载与超时挑战的协同，提出 Recoverable Challenge、Evidence-Carrying Compact Commit 和 DA-Aware Cost Model。"
    "Recoverable Challenge 将超时会话恢复为可重放、可结算状态；Evidence-Carrying Compact Commit 绑定状态根、输出哈希、证明哈希、DA 指针与 DA 根；DA-Aware Cost Model 刻画负载规模、批处理摊销和 DA 路径选择。"
    "基于 Python 原型、本地 EVM gas 测量和 Sepolia 部署记录，实验表明紧凑提交平均约为 406 字节；负载 8192、批大小 1000 时，模块化 DA 采样路径的摊销成本估算较全链上调用数据下降 96.33%；合成超时场景中，恢复路径将挑战成功率由不可恢复路径的 0% 提升至 100%。"
    "结论限于模拟 TEE、可验证 DA 注册表和配置化成本模型。"
)


starts = {
    "超时后挑战流程卡死并非普通工程异常": (
        "超时后挑战流程卡死属于协议活性与仲裁连续性问题。传统挑战机制可以关注错误是否能够被发现，但对于高频 DApp 而言，检测只是争议处理的起点：未完成的挑战会占用状态、延迟结算，并使错误提交停留在不可结算的中间状态。因此，本文以故障发生、恢复、重放、仲裁结算为主线，将紧凑提交作为最小可验证证据入口，使正常路径保持轻量，异常路径仍可进入验证与重放。"
    ),
    "表 2 的作用并非给出能力清单": (
        "表 2 用于界定本文机制与既有系统的关系：已有系统通常已经具备挑战机制或 TEE 快速路径，本文关注的增量在于超时后的挑战活性、紧凑提交到数据可用性字段级证明的绑定关系，以及面向高频工作负载的成本解释。"
    ),
    "上述对比表明，本文机制并非单独追求": (
        "上述对比表明，本文机制的主要贡献在于组合轻量提交、证据绑定和可恢复仲裁路径，使高频场景下的异常处理能够继续推进。"
    ),
    "opML 等工作研究如何通过乐观欺诈证明": (
        "opML 等工作研究如何通过乐观欺诈证明支持机器学习或大模型计算，通常包含二分定位、单步仲裁和链上虚拟机验证。本文借鉴交互式争议定位思想，但未实现完整 FPVM 或真实 ML 执行证明，而将二分重放用于 Hybrid TEE-Rollup 原型中的异常仲裁。"
    ),
    "Dynamic Fraud Proof 关注如何缩短": (
        "Dynamic Fraud Proof 关注无争议场景下的最终性压缩，并在发现争议时动态延迟结算。该方向强调挑战窗口和验证者参与机制的动态调整；本文聚焦更窄的超时后挑战会话恢复问题。"
    ),
}


def main():
    doc = Document(str(DOCX))
    set_text(doc.paragraphs[3], abstract)
    for paragraph in doc.paragraphs:
        for start, replacement in starts.items():
            if paragraph.text.startswith(start):
                set_text(paragraph, replacement)
                break
    doc.save(str(DOCX))
    print(DOCX)
    print("abstract_chars", len(abstract))


if __name__ == "__main__":
    main()
