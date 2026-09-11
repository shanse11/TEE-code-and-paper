from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


DOCX = Path("E:/项目/project_code/paper_outputs/draft/final_balanced_low_aigc_academic.docx")


def set_paragraph_text(paragraph, text):
    style = paragraph.style
    alignment = paragraph.alignment
    runs = list(paragraph.runs)
    for run in runs:
        run.clear()
    if runs:
        runs[0].text = text
    else:
        paragraph.add_run(text)
    paragraph.style = style
    paragraph.alignment = alignment


updates = {
    3: (
        "高频去中心化应用（DApp）要求低延迟、低链上开销与公开可验证仲裁并存。"
        "本文研究 Hybrid TEE-Rollup 中轻量链上承诺、链外数据可用性（Data Availability, DA）负载与超时挑战的协同，提出 Recoverable Challenge、Evidence-Carrying Compact Commit 和 DA-Aware Cost Model。"
        "Recoverable Challenge 将超时会话恢复为可重放、可结算状态；Evidence-Carrying Compact Commit 绑定状态根、输出哈希、证明哈希、DA 指针与 DA 根；DA-Aware Cost Model 刻画负载规模、批处理摊销和 DA 路径选择。"
        "基于 Python 原型、本地以太坊虚拟机（Ethereum Virtual Machine, EVM）gas 测量和 Sepolia 部署记录，实验表明紧凑提交平均约为 406 字节；负载 8192、批大小 1000 时，模块化 DA 采样路径的摊销成本估算较全链上调用数据下降 96.33%；合成超时场景中，恢复路径将挑战成功率由不可恢复路径的 0% 提升至 100%。"
        "结论限于模拟 TEE、可验证 DA 注册表和配置化成本模型。"
    ),
    6: (
        "高频去中心化应用的瓶颈来自低价值高频交互与异常可验证性的耦合。以链上 AI 推理服务为例，用户连续提交提示，执行者在可信执行环境中生成响应并向链上提交紧凑提交；完整提示、响应与证明保存在数据可用性层。"
        "正常路径需要低延迟和低链上开销，异常路径则必须支持公开仲裁：当响应与数据可用性负载不一致，或挑战在中途超时时，争议应继续推进到重放与结算。"
    ),
    7: (
        "现有路线分别覆盖了这一需求的不同侧面。全链上执行公开验证能力强，但高频负载、存储和执行成本直接上链；完全链下执行降低延迟和成本，却缺少用户可依赖的公开仲裁路径；纯 TEE 执行降低正常路径验证成本，但仍受硬件和宿主环境影响。"
        "Hybrid TEE-Rollup 将 TEE 快速路径与链上异常路径结合，其风险也集中在异常路径：一旦恢复机制和证据入口缺失，系统会在最需要仲裁时失去协议连续性。"
    ),
    19: "2.1 Rollup、可信执行环境与数据可用性",
    20: (
        "Rollup 的基本思想是将大量执行过程移至链下完成，并将状态承诺、证明或争议入口提交到链上。"
        "ZK-Rollup 通过有效性证明保证提交状态的正确性，验证属性较强，但证明生成、验证电路和系统实现成本较高。"
        "Optimistic Rollup 默认链下执行结果正确，仅在挑战期内出现有效争议时进行仲裁；其链上负担较轻，但争议窗口和交互过程会引入状态不确定性。"
    ),
    21: (
        "乐观挑战是 Rollup 系统处理异常的重要机制。典型流程包括提交状态承诺、开启挑战、交互式缩小争议范围，并在必要时执行单步重放或链上仲裁。"
        "对于长执行轨迹，二分定位能够将争议定位复杂度从线性扫描降低到对数轮次；挑战过程本身也可能遭遇超时、应答缺失或证据不可用，因此活性和恢复性需要进入协议设计。"
    ),
    22: (
        "数据可用性关注链下或链外数据能否被验证者获取。模块化区块链和惰性账本等路线将数据可用性从执行验证中解耦，使共识层主要负责数据排序和可用性保证，执行与验证由独立的 Rollup 或客户端完成。"
        "当执行者将完整负载放在数据可用性层而非链上时，链上提交必须与数据可用性条目形成可验证绑定。"
    ),
    26: (
        "范围说明。本文威胁模型不覆盖真实硬件侧信道、远程证明供应链故障、真实 DA 网络经济安全、主网最大可提取价值以及拥堵导致的费用波动；当前证据支撑协议状态连续性、证据绑定关系、配置化成本趋势和合约层对应关系。"
    ),
    34: "在本文原型中，模拟可信执行环境用于建立“链下执行摘要—链上提交—异常挑战”的接口关系；真实 TEE 安全性在第 8 章讨论。",
    39: (
        "验证层负责处理异常路径。当前原型实现挑战打开、响应、步骤、恢复、重放和结算等状态推进；争议出现后，验证层通过证据加载、二分定位、单步重放和仲裁结算，将异常状态推进为可解释的协议结果。"
    ),
    94: (
        "从对抗角度看，提交者可能替换负载字段、隐藏数据可用性条目、提交与证明不一致的输出，或在争议过程中延迟响应。紧凑提交的作用在于锁定链上最小承诺所需的证据字段，使后续挑战具备重放准备性，而不仅是节省链上字节。"
    ),
    97: (
        "模型比较四类 DA 路径：全链上调用数据、紧凑外部 DA、紧凑类 EIP-4844 方案和紧凑模块化 DA 采样。该比较作为研究原型的成本建模工具，用于刻画不同数据发布路径的结构性差异。"
    ),
    101: (
        "实验评估围绕机制可运行性与对比趋势展开。当前实验使用 Python 原型、模拟 TEE、合成负载、可验证 DA 注册表、本地 EVM gas 测量和 Sepolia 测试网部署记录。"
        "RQ1 关注紧凑提交是否在保留证据关联的同时隔离负载增长；RQ2 关注批处理摊销与 DA 路径选择的影响；RQ3 关注恢复机制是否改善超时场景下的挑战活性；RQ4 关注不同故障类型能否被检测、恢复、重放或罚没；RQ5 关注 Solidity 合约、本地 EVM 实测与 Sepolia 部署能否支撑链上对应物。"
    ),
    104: (
        "其中，C_commit 表示链上紧凑提交或完整负载提交成本，C_DA 表示不同 DA 路径下的数据成本，C_challenge 表示争议发生时挑战、恢复、重放和结算成本，p_challenge 表示争议发生概率。"
    ),
    142: (
        "上述对比表明，本文机制并非单独追求成本下降或检测率提升，而是组合轻量提交、证据绑定和可恢复仲裁路径，使高频场景下的异常处理能够继续推进。"
    ),
    155: "TEEROLLUP 提出了利用异构 TEE 降低 Rollup 验证成本的系统设计。本文受其思路启发，但聚焦异常场景下的挑战恢复和紧凑提交的证据绑定。",
    158: "LazyLedger、模块化 DA 以及轻客户端相关工作强调数据可用性与执行验证解耦。本文使用可验证 DA 注册表和配置化建模分析紧凑提交加数据可用性验证的成本趋势。",
    162: (
        "高频去中心化应用的核心困难在于低成本交互与异常可验证之间的持续权衡。本文在 Hybrid TEE-Rollup 背景下研究一个更窄的问题：在不重做完整 Rollup 系统的前提下，如何让轻量提交、数据可用性证据和可恢复挑战形成可运行的研究原型。"
        "现有证据显示，Python 原型跑通了紧凑提交、数据可用性证明、挑战、恢复、重放与结算；对比型实验区分了检测与完成、压缩与证据绑定；本地 EVM 和 Sepolia 证据说明协议路径具有合约对应物。"
    ),
    163: (
        "从理论层面看，本文的主要价值在于将挑战恢复从工程操作提升为具有协议语义的独立机制，提出并论证仲裁连续性这一协议性质。第 4.4 节系统化了挑战事实不变性、安全性与活性的分离、仲裁连续性和恢复操作正确性四项性质，为该机制进入更复杂的 Rollup 协议栈提供了理论参照。"
    ),
}


cell_replacements = {
    "乐观 TEE 汇总 / OTR": "乐观 TEE Rollup / OTR",
    "利用异构 TEE 降低 Rollup 验证成本，并保留挑战和 DA 惩罚机制": "利用异构 TEE 降低 Rollup 验证成本，并保留挑战和 DA 惩罚机制",
}


def polish():
    doc = Document(str(DOCX))

    for idx, text in updates.items():
        set_paragraph_text(doc.paragraphs[idx], text)

    # Remove stale duplicate heading left from earlier indexed edits, if present.
    for paragraph in list(doc.paragraphs):
        if paragraph.text.strip() == "2.1 汇总、可信执行环境与数据可用性":
            paragraph._element.getparent().remove(paragraph._element)
            break

    # Targeted Rollup terminology cleanup outside non-blockchain uses such as "汇总文件".
    paragraph_terms = {
        "完整汇总系统": "完整 Rollup 系统",
        "汇总协议栈": "Rollup 协议栈",
        "汇总系统": "Rollup 系统",
        "独立的汇总": "独立的 Rollup",
        "汇总验证成本": "Rollup 验证成本",
        "零知识汇总": "ZK-Rollup",
        "乐观汇总": "Optimistic Rollup",
    }
    for paragraph in doc.paragraphs:
        text = paragraph.text
        for old, new in paragraph_terms.items():
            text = text.replace(old, new)
        if text != paragraph.text:
            set_paragraph_text(paragraph, text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text = paragraph.text
                    for old, new in cell_replacements.items():
                        text = text.replace(old, new)
                    for old, new in paragraph_terms.items():
                        text = text.replace(old, new)
                    if text != paragraph.text:
                        set_paragraph_text(paragraph, text)
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(str(DOCX))
    print(DOCX)
    print("abstract_chars", len(doc.paragraphs[3].text))


if __name__ == "__main__":
    polish()
