from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


DOCX = Path("E:/项目/project_code/paper_outputs/draft/final_balanced_low_aigc_academic.docx")


def set_paragraph_text(paragraph, text):
    style = paragraph.style
    alignment = paragraph.alignment
    runs = list(paragraph.runs)
    for run in runs:
        run.clear()
    if runs:
        runs[0].text = text
    else:
        paragraph.add_run(text)
    paragraph.style = style
    paragraph.alignment = alignment


def has_drawing(paragraph):
    return bool(paragraph._p.xpath(".//w:drawing"))


def set_cell_margins(table, margin_twips=90):
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_mar = tc_pr.first_child_found_in("w:tcMar")
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for side in ("top", "left", "bottom", "right"):
                node = tc_mar.find(qn(f"w:{side}"))
                if node is None:
                    node = OxmlElement(f"w:{side}")
                    tc_mar.append(node)
                node.set(qn("w:w"), str(margin_twips))
                node.set(qn("w:type"), "dxa")


def set_table_width_pct(table, pct=95):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(pct * 50))
    tbl_w.set(qn("w:type"), "pct")


def polish():
    doc = Document(str(DOCX))

    # 摘要首次缩写补全，不改变实验事实。
    abstract = doc.paragraphs[3].text
    abstract = abstract.replace(
        "高频 DApp 的矛盾并不只在吞吐量。",
        "高频去中心化应用（Decentralized Application, DApp）的矛盾并不只在吞吐量。",
    )
    abstract = abstract.replace(
        "完整负载位于 DA 层",
        "完整负载位于数据可用性（Data Availability, DA）层",
    )
    abstract = abstract.replace(
        "本地 EVM gas 测量",
        "本地以太坊虚拟机（Ethereum Virtual Machine, EVM）gas 测量",
    )
    set_paragraph_text(doc.paragraphs[3], abstract)

    # 第 8 章开头拆清边界说明与相关工作，避免混写。
    set_paragraph_text(
        doc.paragraphs[144],
        "本文结论应理解为机制层研究原型结论。下面先从执行环境、数据可用性、成本估算和链上实现四个方面集中说明证据边界，随后单列相关工作，以避免将实验边界与文献定位混写。",
    )

    # 表 3 的标题保持为表题，说明文字留给正文上下文。
    if doc.paragraphs[89].text.startswith("表 3 汇总了"):
        set_paragraph_text(doc.paragraphs[89], "表 3 协议不变量与保持性质")

    # 表 2 后的说明句不是表题，恢复为正文格式。
    if doc.paragraphs[16].text.startswith("表 2 的作用"):
        doc.paragraphs[16].style = doc.styles["Normal"]
        doc.paragraphs[16].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # 修复个别中英文空格。
    for paragraph in doc.paragraphs:
        text = paragraph.text.replace("链上DA Merkle 根", "链上 DA Merkle 根")
        if text != paragraph.text:
            set_paragraph_text(paragraph, text)

    # 图题居中；图片所在段落居中；适度放大图片但不超过页面可读宽度。
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("图 "):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(6)
        elif re.match(r"^表\s*\d+\s+(?!的作用)", text):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(3)
        if has_drawing(paragraph):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    max_width = Inches(5.65)
    for shape in doc.inline_shapes:
        old_w, old_h = shape.width, shape.height
        target_w = min(int(old_w * 1.08), max_width)
        if target_w > old_w:
            shape.width = target_w
            shape.height = int(old_h * target_w / old_w)

    # 表格居中、适当扩展到可用宽度，增加单元格内边距，减少拥挤感。
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        set_table_width_pct(table, 95)
        set_cell_margins(table, 100)
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)

    doc.save(str(DOCX))


if __name__ == "__main__":
    polish()
    print(DOCX)
