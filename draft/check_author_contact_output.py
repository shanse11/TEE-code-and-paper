from __future__ import annotations

import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document

P = Path(__file__).resolve().parent / "论文终版" / "Hybrid_TEE_Rollup_作者信息与联系方式修订版.docx"
doc = Document(str(P))

print("paragraphs", len(doc.paragraphs))
print("FIRST")
for i, p in enumerate(doc.paragraphs[:14], 1):
    print(f"{i:03d}: {p.text!r}")
print("TAIL")
for i, p in list(enumerate(doc.paragraphs, 1))[-20:]:
    print(f"{i:03d}: {p.text!r}")

text = "\n".join(p.text for p in doc.paragraphs)
print("has_contact", all(s in text for s in ["联系人：杨帆", "2991777340@qq.com", "13319693336"]))
print("has_placeholders", any(s in text for s in ["请您在文后给出", "通信方式如有变化", "作者名（出生年", "1 寸数字照片", "作者名1", "作者名2", "单位全名"]))
print("ref_count", len(re.findall(r"^\[\d+\]", text, re.M)))
print("missing_refs", sorted(set(map(int, re.findall(r"\[(\d+)\]", text))) - set(range(1, 13))))

with zipfile.ZipFile(P) as z:
    print("zip_bad", z.testzip())
    names = z.namelist()
    print("duplicate_parts", sorted(n for n in set(names) if names.count(n) > 1))
