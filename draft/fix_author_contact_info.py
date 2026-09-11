from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


BASE = Path(__file__).resolve().parent / "论文终版"
SRC = BASE / "Hybrid_TEE_Rollup.docx"
OUT = BASE / "Hybrid_TEE_Rollup_作者信息与联系方式修订版.docx"


def set_east_asian_font(run, name: str) -> None:
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def clear_para(paragraph) -> None:
    paragraph.clear()


def add_text_run(paragraph, text: str, font: str, size: float, east_asia: str | None = None, superscript: bool = False):
    run = paragraph.add_run(text)
    run.font.name = font
    if east_asia:
        set_east_asian_font(run, east_asia)
    run.font.size = Pt(size)
    run.font.superscript = superscript
    return run


def set_spacing(paragraph, before: float = 0, after: float = 0, line_spacing=1) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing


def insert_paragraph_after(paragraph, text: str = ""):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p = new_p
    if text:
        new_para.add_run(text)
    return new_para


def insert_after(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return paragraph._parent.paragraphs[-1]._parent(new_p, paragraph._parent)


def paragraph_after(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    return Paragraph(new_p, paragraph._parent)


def format_chinese_author(paragraph) -> None:
    clear_para(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(paragraph, 0, 0, 1)
    add_text_run(paragraph, "杨帆", "KaiTi", 12, "楷体")
    add_text_run(paragraph, "1", "Times New Roman", 9, "Times New Roman", superscript=True)


def format_chinese_unit(paragraph) -> None:
    clear_para(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(paragraph, 0, 0, 1)
    add_text_run(paragraph, "1. 华东师范大学 数据科学与工程学院，上海 200062", "KaiTi", 10.5, "楷体")


def format_english_author(paragraph) -> None:
    clear_para(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(paragraph, 0, 0, 1)
    add_text_run(paragraph, "YANG Fan", "Times New Roman", 12, None)
    add_text_run(paragraph, "1", "Times New Roman", 9, None, superscript=True)


def format_english_unit(paragraph) -> None:
    clear_para(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(paragraph, 0, 0, 1)
    add_text_run(
        paragraph,
        "1. School of Data Science and Engineering, East China Normal University, Shanghai 200062, China",
        "Times New Roman",
        10.5,
        None,
    )


def format_contact(paragraph, text: str) -> None:
    clear_para(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_spacing(paragraph, 0, 0, 1)
    paragraph.paragraph_format.left_indent = Pt(0)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.right_indent = Pt(0)
    run = add_text_run(paragraph, text, "Times New Roman", 9, "宋体")
    run.bold = False


def main() -> None:
    doc = Document(str(SRC))
    paras = doc.paragraphs

    if len(paras) < 9:
        raise RuntimeError("Unexpected document structure")

    format_chinese_author(paras[1])
    if not paras[2].text.strip().startswith("1. 华东师范大学"):
        cn_unit = paragraph_after(paras[1])
        format_chinese_unit(cn_unit)
    else:
        format_chinese_unit(paras[2])

    # Refresh paragraph list after insertion.
    paras = doc.paragraphs
    english_author_idx = None
    for i, p in enumerate(paras[:15]):
        if p.text.strip() in {"YANG Fan", "YANG Fan1", "YANG Fan¹"}:
            english_author_idx = i
            break
    if english_author_idx is None:
        for i, p in enumerate(paras[:15]):
            if "YANG Fan" in p.text:
                english_author_idx = i
                break
    if english_author_idx is None:
        raise RuntimeError("English author paragraph not found")
    format_english_author(paras[english_author_idx])
    paras = doc.paragraphs
    after_author = paras[english_author_idx + 1]
    if not after_author.text.strip().startswith("1. School of Data Science"):
        en_unit = paragraph_after(paras[english_author_idx])
        format_english_unit(en_unit)
    else:
        format_english_unit(after_author)

    paras = doc.paragraphs
    # Remove known template prompt paragraphs if they exist.
    prompt_fragments = [
        "请您在文后给出以下内容",
        "通信方式如有变化",
        "作者名（出生年",
        "1 寸数字照片",
        "1. 联系人",
        "电子信箱、电话",
    ]
    for p in paras:
        if any(fragment in p.text for fragment in prompt_fragments):
            clear_para(p)

    paras = doc.paragraphs
    last_ref_idx = None
    for i, p in enumerate(paras):
        if p.text.strip().startswith("[12]"):
            last_ref_idx = i
    if last_ref_idx is None:
        raise RuntimeError("Last reference not found")

    # Avoid duplicate contact block if rerun.
    existing_contact = any("联系人：杨帆" in p.text for p in paras[last_ref_idx + 1 :])
    if not existing_contact:
        p1 = paragraph_after(paras[last_ref_idx])
        format_contact(p1, "联系人：杨帆")
        p2 = paragraph_after(p1)
        format_contact(p2, "通信地址：上海市普陀区中山北路3663号华东师范大学数据科学与工程学院，200062")
        p3 = paragraph_after(p2)
        format_contact(p3, "电子信箱：2991777340@qq.com")
        p4 = paragraph_after(p3)
        format_contact(p4, "电话：13319693336")

    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
