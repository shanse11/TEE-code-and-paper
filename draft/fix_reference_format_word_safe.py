from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


BASE = Path(__file__).resolve().parent
SRC = BASE / "Hybrid_TEE_Rollup_修订版.docx"
OUT = BASE / "Hybrid_TEE_Rollup_参考文献格式修订版.docx"

REFS = [
    "[1] WEN X, FENG Q, LYU H, et al. TEEROLLUP: efficient Rollup design using heterogeneous TEE[EB/OL]. arXiv:2409.14647, 2024 [2026-06-15]. https://arxiv.org/abs/2409.14647.",
    "[2] BUTERIN V. An incomplete guide to Rollups[EB/OL]. 2021 [2026-06-15]. https://vitalik.ca/general/2021/01/05/rollup.html.",
    "[3] ETHEREUM FOUNDATION. Optimistic Rollups[EB/OL]. 2024 [2026-06-15]. https://ethereum.org/en/developers/docs/scaling/optimistic-rollups/.",
    "[4] ARBITRUM FOUNDATION. Arbitrum Nitro technical documentation[EB/OL]. 2024 [2026-06-15]. https://docs.arbitrum.io/.",
    "[5] OP LABS. Optimism documentation: fault proofs and dispute games[EB/OL]. 2024 [2026-06-15]. https://docs.optimism.io/.",
    "[6] PICCO G, FORTUGNO A. Dynamic fraud proof[EB/OL]. arXiv:2502.10321, 2025 [2026-06-15]. https://arxiv.org/abs/2502.10321.",
    "[7] AL-BASSAM M. LazyLedger: a distributed data availability ledger with client-side smart contracts[EB/OL]. arXiv:1905.09274, 2019 [2026-06-15]. https://arxiv.org/abs/1905.09274.",
    "[8] TAS E N, TSE D, YANG L, et al. Light clients for lazy blockchains[EB/OL]. arXiv:2203.15968, 2022 [2026-06-15]. https://arxiv.org/abs/2203.15968.",
    "[9] EIGENLABS. EigenDA documentation[EB/OL]. 2024 [2026-06-15]. https://docs.eigencloud.xyz/products/eigenda/.",
    "[10] CELESTIA. Celestia: modular blockchain network[EB/OL]. 2024 [2026-06-15]. https://docs.celestia.org/.",
    "[11] INTEL CORPORATION. Intel Software Guard Extensions developer guide[EB/OL]. 2024 [2026-06-15]. https://www.intel.com/content/www/us/en/developer/tools/software-guard-extensions/overview.html.",
    "[12] ETHEREUM FOUNDATION. EIP-4844: shard blob transactions[EB/OL]. 2024 [2026-06-15]. https://eips.ethereum.org/EIPS/eip-4844.",
]


def set_east_asian_font(run, name: str) -> None:
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def normalize_citation_spacing(text: str) -> str:
    text = text.replace("\u00ad", "").replace("\ufffe", "")
    return re.sub(r"\]\s+\[", "][", text)


def clear_and_add_text(paragraph, text: str, size_pt: float | None = None, bold: bool | None = None):
    paragraph.clear()
    run = paragraph.add_run(text)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    return run


def set_reference_paragraph_format(paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.left_indent = Pt(18)
    fmt.first_line_indent = Pt(-18)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_reference_text(paragraph, text: str) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    set_east_asian_font(run, "宋体")
    run.font.size = Pt(9)
    set_reference_paragraph_format(paragraph)


def superscript_body_citations(paragraph) -> None:
    text = normalize_citation_spacing(paragraph.text)
    if not re.search(r"\[\d+\]", text):
        return
    paragraph.clear()
    for part in re.split(r"(\[\d+\])", text):
        if not part:
            continue
        run = paragraph.add_run(part)
        if re.fullmatch(r"\[\d+\]", part):
            run.font.superscript = True


def main() -> None:
    doc = Document(str(SRC))
    paragraphs = doc.paragraphs

    ref_idx = None
    for i, p in enumerate(paragraphs):
        if p.text.strip().replace(" ", "") in {"参考文献", "参考文献："}:
            ref_idx = i
            break
    if ref_idx is None:
        raise RuntimeError("Reference heading not found")

    head_run = clear_and_add_text(paragraphs[ref_idx], "参考文献：", size_pt=10.5, bold=True)
    set_east_asian_font(head_run, "黑体")

    ref_paras = []
    for p in paragraphs[ref_idx + 1 :]:
        if re.match(r"^\[\d+\]", p.text.strip()):
            ref_paras.append(p)
    if len(ref_paras) != len(REFS):
        raise RuntimeError(f"Expected {len(REFS)} references, found {len(ref_paras)}")
    for p, text in zip(ref_paras, REFS):
        add_reference_text(p, text)

    for p in paragraphs[:ref_idx]:
        superscript_body_citations(p)

    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
