# -*- coding: utf-8 -*-
"""Phase 4: Final aggressive pass on remaining narrative English."""
from docx import Document
import re

doc = Document(r"E:\项目\project_code\paper_outputs\draft\final_chinese_academic_format.docx")
def st(idx, text):
    doc.paragraphs[idx].text = text

# Final targeted replacements
reps3 = [
    # Table caption cleanup
    ("（Mechanism-Oriented Comparison）", ""),
    # Heading cleanup
    ("（Reproducibility and Evidence Types）", ""),
    # Narrative English that slipped through
    ("Optimistic challenge 是 Rollup", "乐观挑战是汇总"),
    ("Hybrid TEE-Rollup 如何在", "混合 TEE 汇总如何在"),
    ("Hybrid TEE-Rollup 思路下", "混合 TEE 汇总思路下"),
    ("Hybrid TEE-Rollup 面向", "混合 TEE 汇总面向"),
    ("Hybrid TEE-Rollup 原型", "混合 TEE 汇总原型"),
    ("Hybrid TEE-Rollup 中", "混合 TEE 汇总中"),
    ("Hybrid TEE-Rollup 可恢复", "混合 TEE 汇总可恢复"),
    # response/attestation in narrative
    ("下的 response 与", "下的响应与"),
    ("输出 response 与", "输出响应与"),
    ("保存 replay", "保存重放"),
    ("进入 replay", "进入重放"),
    ("的 response 被", "的响应被"),
    ("response 字段及", "响应字段及"),
    ("prompt、response、attestation、", "提示、响应、远程证明、"),
    ("prompt、response、", "提示、响应、"),
    ("response、attestation", "响应、远程证明"),
    ("的 response，", "的响应，"),
    ("prompt 长度", "提示长度"),
    ("prompt、", "提示、"),
    ("的 payload", "的负载"),
    ("完整 payload", "完整负载"),
    ("将 payload", "将负载"),
    ("在 payload", "在负载"),
    ("challenge open、respond、step、recover、replay 和 resolve",
     "挑战打开、响应、步骤、恢复、重放和结算"),
    # fault taxonomy cleanup
    ("（attestation fault）", "（远程证明故障）"),
    ("、timeout 故障、replay 故障", "、超时故障、重放故障"),
    ("trace inconsistency", "轨迹不一致"),
    ("payload、payload", "负载、负载"),
    # P59 cleanup
    ("重写 commit、", "重写紧凑提交、"),
    ("trace 或", "轨迹或"),
    ("slashing 语义", "罚没语义"),
    ("correctness：后续 replay 仍", "正确性：后续重放仍"),
    # P76 bullet cleanup
    ("（状态根）", "（State Root）"),
    ("（Payload Hash）", "（Payload Hash）"),
    ("（Execution Trace）", "（Execution Trace）"),
    # P118-120 narrative 
    ("（synthetic 账本条目）", "（合成账本条目）"),
    ("proof 和", "证明和"),
    ("pointer 和", "指针和"),
    ("replay 准备能力。", "重放准备能力。"),
    ("完整 prompt、response、attestation", "完整提示、响应和远程证明"),
    ("字段都作为 payload 的一部分", "字段都作为负载的一部分"),
    ("状态根、output_hash、proof_hash、da_pointer、da_merkle_root 和 state_root", 
     "状态根、输出哈希、证明哈希、数据可用性指针和数据可用性 Merkle 根"),
    # P210 
    ("recovery_success 与", "恢复成功率与"),
    ("DA-aware 成本模型", "数据可用性感知成本模型"),
    ("slashing 闭环", "罚没闭环"),
    # P214
    ("submit、challenge open、recover、replay、resolve 和 finalize", 
     "提交、挑战打开、恢复、重放、结算和最终确认"),
    # P216  
    ("challenge + slashing 路径", "挑战加罚没路径"),
    ("fault 仍需要", "故障仍需要"),
    ("watchdog 需要", "观察节点需要"),
    # P217
    ("evidence 仍由", "证据仍由"),
    ("trace evidence 和", "轨迹证据和"),
    ("replay arbitration", "重放仲裁"),
    ("evidence 在", "证据在"),
    ("replay 和 replay evidence 之间的", "重放和重放证据之间的"),
    # P218
    ("evidence 替代 simulated attestation 并建立更完整的 challenge 失败到 slashing 的闭环。",
     "证据替代模拟远程证明并建立更完整的挑战失败到罚没的闭环。"),
    # P220
    ("Evidence-Carrying 紧凑提交", "携带证据的紧凑提交"),
    ("Binding，使", "绑定，使"),
    # P232  
    ("commit、数据可用性证明、challenge、recovery、replay 与 resolution",
     "紧凑提交、数据可用性证明、挑战、恢复、重放与结算"),
    # P233
    ("evidence 的链上", "证据的链上"),
    ("replay 对链下", "重放对链下"),
    # P28
    ("payload 线性", "负载线性"),
    ("timeout 不会", "超时不会"),
    # P25 additional
    ("challenge 本身", "挑战本身"),
    ("challenge 的", "挑战的"),
    # P19
    ("timeout 后的", "超时后的"),
    ("commit 到", "紧凑提交到"),
    ("workload 的", "工作负载的"),
    # P221  
    ("challenge 会话", "挑战会话"),
    ("TEE-Rollup 方向", "TEE 汇总方向"),
    ("timeout 恢复", "超时恢复"),
    # P222
    ("Optimistic TEE-Rollups", "乐观 TEE 汇总"),
    ("optimistic challenge 可", "乐观挑战可"),
    ("challenge 触发", "挑战触发"),
    ("payload 改变", "负载改变"),
    # P223
    ("optimistic 欺诈证明", "乐观欺诈证明"),
    ("bisection replay 用于", "二分重放用于"),
    # P225
    ("Hybrid TEE-Rollup 场景", "混合 TEE 汇总场景"),
    # P226
    ("transaction 引入", "交易引入"),
    ("profile 只抽象", "配置只抽象"),
    ("full 调用数据", "完整调用数据"),
    # P229
    ("optimistic 欺诈证明", "乐观欺诈证明"),
    # P231
    ("challenge 形成可运行的", "挑战形成可运行的"),
]

for i, p in enumerate(doc.paragraphs):
    text = p.text
    for eng, chn in reps3:
        if eng in text:
            text = text.replace(eng, chn)
    if text != p.text:
        doc.paragraphs[i].text = text

# Also fix table cells
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                text = para.text
                for eng, chn in reps3:
                    if eng in text:
                        text = text.replace(eng, chn)
                if text != para.text:
                    para.text = text

doc.save(r"E:\项目\project_code\paper_outputs\draft\final_chinese_academic_format.docx")
print("Phase 4 complete. Final clean pass done.")
