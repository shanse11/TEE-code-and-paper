from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(r"E:\项目\meetings\weekly_reports\docs\26春0520第十二周组会汇报词-杨帆.docx")


def set_east_asia(run, font_name="宋体"):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    set_east_asia(r)
    r.font.size = Pt(10.5)
    r.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_q_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0].cells
    set_cell_text(header[0], "导师可能追问的问题", True)
    set_cell_text(header[1], "回答建议", True)
    shade_cell(header[0], "D9EAF7")
    shade_cell(header[1], "D9EAF7")
    for question, answer in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], question)
        set_cell_text(cells[1], answer)
    doc.add_paragraph()


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        r = p.add_run(item)
        set_east_asia(r)


def add_paragraphs(doc, text):
    for para in text.split("\n\n"):
        p = doc.add_paragraph()
        r = p.add_run(para)
        set_east_asia(r)


def setup_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)

    style_specs = [
        ("Title", "黑体", 18, True, "000000"),
        ("Heading 1", "黑体", 15, True, "1F4E79"),
        ("Heading 2", "黑体", 13, True, "000000"),
        ("Heading 3", "黑体", 12, True, "000000"),
    ]
    for name, font, size, bold, color in style_specs:
        style = styles[name]
        style.font.name = font
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(6)
    return doc


SLIDES = [
    {
        "title": "第1页：封面",
        "goal": "说明本次汇报主题是第十二周阶段性进展，重点围绕实验整理和论文质量提升。",
        "speech": "这一周的工作是在第十一周已经完成较多实验和链上部署的基础上，进一步把这些材料转化成论文中可以成立的论证。前面几周我更多是在做功能实现、实验运行和数据收集；这周的重点转向“如何把原型工作讲清楚”。\n\n具体来说，我主要围绕 Hybrid TEE-Rollup 这个方向，重新梳理论文中的实验章节、系统架构表达和理论创新表述。尤其是导师前面提到的一个问题：不能只把实验结果当作创新点，所以我这周也在区分实验贡献和理论建模贡献。",
        "qa": [
            ("这周相比上周的新进展是什么？", "不是完全新增一套实验，而是把已有实验材料转化为论文论证结构，同时补充图示、分类和局限性表达。"),
            ("你是继续做实验，还是主要改论文？", "两者都有，但重心是把已有实验整理成可支撑论文结论的证据链。"),
            ("目前工作定位到底是系统实现还是论文整理？", "定位为 Hybrid TEE-Rollup 研究原型的阶段性完善，系统实现和论文表达同步推进。"),
        ],
    },
    {
        "title": "第2页：主要任务",
        "goal": "说明本周工作的主线：把已有实验和论文材料整理成更符合科研论文逻辑的表达。",
        "speech": "这一页主要说明我这周的总体任务。我的中长期目标没有变化，还是研究在保证去中心化安全边界的前提下，如何提升高频 DApp 场景下 Rollup 系统的性能。\n\n这里的核心问题是：如果所有状态变化和数据都直接放到链上，链上 calldata 和执行成本都会随着交易规模快速增长，这对于元宇宙、去中心化社交这类高频应用并不合适。因此需要考虑链下执行、链下数据可用性和链上轻量验证之间的折中。\n\n本周我主要做了两件事。第一，把 Evaluation 章节按照 RQ1 到 RQ5 五个研究问题重新组织。这样做的原因是，实验章节不能只是把成本实验、挑战实验、失败场景实验逐个堆出来，而应该说明每个实验回答什么问题、采用什么指标、能支撑什么结论、不能外推到哪里。第二，我补充了系统架构图和协议流程图，用来更清楚地表达 Hybrid TEE-Rollup 原型中 simulated TEE、mock DA、compact commit 和 challenge 机制之间的关系。",
        "qa": [
            ("为什么不能全部上链？", "链上存储和 calldata 成本随 payload 增长，难以适配高频应用。"),
            ("为什么需要链下执行？", "链下执行可承载高频计算和状态推进，链上只保留 compact commit、验证入口和争议仲裁。"),
            ("Evaluation 为什么要按 RQ 重构？", "RQ 能把实验设置、指标、结果和结论一一对应，避免结果堆叠和过度外推。"),
        ],
    },
    {
        "title": "第3页：Evaluation 从结果堆叠改为 RQ 驱动",
        "goal": "说明实验章节如何从“列结果”变成“回答研究问题”。",
        "speech": "这一页是本周论文修改中比较核心的部分。前期我已经有三类主要实验：成本实验 3200 条记录、挑战实验 600 条记录、失败场景实验 800 条记录。但如果只是把这些结果按实验类型罗列出来，读者很难知道它们分别支撑论文中的哪个判断。\n\n所以我把 Evaluation 重新组织为五个研究问题。RQ1 关注 compact commit 是否确实降低链上提交规模；RQ2 关注 batch size 增大后，在 DA-aware cost model 下单笔摊销成本是否下降；RQ3 关注 recoverable challenge 是否改善 timeout 情况下的仲裁活性；RQ4 关注不同 failure scenario 中，哪些只是可检测，哪些可以进入 challenge 和 slashing 闭环；RQ5 关注 Solidity 原型、本地 EVM gas 和 Sepolia 部署是否能够支撑“链上对应物”这一阶段性主张。\n\n这里我也特别区分了四类证据边界。Python prototype 主要说明机制能跑通；DA cost model 说明配置化参数下的趋势；Hardhat gasUsed 只是本地 EVM 执行基线；Sepolia 只能说明合约可以部署到公开测试网，不能说明主网性能或生产可用性。",
        "qa": [
            ("RQ1 到 RQ5 之间的逻辑关系是什么？", "RQ1、RQ2 讨论链上提交规模和成本趋势；RQ3、RQ4 讨论挑战机制和故障处理；RQ5 讨论链上对应实现。"),
            ("这些实验是否足够支撑论文结论？", "足够支撑研究原型机制可行和趋势性分析，但不能支撑生产级结论。"),
            ("为什么 Sepolia 不能作为性能证据？", "Sepolia 是测试网环境，网络条件、gas 市场和真实负载都不等同主网。"),
        ],
    },
    {
        "title": "第4页：成本模型与 compact commit",
        "goal": "说明 compact commit 的意义，以及轻量 DA 成本模型如何支撑趋势性结论。",
        "speech": "这一页主要对应 RQ1 和 RQ2。传统做法如果把 full payload 直接作为链上 calldata 提交，那么 payload 越大，链上提交规模基本会线性增长。对高频 DApp 来说，这会导致链上成本迅速放大。\n\n我当前原型中的 compact commit 思路，是链上不直接承载完整 payload，而是提交一个包含状态根、DA commitment、TEE evidence 等信息的紧凑承诺。在当前实现里，compact commit 大约是 406 bytes，并且基本不随 payload size 增长。这个结果的意义不是说系统已经实现了真实主网降本，而是说明“链上承诺大小”和“链下 payload 大小”可以在机制上解耦。\n\n成本模型方面，我把论文中的表达整理为：C_total = C_commit + C_DA + p_challenge × C_challenge，C_amortized = C_total / batch_size。这里 C_commit 表示链上提交 compact commit 的成本，C_DA 表示数据可用性层成本，p_challenge × C_challenge 表示挑战概率下的期望挑战成本。当前论文比较的是 full_onchain_calldata、compact_external_da、compact_eip4844_like 和 compact_modular_da_sampling 四类 profile。需要强调的是，96.33% 这类降本幅度只是在给定 cost model 参数下的趋势，不是主网真实价格。",
        "qa": [
            ("compact commit 为什么有理论意义？", "它将链上验证入口和链下数据规模分离，使链上只保存可验证承诺。"),
            ("406 bytes 是否固定？", "不是协议常数，而是当前原型编码下的测量结果。"),
            ("96.33% 是否可信？", "它可信的范围是配置化模型下的趋势比较，不能解释为真实主网降本。"),
        ],
    },
    {
        "title": "第5页：Recoverable Challenge 与失败场景分类",
        "goal": "说明 recoverable challenge 是活性补强，不是新的正确性证明；同时区分 detection 和 slashing。",
        "speech": "这一页对应 RQ3 和 RQ4。前期实验里 recover 的结果看起来是 timeout 场景下成功率从 0% 提升到 100%，但如果只这样写，会显得像是在说 recover 证明了系统正确性。实际上 recoverable challenge 更准确的定位是 liveness enhancement，也就是活性补强。\n\n它并不改变 expected trace、claimed trace，也不改变 replay mismatch 的事实判断。它解决的是 timeout 之后 challenge session 无法继续推进的问题，让争议过程可以恢复到可继续仲裁的状态。因此 recover 改善的是挑战流程的推进能力，而不是提供新的安全证明。\n\n另外，我这周把 failure scenario 分成四类。第一类是 normal path；第二类是 attestation_invalid、da_unavailable、da_proof_invalid，这些目前可以检测或拒绝，但还没有统一 slashing 闭环；第三类是 response_tampered 和 trace_length_mismatch，可以进入 challenge + slashing；第四类是 challenge_timeout_no_recover 和 challenge_timeout_recover，用于比较 timeout 场景下的活性差异。这里最重要的一点是：failure detection 不等于 slashing。检测到异常只说明系统识别了问题，但是否能罚没、由谁罚没、证据如何上链验证，是更强的仲裁闭环，目前还没有对所有 fault 类型统一完成。",
        "qa": [
            ("recover 会不会改变挑战结果？", "不会，它不改变争议事实，只恢复挑战流程推进能力。"),
            ("为什么说 recover 是 liveness 而不是 correctness？", "它解决 timeout 后流程卡住的问题，不提供新的正确性证明。"),
            ("detection 和 slashing 的区别是什么？", "detection 是发现异常，slashing 是基于链上可验证证据执行惩罚。"),
        ],
    },
    {
        "title": "第6页：系统架构图与协议流程图",
        "goal": "说明本周补充图示的作用：让论文中的系统边界和协议状态推进更清晰。",
        "speech": "这一页是论文表达层面的补强。因为 Hybrid TEE-Rollup 涉及链下执行、TEE evidence、DA commitment、链上合约和 challenge 机制，如果只用文字描述，读者不容易把各个模块之间的关系看清楚。\n\n所以我新增了系统架构图，把高频 DApp、simulated TEE、mock/verifiable DA、compact commit 和链上验证仲裁层放在同一张图里。这样可以更清楚地表达：链下负责执行和生成证据，DA 层负责 payload 的可用性承诺，链上负责接收 compact commit 和处理争议。\n\n另外我新增了 Recoverable Challenge 协议流程图，重点展示 open、respond、step、recover、replay、resolve 这些状态推进。这个图的意义是帮助论文说明 recover 并不是额外绕过挑战，而是在 timeout 后让 challenge session 回到可继续推进的状态。\n\n同时我也在图中保持边界表达：这里的 TEE 是 simulated TEE，不是真实 SGX、TDX 或 CSV 的安全证明；DA 也是 mock/verifiable DA，不是真实 DA 网络。",
        "qa": [
            ("架构图中哪些模块是真实实现，哪些是模拟？", "Python prototype、部分 Solidity 合约路径、本地 EVM gas 和 Sepolia 部署是实现证据；TEE 和 DA 仍是模拟或 mock 层。"),
            ("simulated TEE 和真实 TEE 差别在哪里？", "simulated TEE 只模拟 evidence 和 attestation 流程，不提供真实硬件隔离安全。"),
            ("mock DA 是否会削弱论文贡献？", "会限制结论外推，但不影响原型层面对 DA proof 绑定和成本模型趋势的验证。"),
        ],
    },
    {
        "title": "第7页：本周学习体会",
        "goal": "总结本周在科研方法上的收获：实验、结论和理论创新必须匹配。",
        "speech": "这一页是我这周比较重要的反思。前面做实验时，我更多关注“能不能跑通”和“结果是否明显”。但在整理成论文时发现，实验结果本身并不自动等于理论创新。\n\n第一点体会是，Evaluation 要由研究问题驱动。也就是说，每个实验都要回答一个明确问题，而不是单纯展示结果。比如成本实验回答链上提交规模和摊销成本，挑战实验回答 recover 对活性的影响，failure 实验回答不同 fault 的处理边界。\n\n第二点是，结论必须匹配证据类型。Python prototype 不能证明真实 TEE 安全；DA cost model 不能证明真实主网费用；Hardhat gas 不能代表公网成本；Sepolia 部署也不能代表生产可用性。如果不区分这些边界，论文表述就会过强。\n\n第三点是，原型研究要保持克制。当前工作的价值在于提出一种 Hybrid TEE-Rollup 的机制组织方式，包括 compact commit、recoverable challenge、fault taxonomy 和轻量 DA 成本模型。但它还不是完整生产级 Rollup 系统，后续还需要真实 TEE、真实 DA、形式化分析和更完整的 fault arbitration。",
        "qa": [
            ("你的理论创新到底是什么？", "理论创新在问题建模和协议抽象，包括 compact commit、recoverable challenge、fault taxonomy 和 DA-aware cost model。"),
            ("你这周最大的认识变化是什么？", "认识到实验结果必须转化为有边界的证据，而不是直接写成创新点。"),
            ("当前论文最薄弱的地方是什么？", "部分 fault 还没有统一 slashing 闭环，真实 TEE 和真实 DA 仍未接入。"),
        ],
    },
    {
        "title": "第8页：下周科研计划",
        "goal": "说明后续工作不是简单“继续补实验”，而是补齐 arbitration、Sepolia interaction 和论文贡献表达。",
        "speech": "下周我计划从三个方向继续推进。\n\n第一是深化 fault arbitration。目前有些 fault，比如 attestation_invalid、da_unavailable、da_proof_invalid，已经可以检测或拒绝，但还没有统一进入 slashing 闭环。下周我会继续梳理这些 fault 类型需要什么证据、证据是否能链上验证，以及能否设计更明确的仲裁路径。\n\n第二是补充 Sepolia interaction。目前 Sepolia 主要是部署证据，说明合约可以部署到公开测试网。后续我希望补充少量交互样本和源码验证重试，但仍然会把它定位为 deployability evidence，而不是性能证据。\n\n第三是完善论文贡献段落。尤其是把理论创新正式写入 Introduction、Contribution 和 Discussion 中，包括 compact commit、recoverable challenge、fault taxonomy 和 DA-aware cost model，同时统一全文中过强的表述，避免把研究原型写成生产系统。",
        "qa": [
            ("下周最优先做什么？", "优先补 fault arbitration，因为这是从“可检测”走向“可仲裁”的关键。"),
            ("Sepolia interaction 有必要吗？", "有必要，但只作为测试网可部署和可交互证据，不扩大解释。"),
            ("论文贡献段落怎么写才不像实验总结？", "从机制抽象、协议语义和安全边界建模来写，而不是从实验数值出发。"),
        ],
    },
]


def main():
    doc = setup_doc()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("第十二周组会汇报词")
    set_east_asia(r, "黑体")
    r.font.size = Pt(20)
    r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Hybrid TEE-Rollup 研究原型实验与论文进展")
    set_east_asia(r)
    r.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("汇报人：杨帆    时间：2026年5月20日")
    set_east_asia(r)
    r.font.size = Pt(11)

    doc.add_heading("开场白（约30秒）", level=1)
    add_paragraphs(
        doc,
        "各位老师、同学好，我这周主要做了两部分工作：一是把前期实验结果重新整理成论文中更清晰的 Evaluation 论证结构，二是继续补强 Hybrid TEE-Rollup 原型的系统表达，包括系统架构图、协议流程图、failure 场景分类以及论文中对实验边界和局限性的说明。\n\n"
        "这周我特别关注的不是再单纯增加实验数量，而是思考这些实验分别能支撑什么结论、不能支撑什么结论。因为目前的工作仍然是一个研究原型，不能把 Python 原型、成本模型、本地 EVM gas 或 Sepolia 部署证据直接写成生产级系统结论。",
    )

    for slide in SLIDES:
        doc.add_heading(slide["title"], level=1)
        doc.add_heading("1. 本页核心目标", level=2)
        doc.add_paragraph(slide["goal"])
        doc.add_heading("2. 自然口语化汇报词", level=2)
        add_paragraphs(doc, slide["speech"])
        doc.add_heading("3. 导师可能追问的问题与回答建议", level=2)
        add_q_table(doc, slide["qa"])

    doc.add_heading("结束总结（约1分钟）", level=1)
    add_paragraphs(
        doc,
        "总体来说，这周我的工作重点是把 Hybrid TEE-Rollup 原型从“能运行、有结果”进一步整理为“论文中可以被清楚论证的研究工作”。我重新组织了 Evaluation，使其围绕五个研究问题展开；进一步明确了 compact commit、recoverable challenge、failure taxonomy 和轻量 DA 成本模型的论文表达；同时也补充了系统架构图和协议流程图，帮助说明系统边界和协议状态推进。\n\n"
        "目前我对这项工作的定位更加明确：它不是完整生产级 Rollup，也不能声称真实 TEE 安全或主网级降本。它的价值主要在于提出并验证一种 Hybrid TEE-Rollup 研究原型中的机制组织方式，并用 Python prototype、DA cost model、本地 EVM gas 和 Sepolia deployability 从不同层面提供阶段性证据。后续我会重点补强 fault arbitration、Sepolia interaction 和论文贡献段落，使论文结论和证据边界更加一致。",
    )

    doc.add_heading("导师最可能重点质疑的地方", level=1)
    add_q_table(
        doc,
        [
            ("你的创新点是不是只是实验结果？", "不是。实验只是验证材料，理论创新在于 compact commit 证据结构、recoverable challenge 活性语义、fault taxonomy、DA-aware cost model 和 Hybrid TEE-Rollup 的问题建模。"),
            ("simulated TEE 怎么能说明 TEE 安全？", "不能说明真实 TEE 安全。它只用于模拟 evidence-carrying commit 和 attestation 验证流程，真实硬件安全需要 SGX/TDX/CSV 接入和额外安全分析。"),
            ("DA 是 mock 的，成本结论可信吗？", "成本结论只能解释为配置化模型下的趋势，不是主网价格。它支撑的是比较框架和趋势分析，而不是实际费用承诺。"),
            ("Sepolia 部署有什么意义？", "Sepolia 说明合约具备公开测试网 deployability，证明不是纯 Python 模型；但不说明生产可用性、主网性能或真实安全。"),
            ("detection 为什么不等于 slashing？", "detection 是识别异常，slashing 是基于可验证证据执行惩罚。后者需要更强的证据上链、责任归属和仲裁规则。"),
        ],
    )

    doc.add_heading("如何回答“你的创新点到底是什么”", level=1)
    doc.add_paragraph("可以这样回答：我的创新点不是某个实验数值，而是对 Hybrid TEE-Rollup 原型中的几个关键机制进行了建模和组织。")
    add_numbered(
        doc,
        [
            "提出 evidence-carrying compact commit 的结构，把链下执行结果、TEE evidence、DA commitment 和链上验证入口绑定起来，使链上提交不再随 payload 线性增长。",
            "把 recoverable challenge 明确建模为 timeout 场景下的 liveness enhancement。它不改变争议事实和正确性判断，而是恢复 challenge session 的推进能力。",
            "对 failure scenario 做分类，区分 normal path、detectable but not unified slashing、challenge + slashing 和 timeout liveness comparison，避免把“能检测”直接等同于“能惩罚”。",
            "整理轻量 DA-aware cost model，用 C_total 和 C_amortized 描述 compact commit、DA 成本和挑战成本之间的关系，为不同 DA profile 的比较提供统一框架。",
        ],
    )

    doc.add_heading("如何回答“你这个和传统 Rollup 有什么区别”", level=1)
    add_paragraphs(
        doc,
        "可以这样回答：传统 optimistic rollup 主要依赖链下执行、链上提交状态根和 fraud proof 争议机制。我的工作仍然借鉴 optimistic rollup 的思路，但区别在于引入了 TEE evidence 和 DA commitment 的组合。\n\n"
        "也就是说，链上提交的不是完整数据，而是 compact commit；链下执行不仅给出状态承诺，还附带 simulated TEE evidence；数据可用性不直接等同于链上 calldata，而是通过不同 DA profile 建模；争议机制中还加入了 recoverable challenge，用于处理 timeout 下的活性问题。\n\n"
        "所以这项工作不是替代传统 Rollup，而是在 Hybrid TEE-Rollup 方向上探索 TEE evidence、DA commitment 和 challenge arbitration 如何组合。",
    )

    doc.add_heading("如何回答“为什么你的实验足够支撑结论”", level=1)
    add_paragraphs(
        doc,
        "可以这样回答：我的实验支撑的是阶段性研究原型结论，而不是生产系统结论。不同实验对应不同证据类型。\n\n"
        "Python prototype 支撑机制可运行，说明 compact commit、DA proof、challenge 和 recover 流程可以被串联起来。DA cost model 支撑配置化趋势分析，说明在给定参数下 compact+DA 相比 full-onchain calldata 有成本下降趋势。本地 Hardhat gasUsed 支撑 Solidity 链上路径存在，并提供本地执行基线。Sepolia 部署支撑 deployability，说明合约可以部署到公开测试网。\n\n"
        "所以实验足够支撑“研究原型机制可行、趋势存在、链上对应物存在”这些克制结论；但不能外推为真实 TEE 安全、真实 DA 网络性能、主网级降本或生产可用性。",
    )

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("第十二周组会汇报词 - Hybrid TEE-Rollup 研究原型")
        set_east_asia(run)
        run.font.size = Pt(9)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
