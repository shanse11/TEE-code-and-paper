# -*- coding: utf-8 -*-
"""对 Hybrid_TEE_Rollup.docx 进行研究生阶段学术增强（不改实验数据与核心结论）。"""

from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
DOCX_PATH = ROOT / "project_code" / "paper_outputs" / "draft" / "Hybrid_TEE_Rollup.docx"
BACKUP_DIR = ROOT / "project_code" / "paper_outputs" / "draft" / "backups"

NEW_REFERENCES = [
    "[8] Intel Corporation. Intel Software Guard Extensions (Intel SGX) Developer Guide[EB/OL]. 2024.",
    "[9] Intel Corporation. Intel Trust Domain Extensions (Intel TDX) Architecture Specification[EB/OL]. 2024.",
    "[10] Offchain Labs. Arbitrum: A Next-Generation Layer 2 for Ethereum[EB/OL]. 2021.",
    "[11] Cartesi. Cartesi Rollups Documentation[EB/OL]. 2024.",
    "[12] EigenLayer. EigenDA Documentation[EB/OL]. 2024.",
    "[13] Celestia. Celestia: Modular Blockchain Network[EB/OL]. 2024.",
    "[14] RISC Zero. RISC Zero zkVM Documentation[EB/OL]. 2024.",
    "[15] Succinct Labs. SP1: A Performant, 100% Open-Source zkVM[EB/OL]. 2024.",
    "[16] Arbitrum Foundation. Arbitrum Nitro Technical Documentation[EB/OL]. 2024.",
    "[17] OP Labs. Optimism Documentation: Fault Proofs and Dispute Games[EB/OL]. 2024.",
    "[18] Canetti R, et al. Universally Composable Security: A New Paradigm for Cryptographic Protocols[EB/OL]. 2000.",
    "[19] Goldwasser S, Micali S, Rackoff R. The Knowledge Complexity of Interactive Proof Systems[J]. SIAM Journal on Computing, 1989, 18(1): 186-208.",
    "[20] NIST. Recommendation for Key Management: Part 1 - General (Rev. 5)[EB/OL]. 2020.",
]


def insert_after(anchor: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    anchor._element.addnext(new_p)
    return Paragraph(new_p, anchor._parent)


def insert_before(anchor: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    anchor._element.addprevious(new_p)
    return Paragraph(new_p, anchor._parent)


def find_para(doc: Document, contains: str) -> Paragraph | None:
    for p in doc.paragraphs:
        if contains in p.text:
            return p
    return None


def find_para_exact(doc: Document, text: str) -> Paragraph | None:
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    return None


def set_run_font(run, size=10.5, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def add_body(doc: Document, anchor: Paragraph, text: str, *, first_line=True) -> Paragraph:
    p = insert_after(anchor)
    anchor = p
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_heading(doc: Document, anchor: Paragraph, text: str, level: int = 2) -> Paragraph:
    p = insert_after(anchor)
    p.style = doc.styles[f"Heading {min(level, 3)}"]
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, 11.5 if level == 2 else 10.5, bold=True)
    r.font.color.rgb = RGBColor(31, 78, 121)
    return p


def set_cell_shading(cell, fill: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shd)


def add_table_after(
    doc: Document,
    anchor: Paragraph,
    caption: str,
    headers: list[str],
    rows: list[list[str]],
    col_widths: list[float] | None = None,
) -> Paragraph:
    cap = insert_after(anchor)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    cr = cap.add_run(caption)
    set_run_font(cr, 10, bold=True)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    cap._element.addnext(table._tbl)

    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, 9, bold=True)
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(cell, "1F4E79")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for row_data in rows:
        row = table.add_row()
        for i, val in enumerate(row_data):
            cell = row.cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            set_run_font(r, 9)

    if col_widths:
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = Cm(col_widths[min(i, len(col_widths) - 1)])

    last = insert_after(cap)
    return last


def renumber_tables(doc: Document):
    replacements = [
        ("表 5  Related Work", "表 6  Related Work"),
        ("表 4  Sepolia", "表 5  Sepolia"),
        ("表 3  本地 EVM", "表 4  本地 EVM"),
        ("表 2  失败场景", "表 3  失败场景"),
    ]
    for p in doc.paragraphs:
        for old, new in replacements:
            if old in p.text:
                p.text = p.text.replace(old, new)


def add_contributions_and_capability_table(doc: Document):  # uses doc for table creation
    intro_anchor = find_para(doc, "本文所有结论均限定为研究原型结果")
    if intro_anchor is None:
        intro_anchor = find_para(doc, "围绕上述问题")
    if intro_anchor is None:
        return

    h = add_heading(doc, intro_anchor, "1.1 Contributions（本文贡献）", level=2)
    p = h
    items = [
        "（1）Recoverable Challenge Protocol。提出可恢复挑战机制（Recoverable Challenge Protocol），将传统 challenge 从错误检测流程扩展为具备异常恢复能力的活性保障机制，使 timeout 或中断后的 challenge session 能够继续推进至 replay 与 resolve 阶段。",
        "（2）Evidence-Carrying Compact Commit。提出 evidence-carrying compact commit，将 state root、output hash、proof hash、DA root 与 payload proof 进行绑定，为轻量提交建立最小可验证证据入口。",
        "（3）DA-Aware Cost Model。构建面向高频 DApp 的 DA-aware cost model，用于分析 payload size、batch size 与 DA route 对成本摊销的联合影响。",
    ]
    for item in items:
        p = add_body(doc, p, item)

    # 能力对比表放在「表 1」之后，避免表号顺序错乱
    table1_cap = find_para(doc, "表 1  本文机制推进点与研究作用")
    if table1_cap is None:
        table1_cap = p
    else:
        el = table1_cap._element
        while el.getnext() is not None and el.getnext().tag.endswith("tbl"):
            el = el.getnext()
        table1_cap = Paragraph(el, table1_cap._parent)

    p = add_table_after(
        doc,
        table1_cap,
        "表 2  Capability Comparison（系统能力对比）",
        ["System", "Interactive Challenge", "Recovery", "DA Binding", "Cost Awareness"],
        [
            ["Optimism", "√", "×", "×", "×"],
            ["Arbitrum", "√", "×", "×", "×"],
            ["OTR", "√", "×", "×", "×"],
            ["TEERollup", "√", "△", "△", "×"],
            ["This Work", "√", "√", "√", "√"],
        ],
        col_widths=[2.8, 3.2, 2.0, 2.4, 2.8],
    )
    add_body(
        doc,
        p,
        "需要强调的是，本文创新点并非重新设计 Rollup 系统，而是在已有 Hybrid TEE-Rollup 架构下补强 Recovery、DA Binding 与 Cost Awareness 三项关键能力。上述对比用于说明机制层增强边界，而非声称全面优于既有系统。",
    )


def add_formal_definitions(doc: Document):
    anchor = find_para_exact(doc, "6 Recoverable Challenge Protocol")
    if anchor is None:
        return

    h = insert_before(anchor)
    h.style = doc.styles["Heading 2"]
    h.paragraph_format.space_before = Pt(8)
    r = h.add_run("5.4 Formal Definitions（形式化定义）")
    set_run_font(r, 11.5, bold=True)
    r.font.color.rgb = RGBColor(31, 78, 121)

    p = h
    p = add_body(
        doc,
        p,
        "为在 Recoverable Challenge Protocol 之前建立清晰的理论接口，本节给出三个核心定义，并说明其与恢复安全性之间的关系。",
    )

    p = add_heading(doc, p, "Definition 1 (Challenge Session)", level=3)
    p = add_body(doc, p, "挑战会话定义为：cs = (id, s, F, T)。", first_line=False)
    p = add_body(doc, p, "其中，id 表示会话标识；s 表示当前状态；F 表示 Challenge Facts（挑战事实集合）；T 表示超时参数。", first_line=False)

    p = add_heading(doc, p, "Definition 2 (Challenge State Machine)", level=3)
    p = add_body(doc, p, "挑战状态机定义为：M = (S, E, δ)。", first_line=False)
    p = add_body(doc, p, "其中，S 为状态集合；E 为事件集合；δ: S × E → S 为状态转移函数。", first_line=False)

    p = add_heading(doc, p, "Definition 3 (Recovery Function)", level=3)
    p = add_body(doc, p, "恢复函数定义为：Recover(cs) = cs'。", first_line=False)
    p = add_body(doc, p, "恢复仅改变活性相关状态，不改变 Challenge Facts，即满足：F(cs) = F(Recover(cs))。", first_line=False)

    p = add_heading(doc, p, "Property and Discussion", level=3)
    p = add_body(
        doc,
        p,
        "基于上述定义，可将恢复操作视为活性层变换而非正确性裁决：Recover 允许会话从 TIMED_OUT 等停滞状态回到可推进状态，但不重写 commit、DA root、payload hash、trace 或 mismatch evidence。由此，Definition → Property → Discussion 形成完整理论链条：定义给出状态与事实边界，性质约束恢复不破坏证据上下文，讨论部分则将其映射到第 6 章协议状态机与第 8 章实验观测。",
    )


def add_recovery_overhead_analysis(doc: Document):
    lim = find_para(doc, "Limitation. 当前 challenge trace")
    if lim is None:
        lim = find_para(doc, "Takeaway. 该结果说明 recover 的核心价值")
    if lim is None:
        return

    h = add_heading(doc, lim, "8.3.1 Recovery Overhead Analysis（恢复开销分析）", level=3)
    p = add_body(
        doc,
        h,
        "利用表 4（原表 3）中已有本地 EVM 数据，challenge step 的平均 gasUsed 为 93807，challenge recover 为 156891。由此可得恢复相对单轮 challenge step 的开销比约为：Recovery Overhead = 156891 / 93807 ≈ 1.67。",
    )
    p = add_body(
        doc,
        p,
        "该结果表明 Recover 的单次链上开销高于一次 challenge step。然而，Recover 属于低频异常路径操作，其作用是在 timeout 已被检测但流程停滞时恢复 challenge liveness，使会话能够继续进入 replay 与 resolve。对于高频 DApp 的正常吞吐路径，该额外成本通常可接受。",
    )
    add_body(
        doc,
        p,
        "需要再次强调：上述数值仅来自 Hardhat 本地 EVM 原型测量，不代表主网 gas 价格、拥堵条件或生产级合约优化后的成本；本文不将其外推为链上部署的通用性能结论。",
    )


def restructure_threats_to_validity(doc: Document):
    anchor = find_para_exact(doc, "9 Discussion and Limitations")
    if anchor is None:
        return
    anchor.text = "9 Threats to Validity（效度威胁与局限）"

    # 在标题后插入结构化引导段
    intro = insert_after(anchor)
    intro.paragraph_format.first_line_indent = Cm(0.74)
    r = intro.add_run(
        "为避免将原型结果误解为生产系统结论，本文将原有讨论整理为三类效度威胁：Internal Validity、External Validity 与 Construct Validity。"
    )
    set_run_font(r)

    h1 = add_heading(doc, intro, "Internal Validity（内部效度）", level=3)
    p = add_body(
        doc,
        h1,
        "内部效度主要受 simulated TEE、mock/verifiable DA registry、Python 原型逻辑与 Hardhat 本地 EVM 测量方式影响。"
        "原型中的 attestation 仅建模输入/输出绑定，DA 层未接入真实采样与经济安全机制，因此机制验证与实现细节耦合。"
        "本地 gasUsed 也可能因合约简化、测试数据规模固定而与真实部署存在偏差。",
    )

    h2 = add_heading(doc, p, "External Validity（外部效度）", level=3)
    p = add_body(
        doc,
        h2,
        "外部效度限制来自实验环境到主网/生产网络的泛化边界。"
        "本文未覆盖主网 MEV、拥堵费率波动、真实 TEE 供应链攻击与真实 DA 网络运营风险。"
        "Sepolia 部署仅提供 deployability evidence，不能支持主网性能或安全强度外推。",
    )

    h3 = add_heading(doc, p, "Construct Validity（构念效度）", level=3)
    add_body(
        doc,
        h3,
        "构念效度关注指标是否准确刻画论文主张。"
        "challenge success、recovery_success 与 slashed_rate 衡量的是原型状态机下的仲裁推进，而非完整 Rollup 经济安全；"
        "DA-aware cost model 为配置化趋势模型，不等价于真实 calldata/blob 定价；"
        "fault taxonomy 亦尚未形成统一 slashing 闭环。"
        "因此，本文结论应被理解为机制层研究原型结论，而非生产系统审计结果。",
    )


def enhance_related_work_opening(doc: Document):
    anchor = find_para_exact(doc, "10 Related Work")
    if anchor is None:
        return
    p = insert_after(anchor)
    p.paragraph_format.first_line_indent = Cm(0.74)
    r = p.add_run(
        "本文与现有工作的关系属于“Mechanism Enhancement（机制增强）”，而非“System Replacement（系统替代）”。"
        "本文并不试图替代 TEERollup、OTR 或 opML 等系统路线，而是在其异常路径上补充 Recoverable Challenge 与 Evidence-Carrying DA Binding，"
        "使 Hybrid TEE-Rollup 在 timeout、DA 异常与轻量提交场景下具备更可解释的仲裁推进能力。"
    )
    set_run_font(r)


def add_theoretical_implication(doc: Document):
    anchor = find_para(doc, "后续工作应继续推进真实 TEE attestation")
    if anchor is None:
        anchor = find_para(doc, "当前工作仍是阶段性研究原型")
    if anchor is None:
        return
    h = add_heading(doc, anchor, "11.1 Theoretical Implication（理论启示）", level=3)
    add_body(
        doc,
        h,
        "本文最大的贡献并非提出新的 Rollup 系统，而是将 Challenge Recovery 从工程实现细节提升为具有协议语义的独立机制。"
        "通过 Recoverable Challenge、Evidence-Carrying Commit 与 DA-Aware Cost Model 的结合，"
        "本文为 Hybrid TEE-Rollup 在高频 DApp 场景下的异常处理与低成本验证提供了一种可讨论的研究视角。"
        "需要保持克制的是：该视角建立在研究原型与阶段性证据之上，尚不能替代完整系统安全证明或生产级性能评估。",
    )


def append_references(doc: Document):
    anchor = find_para_exact(doc, "参考文献")
    if anchor is None:
        return
    p = anchor
    for ref in NEW_REFERENCES:
        p = add_body(doc, p, ref, first_line=False)


def enhance():
    backup = BACKUP_DIR / f"Hybrid_TEE_Rollup_before_academic_enhance_{datetime.now():%Y%m%d_%H%M%S}.docx"
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(DOCX_PATH, backup)

    doc = Document(str(DOCX_PATH))
    renumber_tables(doc)
    add_contributions_and_capability_table(doc)
    add_formal_definitions(doc)
    add_recovery_overhead_analysis(doc)
    restructure_threats_to_validity(doc)
    enhance_related_work_opening(doc)
    add_theoretical_implication(doc)
    append_references(doc)

    doc.save(str(DOCX_PATH))
    return backup


if __name__ == "__main__":
    backup = enhance()
    doc = Document(str(DOCX_PATH))
    refs = sum(1 for p in doc.paragraphs if p.text.strip().startswith("["))
    print(f"Done. References count (approx): {refs}")
    print(f"Backup: {backup}")
    print(f"Saved: {DOCX_PATH}")
