# -*- coding: utf-8 -*-
"""最后一轮图表可读性优化：重绘图1-4、放大图5-9、表2原生表格、统一版式。"""

from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path

import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
DRAFT = ROOT / "project_code" / "paper_outputs" / "draft"
DOCX_PATH = DRAFT / "Hybrid_TEE_Rollup_中文论文初稿.docx"
SUMMARY_PATH = ROOT / "project_code" / "paper_outputs" / "summary.json"
ASSETS = DRAFT / "readability_assets_20260528"
BACKUP_DIR = DRAFT / "backups"

CONTENT_WIDTH_IN = 6.83
FIG_DPI = 200

# 版式目标宽度（英寸）
W_ARCH = CONTENT_WIDTH_IN * 0.80
W_THREAT = CONTENT_WIDTH_IN * 0.80
W_SEQ = CONTENT_WIDTH_IN * 0.75
W_STATE = CONTENT_WIDTH_IN * 0.85
W_EXP = CONTENT_WIDTH_IN * 0.75
W_TABLE2 = CONTENT_WIDTH_IN * 0.88

CAPABILITY_ROWS = [
    ["System", "Interactive Proof", "TEE Support", "Recovery", "Fact Preservation", "Arbitration Liveness"],
    ["Arbitrum", "✓", "✗", "✗", "✗", "△"],
    ["Optimism", "✓", "✗", "✗", "✗", "△"],
    ["Cartesi", "✓", "✗", "✗", "✗", "△"],
    ["opML", "✓", "✗", "✗", "✗", "△"],
    ["RCP (Ours)", "✓", "✓", "✓", "✓", "✓"],
]

FIGURE_FILES = {
    "fig1_threat_model": "fig01_threat_model.png",
    "fig2_architecture": "fig02_architecture.png",
    "fig3_sequence": "fig03_sequence.png",
    "fig4_state_machine": "fig04_state_machine.png",
    "fig5_recovery": "fig05_recovery_success.png",
    "fig6_challenge_rounds": "fig06_challenge_rounds.png",
    "fig7_payload": "fig07_cost_payload.png",
    "fig8_da_gas": "fig08_da_gas.png",
    "fig9_failure": "fig09_failure_detection.png",
}

CAPTION_TO_FILE = {
    "图 5 恢复对超时场景挑战完成率的影响": "fig05_recovery_success.png",
    "图 6 二分挑战轮次与 trace steps 的关系": "fig06_challenge_rounds.png",
    "图 7 full payload 与 compact commit 的字节规模对比": "fig07_cost_payload.png",
    "图 8 不同 DA profile 下的摊销 gas 趋势": "fig08_da_gas.png",
    "图 9 失败场景检测与仲裁结果": "fig09_failure_detection.png",
}

REPLACE_BY_CAPTION = {
    "图 1 Threat Model of RCP": ("fig01_threat_model.png", W_THREAT),
    "图 2 可恢复挑战驱动的 Hybrid TEE-Rollup 原型架构": ("fig02_architecture.png", W_ARCH),
    "图 3 RCP 协议交互时序": ("fig03_sequence.png", W_SEQ),
    "图 4 可恢复挑战协议状态机": ("fig04_state_machine.png", W_STATE),
}


def setup_matplotlib():
    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Arial", "DejaVu Sans", "Microsoft YaHei", "SimHei"],
            "font.size": 11,
            "axes.unicode_minus": False,
        }
    )


def save_fig(fig, name: str):
    path = ASSETS / name
    fig.savefig(path, dpi=FIG_DPI, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def draw_round_box(ax, xy, w, h, text, fc, ec, fontsize=11, weight="normal"):
    box = FancyBboxPatch(
        xy,
        w,
        h,
        boxstyle="round,pad=0.02,rounding_size=0.08",
        linewidth=1.6,
        edgecolor=ec,
        facecolor=fc,
    )
    ax.add_patch(box)
    ax.text(xy[0] + w / 2, xy[1] + h / 2, text, ha="center", va="center", fontsize=fontsize, weight=weight)


def arrow_down(ax, x, y1, y2):
    ax.annotate("", xy=(x, y2), xytext=(x, y1), arrowprops=dict(arrowstyle="-|>", color="#333333", lw=1.8))


def generate_fig1_threat_model():
    fig, ax = plt.subplots(figsize=(5.5, 11))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 22)
    ax.axis("off")
    cx = 5.0
    bw, bh = 4.2, 0.9
    y = 20.5
    draw_round_box(ax, (cx - bw / 2, y), bw, bh, "Attacker", "#fde8e8", "#b42318", fontsize=12, weight="bold")
    threats = ["Fake Response", "DA Tampering", "Invalid Attestation", "Timeout Attack"]
    y = 17.2
    for label in threats:
        draw_round_box(ax, (cx - bw / 2, y), bw, bh, label, "#fff4e5", "#b54708", fontsize=11)
        arrow_down(ax, cx, y + bh, y - 0.35)
        y -= 1.35
    arrow_down(ax, cx, 12.8, 11.6)
    draw_round_box(ax, (cx - 1.6, 10.4), 3.2, 1.0, "RCP", "#e8f1ff", "#175cd3", fontsize=13, weight="bold")
    outcomes = [
        ("Recover", "#e8f8ef", "#067647"),
        ("Replay", "#eef4ff", "#3538cd"),
        ("Resolve", "#e8f8ef", "#067647"),
        ("Slash", "#fdecec", "#b42318"),
    ]
    y = 8.0
    for label, fc, ec in outcomes:
        draw_round_box(ax, (cx - bw / 2, y), bw, bh, label, fc, ec, fontsize=11)
        if y > 4.5:
            arrow_down(ax, cx, y, y - 0.35)
        y -= 1.35
    ax.text(5, 21.6, "Threat Model of RCP", ha="center", fontsize=14, weight="bold")
    return save_fig(fig, FIGURE_FILES["fig1_threat_model"])


def generate_fig2_architecture():
    fig, ax = plt.subplots(figsize=(6.8, 8.5))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 12)
    ax.axis("off")
    layers = [
        ("Execution Layer", ["Proposer", "Executor", "TEE Attestation"], "#e8f1ff", "#175cd3", 9.2),
        ("Data Availability Layer", ["Payload", "DA Root", "Merkle Proof"], "#fff6e8", "#b54708", 5.8),
        ("Verification Layer", ["Commit", "Challenge", "Recover", "Replay", "Resolve"], "#e8f8ef", "#067647", 1.8),
    ]
    for title, items, fc, ec, y0 in layers:
        draw_round_box(ax, (0.5, y0), 9.0, 2.6, "", fc, ec)
        ax.text(5.0, y0 + 2.25, title, ha="center", fontsize=12, weight="bold")
        n = len(items)
        gap = 8.4 / n
        for i, item in enumerate(items):
            x = 0.9 + i * gap
            draw_round_box(ax, (x, y0 + 0.35), gap - 0.25, 1.35, item, "white", ec, fontsize=10)
        if y0 > 2.0:
            arrow_down(ax, 5.0, y0, y0 - 0.55)
    ax.text(5, 11.5, "Hybrid TEE-Rollup Prototype Architecture", ha="center", fontsize=13, weight="bold")
    return save_fig(fig, FIGURE_FILES["fig2_architecture"])


def generate_fig3_sequence():
    fig, ax = plt.subplots(figsize=(7.0, 5.5))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 8)
    ax.axis("off")
    actors = ["Proposer", "Watcher", "Smart Contract", "DA Layer"]
    xs = [1.2, 3.6, 6.0, 8.4]
    for x, name in zip(xs, actors):
        ax.text(x, 7.5, name, ha="center", fontsize=11, weight="bold")
        ax.plot([x, x], [0.8, 7.0], color="#98a2b3", lw=1.2, linestyle="--")
    messages = [
        (0, 2, 6.4, "Submit"),
        (1, 2, 5.7, "Challenge"),
        (2, 0, 5.0, "Timeout"),
        (2, 1, 4.3, "Recover"),
        (1, 2, 3.6, "Replay"),
        (2, 3, 2.9, "Fetch DA"),
        (2, 1, 2.2, "Resolve"),
    ]
    for frm, to, y, label in messages:
        x1, x2 = xs[frm], xs[to]
        style = "-|>" if frm < to else "<|-"
        ax.annotate(
            "",
            xy=(x2, y),
            xytext=(x1, y),
            arrowprops=dict(arrowstyle=style, color="#344054", lw=1.6),
        )
        ax.text((x1 + x2) / 2, y + 0.12, label, ha="center", fontsize=10, weight="bold")
    ax.text(5, 7.85, "RCP Protocol Interaction (Simplified)", ha="center", fontsize=12, weight="bold")
    return save_fig(fig, FIGURE_FILES["fig3_sequence"])


def generate_fig4_state_machine():
    fig, ax = plt.subplots(figsize=(14.0, 4.2))
    ax.set_xlim(0, 28)
    ax.set_ylim(0, 6)
    ax.axis("off")
    main_states = [
        "OPEN",
        "RESPONDED",
        "NARROWING",
        "TIMED_OUT",
        "RECOVERED",
        "READY_FOR_REPLAY",
        "REPLAYED",
    ]
    x = 0.3
    bw, bh = 3.0, 1.05
    gap = 0.45
    centers = []
    for state in main_states:
        label = state.replace("_", "\n") if len(state) > 10 else state
        draw_round_box(ax, (x, 2.5), bw, bh, label, "#eef4ff", "#3538cd", fontsize=9.5, weight="bold")
        centers.append(x + bw / 2)
        if state != main_states[-1]:
            ax.annotate(
                "",
                xy=(x + bw + gap, 3.02),
                xytext=(x + bw, 3.02),
                arrowprops=dict(arrowstyle="-|>", color="#344054", lw=1.6),
            )
        x += bw + gap
    fork_x = centers[-1]
    draw_round_box(ax, (fork_x + 2.0, 4.05), 2.6, 0.95, "RESOLVED", "#e8f8ef", "#067647", fontsize=10, weight="bold")
    draw_round_box(ax, (fork_x + 2.0, 1.25), 2.6, 0.95, "SLASHED", "#fdecec", "#b42318", fontsize=10, weight="bold")
    ax.annotate("", xy=(fork_x + 1.2, 4.5), xytext=(fork_x + 0.4, 3.5), arrowprops=dict(arrowstyle="-|>", color="#344054", lw=1.6))
    ax.annotate("", xy=(fork_x + 1.2, 1.7), xytext=(fork_x + 0.4, 2.6), arrowprops=dict(arrowstyle="-|>", color="#344054", lw=1.6))
    ax.text(14, 5.55, "Recoverable Challenge State Machine", ha="center", fontsize=12, weight="bold")
    return save_fig(fig, FIGURE_FILES["fig4_state_machine"])


def load_summary():
    with SUMMARY_PATH.open(encoding="utf-8") as f:
        return json.load(f)


def generate_experiment_figures(summary):
    import matplotlib.pyplot as plt

    paths = {}
    cost = summary["cost_summary"]
    challenge = summary["challenge_summary"]
    failure = summary.get("failure_summary", [])

    # Fig 5 - recovery
    no_recover = next((r for r in failure if r["scenario"] == "challenge_timeout_no_recover"), None)
    recover = next((r for r in failure if r["scenario"] == "challenge_timeout_recover"), None)
    fig, ax = plt.subplots(figsize=(7.5, 5.0))
    vals = [
        0.0 if no_recover is None else no_recover["challenge_success_rate"] * 100.0,
        0.0 if recover is None else recover["challenge_success_rate"] * 100.0,
    ]
    bars = ax.bar(["no recover", "recover"], vals, color=["#84a9ff", "#175cd3"], width=0.55)
    ax.set_ylabel("Challenge success rate (%)", fontsize=11)
    ax.set_title("Timeout Recovery Effect", fontsize=13, weight="bold")
    ax.set_ylim(0, 110)
    ax.tick_params(labelsize=10)
    for bar, v in zip(bars, vals):
        ax.text(bar.get_x() + bar.get_width() / 2, v + 2, f"{v:.0f}%", ha="center", fontsize=10)
    ax.grid(axis="y", alpha=0.25, linestyle="--")
    paths["fig5"] = save_fig(fig, FIGURE_FILES["fig5_recovery"])

    default_challenges = sorted(
        [
            row
            for row in challenge
            if row["challenge_timeout"] == min(item["challenge_timeout"] for item in challenge)
            and row["max_bisection_rounds"] == max(item["max_bisection_rounds"] for item in challenge)
        ],
        key=lambda row: row["trace_steps"],
    )
    fig, ax = plt.subplots(figsize=(7.5, 5.0))
    xs = [row["trace_steps"] for row in default_challenges]
    ax.plot(xs, [row["avg_bisection_rounds"] for row in default_challenges], marker="o", linewidth=2.2, label="observed", color="#175cd3")
    ax.plot(xs, [row["avg_expected_log2_rounds"] for row in default_challenges], marker="x", linewidth=2.0, label="log2 baseline", color="#b42318")
    ax.set_xlabel("Trace steps", fontsize=11)
    ax.set_ylabel("Bisection rounds", fontsize=11)
    ax.set_title("Challenge Bisection Rounds", fontsize=13, weight="bold")
    ax.legend(fontsize=10)
    ax.tick_params(labelsize=10)
    ax.grid(alpha=0.25, linestyle="--")
    paths["fig6"] = save_fig(fig, FIGURE_FILES["fig6_challenge_rounds"])

    cost_by_payload = {}
    for row in cost:
        if row["prompt_length"] != min(item["prompt_length"] for item in cost):
            continue
        if row["batch_size"] != 1:
            continue
        cost_by_payload[row["payload_size"]] = row
    payloads = sorted(cost_by_payload)
    full = [cost_by_payload[p]["avg_full_bytes"] for p in payloads]
    compact = [cost_by_payload[p]["avg_compact_bytes"] for p in payloads]
    fig, (ax_main, ax_zoom) = plt.subplots(1, 2, figsize=(10.5, 5.0), gridspec_kw={"width_ratios": [1.45, 1.0]})
    ax_main.plot(payloads, full, marker="o", linewidth=2.2, label="full payload")
    ax_main.plot(payloads, compact, marker="o", linewidth=2.2, label="compact commit")
    ax_main.set_xlabel("Payload size", fontsize=11)
    ax_main.set_ylabel("Bytes", fontsize=11)
    ax_main.set_title("Overall trend", fontsize=12, weight="bold")
    ax_main.legend(fontsize=10)
    ax_main.tick_params(labelsize=10)
    ax_main.grid(alpha=0.25, linestyle="--")
    ax_zoom.plot(payloads, compact, marker="o", color="#ff7f0e", linewidth=2.2, label="compact commit")
    ax_zoom.set_xlabel("Payload size", fontsize=11)
    ax_zoom.set_ylabel("Compact bytes", fontsize=11)
    ax_zoom.set_title("Compact commit zoom-in", fontsize=12, weight="bold")
    ax_zoom.tick_params(labelsize=10)
    ax_zoom.grid(alpha=0.25, linestyle="--")
    cmin, cmax = min(compact), max(compact)
    ax_zoom.set_ylim(max(0, cmin - 30), cmax + 30)
    for x, y in zip(payloads, compact):
        ax_zoom.annotate(f"{y:.0f}B", (x, y), textcoords="offset points", xytext=(0, 6), ha="center", fontsize=9)
    fig.suptitle("Full-onchain vs Compact Commit", fontsize=13, weight="bold")
    fig.tight_layout(rect=[0, 0, 1, 0.94])
    paths["fig7"] = save_fig(fig, FIGURE_FILES["fig7_payload"])

    batch_rows = sorted(
        [
            row
            for row in cost
            if row["payload_size"] == max(item["payload_size"] for item in cost)
            and row["prompt_length"] == min(item["prompt_length"] for item in cost)
        ],
        key=lambda row: row["batch_size"],
    )
    fig, ax = plt.subplots(figsize=(7.5, 5.0))
    for key, label, color in [
        ("avg_full_onchain_calldata_amortized_gas", "full onchain", "#344054"),
        ("avg_compact_external_da_amortized_gas", "external DA", "#175cd3"),
        ("avg_compact_eip4844_like_amortized_gas", "EIP-4844-like", "#7a5af8"),
        ("avg_compact_modular_da_sampling_amortized_gas", "modular DA", "#12b76a"),
    ]:
        ax.plot(
            [row["batch_size"] for row in batch_rows],
            [row[key] for row in batch_rows],
            marker="o",
            linewidth=2.0,
            label=label,
            color=color,
        )
    ax.set_xscale("log")
    ax.set_xlabel("Batch size", fontsize=11)
    ax.set_ylabel("Amortized gas", fontsize=11)
    ax.set_title("DA Profile Amortized Gas", fontsize=13, weight="bold")
    ax.legend(fontsize=9)
    ax.tick_params(labelsize=10)
    ax.grid(alpha=0.25, linestyle="--")
    paths["fig8"] = save_fig(fig, FIGURE_FILES["fig8_da_gas"])

    fig, ax = plt.subplots(figsize=(9.0, 5.5))
    scenarios = [row["scenario"] for row in failure]
    ax.bar(scenarios, [row["detection_rate"] * 100.0 for row in failure], color="#175cd3")
    ax.set_ylabel("Detection rate (%)", fontsize=11)
    ax.set_title("Failure Scenario Detection", fontsize=13, weight="bold")
    ax.tick_params(axis="x", labelsize=9, rotation=28)
    ax.tick_params(axis="y", labelsize=10)
    ax.grid(axis="y", alpha=0.25, linestyle="--")
    fig.tight_layout()
    paths["fig9"] = save_fig(fig, FIGURE_FILES["fig9_failure"])
    return paths


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True
    p_pr = paragraph._element.get_or_add_pPr()
    keep = OxmlElement("w:keepLines")
    keep.set(qn("w:val"), "1")
    p_pr.append(keep)


def clear_paragraph(paragraph):
    element = paragraph._element
    for child in list(element):
        if child.tag.endswith("r") or child.tag.endswith("drawing"):
            element.remove(child)


def set_paragraph_image(paragraph, image_path: Path, width_in: float):
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    set_keep_with_next(paragraph)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_in))


def insert_image_before(paragraph, image_path: Path, width_in: float):
    new_p = OxmlElement("w:p")
    paragraph._element.addprevious(new_p)
    from docx.text.paragraph import Paragraph

    img_para = Paragraph(new_p, paragraph._parent)
    set_paragraph_image(img_para, image_path, width_in)
    return img_para


def style_caption(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(10)
    set_keep_with_next(paragraph)
    for run in paragraph.runs:
        run.bold = True
        run.font.size = Pt(10.5)


def set_cell_shading(cell, fill: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def add_capability_table_after(paragraph, doc: Document):
    """在表2标题后插入原生表格，替换原图片段。"""
    tbl = doc.add_table(rows=len(CAPABILITY_ROWS), cols=len(CAPABILITY_ROWS[0]))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths = [1.05, 1.15, 0.95, 0.85, 1.15, 1.25]
    for ri, row_data in enumerate(CAPABILITY_ROWS):
        for ci, text in enumerate(row_data):
            cell = tbl.rows[ri].cells[ci]
            cell.text = text
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(9 if ri else 10)
                    run.font.name = "Times New Roman"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                    if ri == 0:
                        run.bold = True
            if ri == 0:
                set_cell_shading(cell, "D9E8F8")
            elif row_data[0] == "RCP (Ours)":
                set_cell_shading(cell, "E8F8EF")
    paragraph._element.addnext(tbl._tbl)
    # 设置列宽
    for row in tbl.rows:
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(col_widths[ci])


def format_all_tables(doc: Document):
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, row in enumerate(table.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    for run in p.runs:
                        if run.font.size is None or run.font.size < Pt(9):
                            run.font.size = Pt(9)
                        if ri == 0:
                            run.bold = True


def para_has_image(paragraph) -> bool:
    return "pic:pic" in paragraph._element.xml or "w:drawing" in paragraph._element.xml


def find_prev_image_paragraph(doc: Document, caption_text: str):
    target_idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == caption_text:
            target_idx = i
            break
    if target_idx is None:
        return None
    for j in range(target_idx - 1, -1, -1):
        if para_has_image(doc.paragraphs[j]):
            return doc.paragraphs[j]
        if doc.paragraphs[j].text.strip():
            break
    return None


def apply_docx_updates():
    backup = BACKUP_DIR / f"Hybrid_TEE_Rollup_中文论文初稿_before_readability_{datetime.now():%Y%m%d_%H%M%S}.docx"
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(DOCX_PATH, backup)

    doc = Document(str(DOCX_PATH))

    # 表2：删除标题后的图片段，插入原生表
    table2_caption = None
    table2_image = None
    caption_idx = -1
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t == "表 2 现有方案与 RCP 的能力对比":
            table2_caption = p
            caption_idx = i
        if caption_idx >= 0 and i > caption_idx and para_has_image(p) and table2_image is None:
            table2_image = p
            break
    if table2_image is not None:
        table2_image._element.getparent().remove(table2_image._element)
    if table2_caption is not None:
        style_caption(table2_caption)
        add_capability_table_after(table2_caption, doc)

    # 图1-4：替换紧邻图题前的图片段
    for caption, (fname, width) in REPLACE_BY_CAPTION.items():
        img_path = ASSETS / fname
        img_para = find_prev_image_paragraph(doc, caption)
        cap_para = next((p for p in doc.paragraphs if p.text.strip() == caption), None)
        if cap_para is None:
            continue
        style_caption(cap_para)
        if img_para is not None:
            set_paragraph_image(img_para, img_path, width)
        else:
            insert_image_before(cap_para, img_path, width)

    # 图5-9：在图题前插入（若缺失）
    for caption, fname in CAPTION_TO_FILE.items():
        img_path = ASSETS / fname
        cap_para = next((p for p in doc.paragraphs if p.text.strip() == caption), None)
        if cap_para is None:
            continue
        style_caption(cap_para)
        img_para = find_prev_image_paragraph(doc, caption)
        if img_para is not None:
            set_paragraph_image(img_para, img_path, W_EXP)
        else:
            insert_image_before(cap_para, img_path, W_EXP)

    format_all_tables(doc)
    doc.save(str(DOCX_PATH))
    return backup


def main():
    setup_matplotlib()
    ASSETS.mkdir(parents=True, exist_ok=True)
    print("Generating figures 1-4...")
    generate_fig1_threat_model()
    generate_fig2_architecture()
    generate_fig3_sequence()
    generate_fig4_state_machine()
    print("Generating experiment figures 5-9...")
    summary = load_summary()
    generate_experiment_figures(summary)
    print("Updating DOCX...")
    backup = apply_docx_updates()
    print(f"Done. Backup: {backup}")
    print(f"Output: {DOCX_PATH}")
    print(f"Assets: {ASSETS}")


if __name__ == "__main__":
    main()
