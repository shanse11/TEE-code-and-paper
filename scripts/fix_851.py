from docx import Document
doc = Document(r"E:\项目\project_code\paper_outputs\draft\final_chinese_academic_format.docx")

for i, p in enumerate(doc.paragraphs):
    if "8.5.1" in p.text[:20] and "理论启示" in p.text:
        content_para = doc.paragraphs[i + 1]
        content_para.text = (
            "从理论层面看，本文更有价值的部分不是提出新的 Rollup 系统，而是将挑战恢复从工程补丁提升为具有协议语义的独立机制。"
            "可恢复挑战机制的核心理论贡献在于提出并论证了仲裁连续性这一协议性质："
            "恢复操作不是简单的超时重试，而是在保持挑战事实不变的条件下恢复协议活性的状态转移。"
            "携带证据的紧凑提交与 DA 感知成本模型的组合，"
            "则为混合可信执行环境汇总在高频去中心化应用场景下提供了一条正常路径轻量、异常路径可验证的理论设计路线。"
            "本文通过协议性质分析（第 4.4 节）系统化了挑战事实不变性、安全性与活性的分离、仲裁连续性和恢复操作正确性四项性质，"
            "并给出了协议不变量表格与抗滥用分析，为后续将恢复机制推广到更复杂的汇总协议栈提供了理论参照。"
        )
        print(f"Fixed 8.5.1 at P{i+1}")
        break

doc.save(r"E:\项目\project_code\paper_outputs\draft\final_chinese_academic_format.docx")
print("Saved.")
