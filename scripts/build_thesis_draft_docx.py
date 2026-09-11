# -*- coding: utf-8 -*-

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "project_code" / "paper_outputs" / "draft"
DOCX_PATH = OUT_DIR / "Hybrid_TEE_Rollup_中文论文初稿_标准格式.docx"
FALLBACK_DOCX_PATH = OUT_DIR / "Hybrid_TEE_Rollup_中文论文初稿_标准格式_更新版.docx"
MD_PATH = OUT_DIR / "Hybrid_TEE_Rollup_中文论文初稿_标准格式.md"
FIGURE_DIR = ROOT / "project_code" / "paper_outputs" / "figures"
ARCH_FIGURE_PATH = FIGURE_DIR / "system_architecture.png"
PROTOCOL_FIGURE_PATH = FIGURE_DIR / "recoverable_challenge_flow.png"


TITLE = "面向高频 DApp 的 Hybrid TEE-Rollup 可恢复挑战协议与轻量 DA 成本评估"
EN_TITLE = "Recoverable Challenge Protocol and Lightweight DA Cost Evaluation for Hybrid TEE-Rollups in High-Frequency DApps"


ABSTRACT = (
    "高频 DApp，如元宇宙状态同步、去中心化社交交互和链上 AI 推理服务，要求系统同时具备低交互成本、较低延迟和异常情况下的可验证性。"
    "全链上执行虽然可验证性强，但难以承受高频负载；完全链下执行虽能降低成本，却削弱了公开可验证与可追责能力；单纯依赖 TEE 又会引入硬件信任与可用性风险。"
    "本文不提出一个全新的 Rollup 架构，而是在已有 Hybrid TEE-Rollup 思路下，研究高频 DApp 场景中的可验证低成本交互机制，并将问题明确限定在研究原型和课程论文评估范围内。"
    "具体而言，本文设计并实现了一个研究原型，重点推进三项机制：一是 recoverable challenge，将交互式挑战从错误检测扩展为异常场景下可恢复并完成仲裁的状态机；"
    "二是 evidence-carrying compact commit，将轻量提交与 DA Merkle root、payload hash、Merkle proof 和 challenge evidence 绑定为最小证据入口；"
    "三是 DA-aware cost amortization，从 payload size、batch size 和 DA route 的共同作用下分析成本摊销。"
    "实验表明，compact commit 的字节开销近似固定，而 full payload 随负载规模近似线性增长；在高 payload 设置下，compact commit + DA 相对 full-onchain 呈现明显降本趋势；"
    "recover 机制将 timeout 场景下的 challenge success 从 0% 提升到 100%；二分定位轮次与 log2(trace steps) 基本一致。"
    "同时，本文补充威胁模型、成本公式、实验设置和轻量性质讨论，使实验结果与论文主张形成更清晰的对应关系。"
    "本文还给出本地 EVM gasUsed 基线与 Sepolia 部署记录，作为阶段性链上对应物证据。"
    "需要强调的是，当前系统仍为研究原型，其中 TEE 为 simulated TEE，DA 为 mock/verifiable DA 原型，本地 EVM 与 Sepolia 结果不能直接外推为主网生产性能结论。"
)

KEYWORDS = "Hybrid TEE-Rollup；交互式挑战；数据可用性；高频 DApp；成本摊销；故障恢复"


SECTIONS = [
    (
        "1 Introduction",
        [
            "区块链系统在面向高频 DApp 时面临一个结构性矛盾：应用希望获得接近传统互联网服务的低延迟、高吞吐和低单次交互成本，但底层系统又必须维持公开可验证、去中心化安全和异常可追责能力。对于元宇宙资产交互、去中心化社交状态更新、链上 AI 推理请求等场景而言，单次交互可能价值较低，但请求频率高、状态连续性强。如果每一次交互都以完整链上执行和全量数据上链的方式处理，系统成本会快速放大；如果完全转向链下服务，则执行正确性、数据可用性和异常追责又缺乏足够的公开验证路径。",
            "Rollup 技术为这一矛盾提供了一类中间路线。ZK-Rollup 通过有效性证明获得较强验证能力，但证明生成和系统实现成本较高；Optimistic Rollup 降低了常规执行路径成本，但挑战窗口和争议处理流程会带来延迟。TEE-Rollup 或 Hybrid TEE-Rollup 则试图利用可信执行环境加速链下执行，并通过链上提交和挑战机制保留异常情况下的仲裁能力。然而，TEE 本身并不应被视为绝对可信：硬件实现、宿主环境、I/O 操控、数据可用性和挑战超时都可能成为系统风险来源。因此，Hybrid TEE-Rollup 的关键问题不只是“如何把执行放到 TEE 中”，而是如何在低成本链下执行之后，仍然保留足够的可验证证据入口和异常恢复能力。",
            "本文基于这一背景，将研究切口收敛为两个具体问题。第一，交互式挑战在异常场景下是否能够恢复推进。传统 challenge 机制通常关注错误是否可被发现，而在 timeout、中断、执行轨迹不一致或数据不可用时，争议流程可能停滞。对于高频 DApp 而言，仅能检测异常并不足够，系统还需要将异常推进到可仲裁、可结算的状态。第二，轻量提交如何在降低链上成本的同时保留验证能力。若 compact commit 只是一个孤立哈希，则虽然节省链上空间，却难以支撑后续挑战；因此，轻量提交需要与 DA root、payload hash、Merkle proof 和 challenge evidence 建立明确关联。",
            "围绕上述问题，本文的课程论文贡献包括：设计可恢复交互式挑战原型，补充 evidence-carrying compact commit 与字段级 DA proof，建立 DA-aware 成本模型与实验网格，并给出 Python 原型、本地 EVM gasUsed 和 Sepolia 部署记录构成的阶段性证据链。本文所有结论均限定为研究原型结果，不声称真实 TEE 安全、生产级 DA 或主网性能。",
        ],
    ),
    (
        "2 Background",
        [
            "Rollup 是区块链扩容中的重要路线，其基本思想是将大量执行过程移至链下完成，并将状态承诺、证明或争议入口提交到链上。ZK-Rollup 通过有效性证明保证提交状态的正确性，具有较强的验证属性，但证明生成、验证电路设计和复杂系统实现带来较高成本。Optimistic Rollup 假设链下执行结果默认正确，只有在挑战期内出现争议时才进入 fraud proof 或交互式验证流程，因此常规路径成本较低，但结算延迟和挑战机制复杂度成为关键限制。",
            "TEE attestation 提供了一种链下执行可信性的补充方式。可信执行环境可以在隔离环境中执行程序，并对输入、输出、环境标识和 nonce 生成证明，使外部验证者能够确认某一输出来自指定执行环境。然而，TEE 并非无条件可信。硬件漏洞、实现缺陷、宿主环境操控和可用性问题都可能影响系统安全。因此，TEE 更适合作为降低正常路径验证成本的组件，而不应替代链上挑战、数据可用性和故障追责机制。",
            "Optimistic challenge 是 Rollup 系统中处理异常的重要机制。典型流程包括提交状态承诺、开启挑战、交互式缩小争议范围，并在必要时执行单步重放或链上仲裁。对于长执行轨迹，二分定位能够将争议定位复杂度从线性扫描降低到对数轮次。然而，challenge 本身也可能遭遇 timeout、应答缺失或证据不可用等异常，因此其活性和恢复性同样需要被纳入协议设计。",
            "Data Availability 关注链下或链外数据是否可被需要的验证者获得。模块化区块链和 LazyLedger 等路线强调将数据可用性从执行验证中解耦，使共识层主要负责数据排序和可用性保证。对于 Hybrid TEE-Rollup 而言，DA 的作用不仅是降低存储成本，还要支撑后续 challenge 的证据回溯。因此，compact commit 必须与 DA proof 建立可验证联系。",
        ],
    ),
    (
        "3 Motivation",
        [
            "纯全链上执行并不适合高频 DApp。若将每一次社交状态更新、元宇宙交互或 AI 推理请求都作为完整链上交易执行并存储，则链上 calldata、存储和执行成本会随请求数量快速增长。对于大量小额、高频、状态连续的交互而言，这种成本结构难以长期维持。",
            "纯链下执行同样不足。中心化服务或普通链下节点可以降低成本，但用户和外部观察者难以验证执行结果。一旦执行者篡改 response、隐藏 payload、延迟应答或在争议中失联，系统缺少公开可验证的仲裁路径。对于区块链应用而言，这会削弱去中心化系统最核心的可信边界。",
            "纯 TEE 路线也存在局限。TEE 能够降低链下执行的信任成本，但不能消除所有故障。TEE attestation 可能无效，执行者可能提交与 DA payload 不一致的结果，宿主环境可能影响可用性，挑战过程也可能因 timeout 停滞。因此，将 TEE 作为唯一安全来源并不稳健。更合理的方式是让 TEE 服务于正常路径的低成本执行，同时保留链上提交、DA 证明和 challenge 仲裁作为异常路径。",
            "纯 Optimistic Rollup 机制则面临挑战窗口和异常处理问题。传统 optimistic 机制能够在争议出现时启动验证，但如果 challenge session 因超时或应答中断而停滞，系统可能只能发现异常，却无法完成仲裁。对于高频 DApp，协议活性与可恢复性是实际可用性的组成部分。因此，本文将 recoverable challenge 作为独立研究问题：challenge 不应只是错误检测工具，而应是异常场景下仍可恢复推进的容错状态机。",
        ],
    ),
    (
        "4 Problem Statement and Threat Model",
        [
            "本文关注的问题可以概括为：在高频 DApp 场景中，如何让链下执行结果以较低链上成本提交，同时在异常发生时仍能被公开验证、恢复推进并完成仲裁。该问题同时包含成本、可验证性和活性三个维度。成本维度要求链上提交不随完整 payload 线性增长；可验证性维度要求 compact commit 能指向可检查的 DA payload 与证明；活性维度要求 challenge 在 timeout 后不永久停滞。",
            "本文假设执行者或提交者可能提交错误 response、篡改 DA payload、提供无效 attestation，或在挑战流程中延迟响应。DA 层可能出现 payload 缺失、Merkle proof 损坏或 root 不一致。挑战流程本身也可能因超时、中断或 trace 长度不一致而无法自然进入 resolve 阶段。验证者、挑战者或 watchdog 可以观察链上提交和 DA 证据，并在发现异常时发起或恢复挑战。",
            "本文不覆盖真实硬件 TEE 的侧信道攻击、远程证明供应链、真实 DA 网络的经济安全、主网 MEV 或拥堵环境下的费用波动，也不证明完整 Rollup 安全性。原型中的 TEE 为 simulated TEE，DA 为 mock/verifiable DA registry，Sepolia 部署仅用于说明合约具有公开测试网可部署性。该边界是本文论证成立的前提。",
        ],
    ),
    (
        "5 System Design",
        [
            "本文原型采用三层结构：执行层、数据层和验证层。整体逻辑是链下执行、链上留痕、异常时可恢复验证。该设计不声称实现完整生产级 Rollup，而是用于验证 Hybrid TEE-Rollup 路线下 recoverable challenge 与 verifiable DA 的机制可行性。",
        ],
    ),
    (
        "5.1 执行层",
        [
            "执行层负责处理高频 DApp 请求。当前原型采用 deterministic mock model 或可选模型后端生成响应，并使用 simulated TEE attestation 对输入哈希、输出哈希、nonce 和执行环境标识进行绑定。attestation 的作用是为链下执行结果提供最小可信摘要，使后续验证层能够检查提交结果是否与声明的执行过程一致。",
            "需要明确的是，当前 TEE 为模拟实现，并不代表真实 SGX、TDX 或其他硬件 TEE 的安全保证。本文使用 simulated TEE 的目的，是在研究原型中建立“链下执行摘要—链上提交—异常挑战”的接口关系，而不是证明真实硬件安全。",
        ],
    ),
    (
        "5.2 数据层",
        [
            "数据层的核心是 evidence-carrying compact commit。系统不将完整 payload 全量上链，而是在链上提交精简结构：state_root、output_hash、proof_hash、da_pointer 和 da_merkle_root。其中，state_root 表示输入、输出和环境摘要绑定后的状态承诺；output_hash 绑定执行结果；proof_hash 绑定 attestation 证明；da_pointer 指向链外 DA payload；da_merkle_root 则为字段级 DA proof 提供根承诺。",
            "DA entry 中保存完整 payload 及其可验证结构，包括 prompt、response、attestation、payload hash、leaf hashes、Merkle root 和字段级 Merkle proof。这样，compact commit 不再只是为了降低链上字节数的摘要，而成为后续 challenge 的最小证据入口。当验证者需要检查 response、prompt 或 attestation 是否被篡改时，可以通过 DA pointer 获取 payload，并通过 Merkle proof 与链上 da_merkle_root 进行一致性验证。",
        ],
    ),
    (
        "5.3 验证层",
        [
            "验证层负责处理异常路径。当前原型实现了 challenge open、respond、step、recover、replay 和 resolve 等状态推进。其目标不是将所有执行都搬到链上，而是在出现争议时，通过证据加载、二分定位、单步重放和仲裁结算，将异常状态推进为可解释的结果。",
            "验证层同时承担 fault taxonomy 的职责。本文将故障细分为 attestation fault、DA fault、timeout fault、replay fault 和 trace inconsistency。不同故障对应不同的观察和处理路径。例如，DA unavailable 表明数据层无法提供 payload；DA proof invalid 表明 payload 与链上 root 不一致；trace length mismatch 表明执行轨迹与声明结构不一致；response tampering 则可能同时导致 payload hash、attestation 和 state root mismatch。",
            "在 Solidity 映射中，原型包含 MockTEEVerifier、MockDARegistry、MerkleVerifier 和 HybridTEERollup 等合约，用于表达链上状态推进和最小验证逻辑。本地 EVM 合约并非 Python 原型的逐行移植，而是抽取适合链上表达的提交、验证、挑战、恢复和结算路径。",
        ],
    ),
    (
        "6 Recoverable Challenge Protocol",
        [
            "Recoverable Challenge Protocol 的目标，是将 challenge 从一次性错误检测流程扩展为异常场景下可恢复推进的协议状态机。其基本状态包括 PENDING、OPEN、RESPONDED、NARROWING、RECOVERED、READY_FOR_REPLAY、REPLAYED、RESOLVED、SLASHED 和 FINALIZED。",
            "当交易处于 PENDING 状态时，挑战者可以调用 open 进入 OPEN 状态。系统加载 compact commit、DA status 和相关 mismatch evidence。随后，被挑战方进入 respond 阶段，提供其对争议的回应。若争议涉及执行轨迹，协议进入 step 阶段，通过 bisection narrowing 缩小争议区间。对于长度为 n 的 trace，理想情况下定位轮次接近 log2(n)，这避免了线性扫描长执行轨迹。",
            "当争议区间被缩小到单步时，协议进入 single-step replay。replay 阶段比较 expected hash 与 claimed hash，判断局部执行是否一致。如果重放失败，resolve 阶段可以将交易标记为 SLASHED；如果重放通过，则挑战不成立或进入相应结算逻辑。",
            "recover 是本文强调的关键机制。在普通 challenge 中，如果 session 超时，流程可能停滞在中间状态，导致系统只能检测异常而无法完成仲裁。Recoverable challenge 引入 watchdog/operator 恢复路径，使超时 session 能够被恢复并继续 step、replay 和 resolve。该机制的核心价值是 liveness：它使 challenge 能够从“异常被发现”推进到“异常被处理完”。",
        ],
    ),
    (
        "6.1 Correctness and Property Discussion",
        [
            "性质一，DA proof 绑定性质。compact commit 中的 da_merkle_root 与 DA entry 中的字段级 Merkle proof 绑定，因此 prompt、response 或 attestation 字段被替换后，会导致 payload hash、leaf hash 或 Merkle proof 与链上 root 不一致。该性质不依赖真实 DA 网络，只说明当前 verifiable DA 结构能检测字段篡改。",
            "性质二，recover 的活性性质。recover 操作的目标是恢复超时挑战会话，使其继续进入 step、replay 和 resolve；它不修改 expected trace、claimed trace 或 replay mismatch 的事实。因此，recover 不改变争议判断本身，而是避免 timeout 让争议流程永久停在中间状态。",
            "性质三，二分定位性质。对于长度为 n 的执行 trace，challenge-step 每轮缩小争议区间，理想轮次接近 log2(n)。本文实验中 trace steps 为 4、8、16、32、64、128 时，平均二分轮次分别为 2、3、4、5、6、7，与理论对数趋势一致。",
        ],
    ),
    (
        "7 DA Cost Model",
        [
            "本文的成本模型关注 full-onchain 与 compact commit + DA 在高频场景下的结构性差异。full-onchain 路线将完整 payload 作为链上数据提交，其成本随 payload size 增大近似线性增长。compact commit 路线则只将固定摘要结构提交到链上，将完整 payload 放入 DA 层，并通过 DA root 和 Merkle proof 保留验证能力。",
            "本文比较四类 DA route：full-onchain calldata、compact external DA、compact EIP-4844-like 和 compact modular DA sampling。成本分析不是单一变量问题。payload size 决定 full payload 上链成本增长速度；batch size 决定固定提交开销能否被摊薄；DA route 决定 payload 放在不同数据层时的边际成本。",
            "为便于说明，本文采用如下抽象成本表达：C_total = C_commit + C_DA + p_challenge * C_challenge。其中 C_commit 表示链上 compact commit 或 full payload 提交成本，C_DA 表示不同 DA route 下的数据成本，C_challenge 表示发生争议时的挑战、恢复、重放和结算成本，p_challenge 表示争议发生概率。批处理后单笔摊销成本为 C_amortized = C_total / batch_size。当前实验主要比较不同 route 在相同参数下的结构性趋势，而不估计真实业务中的 p_challenge。",
            "该表达不用于给出主网精确费用，而用于比较不同提交策略在相同参数下的结构性趋势。当前 Python prototype 中的成本实验是配置化估算模型，本地 EVM gasUsed 则用于补充链上状态推进基线，两者在 Evaluation 中需要明确区分。",
        ],
    ),
    (
        "8 Evaluation",
        [
            "本文实验围绕三个理论命题展开：第一，evidence-carrying compact commit 能在保留验证入口的同时降低链上提交规模；第二，recoverable challenge 能改善异常场景下的仲裁活性；第三，Python prototype、本地 EVM 和 Sepolia 部署可以形成阶段性验证链路。",
            "实验分为三组。第一组为 Python 成本估算，覆盖 payload size 128、512、2048、8192，prompt length 64、256，batch size 1、10、100、1000，以及四类 DA route，共生成 3200 条记录。第二组为挑战流程实验，覆盖 trace steps 4、8、16、32、64、128，共生成 600 条记录。第三组为失败场景实验，覆盖 normal、attestation_invalid、challenge_timeout_no_recover、challenge_timeout_recover、da_proof_invalid、da_unavailable、response_tampered 和 trace_length_mismatch，共生成 800 条记录。",
            "成本实验覆盖 payload size、prompt length、batch size 和 DA route，共生成 3200 条记录。实验结果显示，compact commit 的字节规模在当前实现中约为 406 bytes，基本不随 payload size 增长；相对地，full payload bytes 随 payload size 增大近似线性上升。这一结果支撑了本文关于 evidence-carrying compact commit 的第一项判断：轻量提交的链上成本主要由固定摘要结构决定，而不是与完整业务负载等比例绑定。",
            "在 payload 较小时，compact commit 的优势已经可见；当 payload 增大到 8192 级别时，full-onchain 成本显著上升，而 compact + DA 路线保持较低增长。在当前参数设置下，payload=8192 时 modular DA profile 相对 full-onchain 的最佳降本比例约为 96.33%。这一数字应被理解为给定模型和 profile 下的趋势结果，而不是主网真实费用结论。",
            "实验还显示，batch size 增大后，单笔 amortized gas 明显下降。例如，在相同 payload 设置下，batch size 从 1 增加到 10、100、1000 时，单笔成本近似按比例摊薄。这一趋势说明 compact commit + DA 更适合高频 DApp 的批量交互场景：高频请求本身为批处理提供了自然条件，而批处理又能降低每笔请求承担的链上固定成本。",
            "挑战实验覆盖 trace steps 4、8、16、32、64、128，共生成 600 条记录。实验显示，二分定位平均轮次分别为 2、3、4、5、6、7，与 log2(trace steps) 基本一致。这说明当前 challenge-step 实现的是对数级争议定位，而不是线性扫描。更重要的是 timeout 对照实验。challenge_timeout_no_recover 场景下，系统能够检测到异常，但 challenge success 为 0%，slashed 也为 0%。相对地，在 challenge_timeout_recover 场景下，challenge success 和 slashed 均达到 100%。该结果直接支撑 recoverable challenge 的 liveness 命题。",
            "失败场景结果需要分层解释。attestation_invalid、da_unavailable 和 da_proof_invalid 均能被稳定检测，但当前原型将它们主要作为可拒绝或可报告故障，并未全部纳入统一 slashing 闭环。response_tampered、trace_length_mismatch 和 challenge_timeout_recover 则能够进入 challenge、replay 与 resolve，并最终形成 slashing。该分类避免把“可检测”误写成“所有故障均可完整惩罚”。",
            "为避免只依赖 Python 成本估算，本文将关键链上路径映射为 Solidity 原型，并在 Hardhat 本地测试链中测量 gasUsed。平均结果包括：submit rollup 约 299033.67 gas，challenge open 约 283769 gas，challenge recover 约 156891 gas，challenge replay 约 135632 gas，challenge resolve 约 113402 gas，finalize 约 33477 gas。这些结果说明，提交、挑战、恢复、重放和结算等关键路径已经具有链上状态推进对应物。",
            "当前 Solidity 原型已部署到 Sepolia，chain id 为 11155111，部署区块为 10825904。部署合约包括 MockTEEVerifier、MockDARegistry 和 HybridTEERollup。其中 HybridTEERollup 地址为 0xB9B72f10bB8aBC1ed090f1B0443Fe57189c497Fd，MockTEEVerifier 地址为 0xC42CB0Cf0D112Bd59E0f212F2DB2002ca50a8b6A，MockDARegistry 地址为 0xc42C64a7De05bf3ad4f1c75CbD8E0506F7De99Ff。Sepolia 部署说明当前系统已具备公开测试网对应物，而不只是本地模拟。源码验证曾因 block explorer connect timeout 未完成，因此本文只将其作为 deployability evidence，不能解释为主网性能、生产安全性或完整系统成熟度的证明。",
        ],
    ),
    (
        "8.1 Experimental Setup and Reproducibility",
        [
            "Python 原型位于 project_code/tee_rollup_demo，核心逻辑包括 simulated attestation、JSON ledger、DA storage、Merkle proof、challenge 状态机和成本估算。正式实验脚本为 project_code/scripts/run_paper_experiments.py，生成 cost_grid.csv、challenge_grid.csv、failure_scenarios.csv、summary.json 和 figures 目录。",
            "本地 EVM 原型位于 project_code/evm，包含 MockTEEVerifier、MockDARegistry、MerkleVerifier 和 HybridTEERollup 合约。Hardhat 测试覆盖 compact commit 提交、可恢复挑战、finalize、无效 attestation 拒绝和 DA field witness 验证。gasUsed 数据由 scripts/measure_gas.js 生成。",
            "复现实验可使用以下命令：python -m unittest discover -s tests -v；python scripts/run_paper_experiments.py --clean --samples 100 --payload-sizes 128,512,2048,8192 --batch-sizes 1,10,100,1000 --trace-steps 4,8,16,32,64,128；在 evm 目录下运行 npm.cmd test 和 npm.cmd run measure:gas。Sepolia 部署记录位于 project_code/evm/deployments/sepolia.json。",
        ],
    ),
    (
        "9 Discussion",
        [
            "本文工作仍有明确局限。第一，当前 TEE 为 simulated TEE。原型中的 attestation 使用模拟方式绑定输入、输出、nonce 和执行环境标识，其目的在于验证协议接口和异常路径，而不是证明真实 SGX、TDX 或其他硬件 TEE 的安全性。真实 TEE 部署还需要考虑远程证明格式、硬件漏洞、侧信道攻击和多厂商 TEE 异构信任假设。",
            "第二，当前 DA 为 mock/verifiable DA 原型。系统实现了 DA root、payload hash、字段级 Merkle proof 和可用性状态记录，但并未接入真实生产级 DA 网络。因此，本文结论更适合表述为“compact commit 与 verifiable DA 结构能够形成可验证证据入口”，而不能写成已经实现生产级 DA。",
            "第三，Python prototype 主要用于结构性趋势验证。成本实验中的 DA profile 是配置化模型，适合比较 full-onchain 与 compact+DA 的相对趋势，但不能直接等价为真实主网 gas。本文将其与本地 EVM gas baseline 区分，是为了避免将估算结果过度外推。",
            "第四，本地 EVM gasUsed 不等于主网费用。Hardhat 本地测试链提供了链上状态推进的真实 gasUsed 基线，但真实公网费用还取决于网络环境、数据定价和交易拥堵状态。第五，Sepolia 部署仅代表阶段性测试网落地。它证明合约可以部署到公开测试网，并生成地址和 ABI 记录，但不代表系统已经具备生产可用性。",
            "第六，当前 fault taxonomy 尚未全部形成统一 slashing 闭环。例如 DA unavailable、DA proof invalid 和 attestation invalid 在当前实验中更多体现为可检测故障，而 response_tampered 和 trace_length_mismatch 已形成检测、challenge 与 slashing 的闭环。后续工作需要进一步统一不同故障类型的仲裁策略，并区分可直接拒绝、需恢复后仲裁和需重放证明的故障类别。",
        ],
    ),
    (
        "10 Related Work",
        [
            "TEEROLLUP 提出了利用异构 TEE 降低 Rollup 验证成本并缩短提现延迟的系统设计，其核心包括 TEE 委员会、链上状态合约、挑战机制和数据可用性惩罚。本文受到 Hybrid TEE-Rollup 思路启发，但并不重做完整 TEEROLLUP 系统，而是聚焦异常场景下 challenge 的恢复性，以及 compact commit 与 verifiable DA 的成本和证据结构。",
            "OTR 和 Optimistic TEE-Rollups 关注如何将 TEE 与 optimistic verification 结合，尤其适用于链上生成式 AI 或模型推理结果验证。这类工作表明 TEE 可以作为快速路径，而 optimistic challenge 可作为异常路径。本文与其相近之处在于同样关注 TEE 与 challenge 的结合，差异在于本文强调 recoverable challenge 的 liveness，以及 DA payload 与 compact commit 之间的证据绑定。",
            "opML 等工作研究如何通过 optimistic fraud proof 支持机器学习或大模型计算，通常包含二分定位、单步仲裁和链上虚拟机验证。本文借鉴交互式争议定位思想，但没有实现完整 FPVM 或真实 ML 执行证明，而是将 bisection replay 用于 Hybrid TEE-Rollup 原型中的异常仲裁。",
            "Dynamic Fraud Proof 关注如何缩短无争议场景下的最终性，并在发现争议时动态延迟结算。该方向强调 challenge window 和验证者参与机制的动态调整。本文不提出完整动态最终性协议，而是在更小范围内研究 timeout 后 challenge session 如何恢复推进。",
            "LazyLedger 和 Light Clients for Lazy Blockchains 等工作强调将数据可用性从执行验证中解耦，并通过数据可用性抽样支持模块化扩容。本文没有实现真实 DAS 网络，而是使用 mock DA 和 DA profile 建模 compact commit + verifiable DA 的成本与证据结构。本文的重点不是替代 DA 层，而是说明 Hybrid TEE-Rollup 中链上轻量提交必须与 DA 证明建立可验证联系。",
        ],
    ),
    (
        "11 Conclusion",
        [
            "本文面向高频 DApp 场景，在已有 Hybrid TEE-Rollup 思路下研究可验证低成本交互机制。本文没有提出一个全新的 Rollup 体系，而是聚焦 recoverable challenge、evidence-carrying compact commit 和 DA-aware cost amortization 三个具体问题。",
            "通过 Python prototype，本文验证了提交、DA proof、挑战、恢复、重放和结算的基本闭环；通过实验，本文观察到 compact commit 的固定成本特征、payload 增大后的 DA 路线收益、batch size 对 amortized gas 的摊薄作用、recover 对 challenge liveness 的改善，以及 bisection 轮次与理论对数复杂度的一致性；通过本地 EVM 和 Sepolia 部署，本文进一步给出链上关键路径基线和公开测试网对应物。",
            "当前工作仍是阶段性研究原型，不能被解释为真实 TEE 安全证明、生产级 DA 系统或主网性能结论。后续工作应继续推进真实 TEE attestation、真实 DA 网络接入、更完整的链上交易样本、源码验证、形式化安全分析和更统一的 fault arbitration 策略。总体而言，本文为 Hybrid TEE-Rollup 在高频 DApp 中的可验证低成本交互提供了一个克制但可继续扩展的研究基础。",
        ],
    ),
]


TABLES = {
    "contrib": {
        "caption": "表 1  本文机制推进点与研究作用",
        "headers": ["机制推进", "核心作用", "论文定位"],
        "rows": [
            ["Recoverable Challenge", "在 timeout 或中断后恢复争议流程", "提升 challenge liveness"],
            ["Evidence-Carrying Compact Commit", "绑定 state/output/proof/DA root", "形成最小证据入口"],
            ["Fault Taxonomy", "区分 attestation、DA、timeout、replay 和 trace 故障", "明确异常责任边界"],
            ["DA-aware Cost Amortization", "联合 payload、batch 和 DA route 分析成本", "支撑高频 DApp 成本论证"],
        ],
    },
    "gas": {
        "caption": "表 4  本地 EVM 关键路径平均 gasUsed",
        "headers": ["操作", "平均 gasUsed", "解释"],
        "rows": [
            ["submit rollup", "299033.67", "compact commit 提交路径"],
            ["challenge open", "283769.00", "打开争议并加载证据"],
            ["challenge recover", "156891.00", "恢复超时挑战会话"],
            ["challenge replay", "135632.00", "单步重放仲裁"],
            ["challenge resolve", "113402.00", "结算争议结果"],
            ["finalize", "33477.00", "挑战期后最终确认"],
        ],
    },
    "setup": {
        "caption": "表 2  实验设置与输出材料",
        "headers": ["实验模块", "参数设置", "输出材料"],
        "rows": [
            ["成本估算", "payload=128/512/2048/8192；batch=1/10/100/1000；samples=100", "cost_grid.csv；summary.json；成本图"],
            ["挑战流程", "trace steps=4/8/16/32/64/128；timeout=2；max rounds=16", "challenge_grid.csv；challenge_rounds.png"],
            ["失败场景", "8 类场景，每类 100 个样本", "failure_scenarios.csv；failure_detection.png"],
            ["本地 EVM", "Hardhat 本地链；3 种 payload", "measured_gas.csv；measured_gas_report.md"],
            ["Sepolia", "chain id=11155111；部署区块=10825904", "sepolia.json；ABI 地址索引"],
        ],
    },
    "failure": {
        "caption": "表 3  失败场景检测与仲裁结果",
        "headers": ["场景类别", "代表场景", "当前原型结果"],
        "rows": [
            ["正常路径", "normal", "无异常检测；不触发 slashing"],
            ["可检测但未统一 slashing", "attestation_invalid；da_unavailable；da_proof_invalid", "检测率 100%，但当前主要作为拒绝或报告故障"],
            ["可完成 challenge + slashing", "response_tampered；trace_length_mismatch", "检测率、challenge success、slashed 均为 100%"],
            ["timeout 对照", "challenge_timeout_no_recover；challenge_timeout_recover", "无 recover 时不完成仲裁；recover 后 challenge success 与 slashed 为 100%"],
        ],
    },
    "sepolia": {
        "caption": "表 5  Sepolia 部署信息",
        "headers": ["项目", "数值"],
        "rows": [
            ["Network / Chain ID", "Sepolia / 11155111"],
            ["部署区块", "10825904"],
            ["MockTEEVerifier", "0xC42CB0Cf0D112Bd59E0f212F2DB2002ca50a8b6A"],
            ["MockDARegistry", "0xc42C64a7De05bf3ad4f1c75CbD8E0506F7De99Ff"],
            ["HybridTEERollup", "0xB9B72f10bB8aBC1ed090f1B0443Fe57189c497Fd"],
        ],
    },
}


FIGURES = [
    ("project_code/paper_outputs/figures/cost_payload_full_vs_compact.png", "图 1  full payload 与 compact commit 的字节规模对比"),
    ("project_code/paper_outputs/figures/da_profile_amortized_gas.png", "图 2  不同 DA profile 下的摊销 gas 趋势"),
    ("project_code/paper_outputs/figures/challenge_rounds.png", "图 3  二分挑战轮次与 trace steps 的关系"),
    ("project_code/paper_outputs/figures/recovery_success.png", "图 4  timeout recover 对挑战完成率的影响"),
    ("project_code/paper_outputs/figures/failure_detection.png", "图 5  失败场景检测与仲裁结果"),
]


REFERENCES = [
    "[1] WEN X, FENG Q, LYU H, et al. TEEROLLUP: Efficient Rollup Design Using Heterogeneous TEE[EB/OL]. arXiv:2409.14647, 2024.",
    "[2] PICCO G, FORTUGNO A. Dynamic Fraud Proof[EB/OL]. arXiv:2502.10321, 2025.",
    "[3] AL-BASSAM M. LazyLedger: A Distributed Data Availability Ledger with Client-Side Smart Contracts[EB/OL]. arXiv:1905.09274, 2019.",
    "[4] TAS E N, TSE D, YANG L, et al. Light Clients for Lazy Blockchains[EB/OL]. arXiv:2203.15968, 2022.",
    "[5] Ethereum Foundation. EIP-4844: Shard Blob Transactions[EB/OL]. 2024.",
    "[6] Ethereum Foundation. Optimistic Rollups[EB/OL]. Ethereum Documentation, 2024.",
    "[7] Buterin V. An Incomplete Guide to Rollups[EB/OL]. 2021.",
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_cm):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=None, bold=None, color=None, font="宋体"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def paragraph(document, text="", style=None, align=None, first_line=True):
    p = document.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.space_before = Pt(0)
    if text:
        run = p.add_run(text)
        set_run_font(run, 10.5)
    return p


def heading(document, text, level=1):
    p = document.add_paragraph(style=f"Heading {min(level, 3)}")
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, 13 if level == 1 else 11.5, bold=True, color=(31, 78, 121))
    return p


def section_level(title):
    marker = title.split()[0]
    return 2 if "." in marker else 1


def add_table(document, table_spec):
    caption = paragraph(document, table_spec["caption"], align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    caption.paragraph_format.keep_with_next = True
    caption.runs[0].bold = True
    table = document.add_table(rows=1, cols=len(table_spec["headers"]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(table_spec["headers"]):
        cell = hdr.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_run_font(r, 9.5, bold=True, color=(255, 255, 255))
        set_cell_shading(cell, "1F4E79")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row_data in table_spec["rows"]:
        row = table.add_row()
        for i, val in enumerate(row_data):
            cell = row.cells[i]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != len(row_data) - 1 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            set_run_font(r, 9)
    column_count = len(table_spec["headers"])
    if column_count == 2:
        widths = [4.8, 10.2]
    elif column_count == 3:
        widths = [4.2, 4.8, 6.0]
    elif column_count == 4:
        widths = [3.2, 4.2, 4.2, 5.0]
    else:
        widths = [15.0 / max(1, column_count)] * column_count
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(i, len(widths) - 1)])
    paragraph(document, "", first_line=False)


def add_figure(document, rel_path, caption):
    path = ROOT / rel_path
    if not path.exists():
        return
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Cm(12.8))
    cap = paragraph(document, caption, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    cap.runs[0].bold = True


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        st = styles[style_name]
        st.font.name = "宋体"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def build_docx():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    r = title.add_run(TITLE)
    set_run_font(r, 16, bold=True, font="黑体")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    r = subtitle.add_run(EN_TITLE)
    set_run_font(r, 11, bold=False, font="Times New Roman")

    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.space_after = Pt(2)
    r = author.add_run("作者：杨帆（占位）")
    set_run_font(r, 10.5)

    aff = doc.add_paragraph()
    aff.alignment = WD_ALIGN_PARAGRAPH.CENTER
    aff.paragraph_format.space_after = Pt(12)
    r = aff.add_run("单位：待补充；邮箱：待补充")
    set_run_font(r, 10)

    abstract_title = paragraph(doc, "摘要", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    abstract_title.runs[0].bold = True
    paragraph(doc, ABSTRACT)
    kw = paragraph(doc, f"关键词：{KEYWORDS}", first_line=False)
    kw.runs[0].bold = True

    # First page intentionally keeps front matter together; main text starts after it.
    doc.add_section(WD_SECTION.CONTINUOUS)

    inserted_contrib = False
    inserted_system_fig = False
    inserted_eval_figs = False

    for sec_title, paras in SECTIONS:
        level = section_level(sec_title)
        heading(doc, sec_title, level=level)
        for text in paras:
            paragraph(doc, text)
        if sec_title == "1 Introduction" and not inserted_contrib:
            add_table(doc, TABLES["contrib"])
            inserted_contrib = True
        if sec_title == "5.3 验证层" and not inserted_system_fig:
            # Keep the design chapter text-centered; figures here are experiment figures, not architecture.
            inserted_system_fig = True
        if sec_title == "8 Evaluation" and not inserted_eval_figs:
            add_table(doc, TABLES["setup"])
            add_figure(doc, FIGURES[0][0], FIGURES[0][1])
            add_figure(doc, FIGURES[1][0], FIGURES[1][1])
            add_figure(doc, FIGURES[2][0], FIGURES[2][1])
            add_figure(doc, FIGURES[3][0], FIGURES[3][1])
            add_figure(doc, FIGURES[4][0], FIGURES[4][1])
            add_table(doc, TABLES["failure"])
            add_table(doc, TABLES["gas"])
            doc.add_page_break()
            add_table(doc, TABLES["sepolia"])
            inserted_eval_figs = True

    heading(doc, "参考文献", level=1)
    for ref in REFERENCES:
        p = paragraph(doc, ref, first_line=False)
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.74)

    doc.save(DOCX_PATH)


def markdown_table(table_spec):
    lines = [f"\n**{table_spec['caption']}**\n"]
    headers = table_spec["headers"]
    lines.append("\n| " + " | ".join(headers) + " |\n")
    lines.append("| " + " | ".join(["---"] * len(headers)) + " |\n")
    for row in table_spec["rows"]:
        lines.append("| " + " | ".join(row) + " |\n")
    return "".join(lines)


def markdown_figure(rel_path, caption):
    figure_name = Path(rel_path).name
    return f"\n![{caption}](../figures/{figure_name})\n\n**{caption}**\n"


def strip_inline_markdown(text):
    return text.replace("**", "").replace("`", "")


def parse_markdown_table(lines, start):
    caption = None
    j = start - 1
    while j >= 0 and not lines[j].strip():
        j -= 1
    if j >= 0 and lines[j].strip().startswith("**表"):
        caption = strip_inline_markdown(lines[j].strip())
    headers = [strip_inline_markdown(item.strip()) for item in lines[start].strip().strip("|").split("|")]
    rows = []
    i = start + 2
    while i < len(lines) and lines[i].strip().startswith("|"):
        row = [strip_inline_markdown(item.strip()) for item in lines[i].strip().strip("|").split("|")]
        rows.append(row)
        i += 1
    return {"caption": caption or "表", "headers": headers, "rows": rows}, i


def diagram_font(size, bold=False):
    candidates = [
        "C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simhei.ttf",
        "C:/Windows/Fonts/simsun.ttc",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            continue
    return ImageFont.load_default()


def draw_multiline_center(draw, box, text, font, fill="#1f2937", line_gap=5):
    x1, y1, x2, y2 = box
    lines = text.split("\n")
    heights = []
    widths = []
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        widths.append(bbox[2] - bbox[0])
        heights.append(bbox[3] - bbox[1])
    total_h = sum(heights) + line_gap * (len(lines) - 1)
    y = y1 + ((y2 - y1) - total_h) / 2
    for line, w, h in zip(lines, widths, heights):
        draw.text((x1 + ((x2 - x1) - w) / 2, y), line, font=font, fill=fill)
        y += h + line_gap


def draw_round_box(draw, box, label, fill, outline="#36506c", radius=18):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=2)
    draw_multiline_center(draw, box, label, diagram_font(22), fill="#111827")


def draw_arrow(draw, start, end, fill="#374151", width=3):
    draw.line([start, end], fill=fill, width=width)
    sx, sy = start
    ex, ey = end
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex >= sx else -1
        head = [(ex, ey), (ex - direction * 16, ey - 8), (ex - direction * 16, ey + 8)]
    else:
        direction = 1 if ey >= sy else -1
        head = [(ex, ey), (ex - 8, ey - direction * 16), (ex + 8, ey - direction * 16)]
    draw.polygon(head, fill=fill)


def generate_paper_diagrams():
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)

    img = Image.new("RGB", (1600, 900), "#f8fafc")
    draw = ImageDraw.Draw(img)
    title_font = diagram_font(34, bold=True)
    draw.text((60, 38), "Hybrid TEE-Rollup 原型系统架构", font=title_font, fill="#111827")

    boxes = {
        "user": (70, 155, 350, 285),
        "executor": (490, 120, 830, 300),
        "da": (490, 410, 830, 590),
        "chain": (990, 120, 1430, 300),
        "challenge": (990, 410, 1430, 650),
    }
    draw_round_box(draw, boxes["user"], "高频 DApp 请求\nprompt / input", "#dbeafe")
    draw_round_box(draw, boxes["executor"], "链下执行层\nSimulated TEE\nresponse + attestation", "#dcfce7")
    draw_round_box(draw, boxes["da"], "数据可用性层\nmock/verifiable DA\npayload + Merkle proof", "#fef3c7")
    draw_round_box(draw, boxes["chain"], "链上提交层\nHybridTEERollup\ncompact commit", "#ede9fe")
    draw_round_box(draw, boxes["challenge"], "验证与仲裁层\nopen / respond / step\nrecover / replay / resolve", "#fee2e2")

    draw_arrow(draw, (350, 220), (490, 220))
    draw_arrow(draw, (660, 300), (660, 410))
    draw_arrow(draw, (830, 210), (990, 210))
    draw_arrow(draw, (830, 505), (990, 520))
    draw_arrow(draw, (1210, 300), (1210, 410))
    draw_arrow(draw, (990, 585), (830, 560))

    note_font = diagram_font(20)
    draw.text((90, 750), "证据边界：Python prototype / JSON ledger / simulated TEE 用于机制验证；DA profile 用于配置化成本趋势；Hardhat 与 Sepolia 仅支撑链上对应物和可部署性证据。", font=note_font, fill="#374151")
    img.save(ARCH_FIGURE_PATH)

    img = Image.new("RGB", (1700, 780), "#f8fafc")
    draw = ImageDraw.Draw(img)
    draw.text((60, 38), "Recoverable Challenge 协议流程", font=title_font, fill="#111827")
    flow = [
        ("PENDING\n提交等待", 70, 170),
        ("OPEN\n挑战开启", 300, 170),
        ("RESPONDED\n回应争议", 530, 170),
        ("NARROWING\n二分定位", 760, 170),
        ("READY\n单步证据", 990, 170),
        ("REPLAYED\n单步重放", 1220, 170),
        ("RESOLVED\n结算结果", 1450, 170),
    ]
    for label, x, y in flow:
        draw_round_box(draw, (x, y, x + 190, y + 110), label, "#e0f2fe")
    for (_, x1, y1), (_, x2, y2) in zip(flow, flow[1:]):
        draw_arrow(draw, (x1 + 190, y1 + 55), (x2, y2 + 55))

    draw_round_box(draw, (760, 450, 990, 585), "TIMEOUT\n会话停滞", "#ffedd5")
    draw_round_box(draw, (1050, 450, 1280, 585), "RECOVERED\n恢复推进", "#dcfce7")
    draw_arrow(draw, (855, 280), (855, 450))
    draw_arrow(draw, (990, 520), (1050, 520))
    draw_arrow(draw, (1165, 450), (1085, 280))

    draw.text((80, 660), "recover 只恢复挑战活性，不改写 expected trace、claimed trace 或 replay mismatch；最终是否 slashing 仍由 evidence 与 replay 结果决定。", font=note_font, fill="#374151")
    img.save(PROTOCOL_FIGURE_PATH)


def resolve_markdown_image(path_text):
    raw = path_text.strip()
    candidate = (MD_PATH.parent / raw).resolve()
    if candidate.exists():
        return candidate
    return (ROOT / raw).resolve()


def add_markdown_figure(document, image_path, caption):
    path = Path(image_path)
    if not path.exists():
        return
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Cm(12.8))
    cap = paragraph(document, caption, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    cap.runs[0].bold = True


def build_docx_from_markdown():
    if not MD_PATH.exists():
        build_markdown()

    raw_lines = MD_PATH.read_text(encoding="utf-8").splitlines()
    try:
        body_start = next(i for i, line in enumerate(raw_lines) if line.startswith("## 1 "))
    except StopIteration:
        body_start = 0
    lines = raw_lines[body_start:]

    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    r = title.add_run(TITLE)
    set_run_font(r, 16, bold=True, font="黑体")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    r = subtitle.add_run(EN_TITLE)
    set_run_font(r, 11, bold=False, font="Times New Roman")

    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.space_after = Pt(2)
    r = author.add_run("作者：杨帆（占位）")
    set_run_font(r, 10.5)

    aff = doc.add_paragraph()
    aff.alignment = WD_ALIGN_PARAGRAPH.CENTER
    aff.paragraph_format.space_after = Pt(12)
    r = aff.add_run("单位：待补充；邮箱：待补充")
    set_run_font(r, 10)

    abstract_title = paragraph(doc, "摘要", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    abstract_title.runs[0].bold = True
    paragraph(doc, ABSTRACT)
    kw = paragraph(doc, f"关键词：{KEYWORDS}", first_line=False)
    kw.runs[0].bold = True
    doc.add_section(WD_SECTION.CONTINUOUS)

    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("## "):
            heading(doc, strip_inline_markdown(line[3:].strip()), level=1)
            i += 1
            continue
        if line.startswith("### "):
            heading(doc, strip_inline_markdown(line[4:].strip()), level=2)
            i += 1
            continue
        if line.startswith("!["):
            close = line.find("]")
            open_paren = line.find("(", close)
            close_paren = line.rfind(")")
            caption = line[2:close] if close != -1 else "图"
            image_path = line[open_paren + 1 : close_paren] if open_paren != -1 and close_paren != -1 else ""
            add_markdown_figure(doc, resolve_markdown_image(image_path), strip_inline_markdown(caption))
            if i + 1 < len(lines) and strip_inline_markdown(lines[i + 1].strip()) == strip_inline_markdown(f"**{caption}**"):
                i += 2
            else:
                i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            table_spec, i = parse_markdown_table(lines, i)
            if table_spec["caption"].startswith("表 4  Sepolia"):
                doc.add_page_break()
            add_table(doc, table_spec)
            continue
        if line.startswith("**表") or line.startswith("**图"):
            i += 1
            continue
        paragraph(doc, strip_inline_markdown(line))
        i += 1

    try:
        doc.save(DOCX_PATH)
        return DOCX_PATH
    except PermissionError:
        doc.save(FALLBACK_DOCX_PATH)
        return FALLBACK_DOCX_PATH


def build_markdown():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    parts = [f"# {TITLE}\n", f"**英文题目：** {EN_TITLE}\n", "**作者：** 杨帆（占位）\n", "**单位：** 待补充\n"]
    parts.append(f"\n## 摘要\n\n{ABSTRACT}\n")
    parts.append(f"\n**关键词：** {KEYWORDS}\n")
    for title, paras in SECTIONS:
        hashes = "###" if section_level(title) == 2 else "##"
        parts.append(f"\n{hashes} {title}\n")
        for para in paras:
            parts.append(f"\n{para}\n")
        if title == "1 Introduction":
            parts.append(markdown_table(TABLES["contrib"]))
        if title == "8 Evaluation":
            parts.append(markdown_table(TABLES["setup"]))
            for rel_path, caption in FIGURES:
                parts.append(markdown_figure(rel_path, caption))
            parts.append(markdown_table(TABLES["failure"]))
            parts.append(markdown_table(TABLES["gas"]))
            parts.append(markdown_table(TABLES["sepolia"]))
    parts.append("\n## 参考文献\n")
    for ref in REFERENCES:
        parts.append(f"\n{ref}\n")
    MD_PATH.write_text("".join(parts), encoding="utf-8")


if __name__ == "__main__":
    generate_paper_diagrams()
    if not MD_PATH.exists():
        build_markdown()
    output_docx = build_docx_from_markdown()
    print(output_docx)
    print(MD_PATH)
