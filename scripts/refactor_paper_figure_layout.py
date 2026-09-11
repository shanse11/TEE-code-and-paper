# -*- coding: utf-8 -*-
"""
全文图表排版结构重构（Layout Refactoring）
- 重排图1-4至语义锚点；图5-9在保留图题处重建内嵌图
- 强制上下型内嵌、禁止环绕；宽度/最大高度约束
"""

from __future__ import annotations

import shutil
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
DRAFT = ROOT / "project_code" / "paper_outputs" / "draft"
DOCX_PATH = DRAFT / "Hybrid_TEE_Rollup_中文论文初稿.docx"
BACKUP_DIR = DRAFT / "backups"
ASSETS = DRAFT / "readability_assets_20260528"

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

STRUCTURAL_CAPTIONS = {
    "Figure 1 Threat Model of RCP",
    "图 1 Threat Model of RCP",
    "图 2 可恢复挑战驱动的 Hybrid TEE-Rollup 原型架构",
    "图 3 RCP 协议交互时序",
    "图 4 可恢复挑战协议状态机",
}

EXPERIMENT_CAPTIONS = [
    "图 5 恢复对超时场景挑战完成率的影响",
    "图 6 二分挑战轮次与 trace steps 的关系",
    "图 7 full payload 与 compact commit 的字节规模对比",
    "图 8 不同 DA profile 下的摊销 gas 趋势",
    "图 9 失败场景检测与仲裁结果",
]

EXPERIMENT_ASSETS = [
    "fig05_recovery_success.png",
    "fig06_challenge_rounds.png",
    "fig07_cost_payload.png",
    "fig08_da_gas.png",
    "fig09_failure_detection.png",
]

TABLE_PAGE_BREAKS_BEFORE = ["表 5 RCP 状态转移"]


@dataclass
class FigureSpec:
    caption: str
    asset: str
    width_ratio: float
    max_height_ratio: float
    anchor_contains: str
    page_break_before: bool = False


STRUCTURAL_FIGURES = [
    FigureSpec(
        "Figure 1 Threat Model of RCP",
        "fig01_threat_model.png",
        0.55,
        0.35,
        "攻击者可控制提交者、排序者或被挑战方",
        page_break_before=True,
    ),
    FigureSpec(
        "图 2 可恢复挑战驱动的 Hybrid TEE-Rollup 原型架构",
        "fig02_architecture.png",
        0.70,
        0.48,
        "原型采用执行层、数据层和验证层三层结构",
    ),
    FigureSpec(
        "图 4 可恢复挑战协议状态机",
        "fig04_state_machine.png",
        0.65,
        0.38,
        "可恢复挑战协议的目标，是将 challenge 从一次性欺诈检测流程扩展",
        page_break_before=True,
    ),
]

FIG3_SPEC = FigureSpec(
    "图 3 RCP 协议交互时序",
    "fig03_sequence.png",
    0.70,
    0.45,
    "正常路径中，提交者只向链上提交",
    page_break_before=True,
)


def page_sizes(doc: Document):
    s = doc.sections[0]
    cw = (s.page_width - s.left_margin - s.right_margin) / 914400
    ph = s.page_height / 914400
    return cw, ph


def para_has_image(paragraph) -> bool:
    return "pic:pic" in paragraph._element.xml


def is_page_break_paragraph(paragraph) -> bool:
    return not paragraph.text.strip() and paragraph._element.find(f".//{{{NS_W}}}br") is not None


def is_blank_paragraph(paragraph) -> bool:
    return not paragraph.text.strip() and not para_has_image(paragraph) and not is_page_break_paragraph(paragraph)


def strip_wrap(paragraph):
    p_pr = paragraph._element.get_or_add_pPr()
    for child in list(p_pr):
        if child.tag.endswith("framePr"):
            p_pr.remove(child)


def find_paragraph_contains(doc: Document, needle: str) -> Paragraph | None:
    for p in doc.paragraphs:
        if needle in p.text:
            return p
    return None


def find_caption(doc: Document, caption: str) -> Paragraph | None:
    for p in doc.paragraphs:
        t = p.text.strip()
        if t == caption or t.startswith(caption):
            return p
    return None


def insert_after(anchor: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    anchor._element.addnext(new_p)
    return Paragraph(new_p, anchor._parent)


def insert_before(anchor: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    anchor._element.addprevious(new_p)
    return Paragraph(new_p, anchor._parent)


def add_page_break_after(anchor: Paragraph) -> Paragraph:
    p = insert_after(anchor)
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


def add_page_break_before(anchor: Paragraph) -> Paragraph:
    p = insert_before(anchor)
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


def compute_picture_size(asset: Path, cw: float, ph: float, width_ratio: float, max_height_ratio: float):
    from PIL import Image

    with Image.open(asset) as im:
        px_w, px_h = im.size
    aspect = px_h / px_w if px_w else 1.0
    width_in = cw * width_ratio
    height_in = width_in * aspect
    max_h = ph * max_height_ratio
    if height_in > max_h:
        height_in = max_h
        width_in = height_in / aspect
    return width_in, height_in


def fill_image_paragraph(para: Paragraph, asset: Path, width_in: float, height_in: float):
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(12)
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.keep_with_next = True
    strip_wrap(para)
    para.add_run().add_picture(str(asset), width=Inches(width_in), height=Inches(height_in))


def clear_runs(paragraph):
    el = paragraph._element
    for child in list(el):
        if child.tag.endswith("r"):
            el.remove(child)


def fill_caption_paragraph(para: Paragraph, caption: str):
    clear_runs(para)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(12)
    para.paragraph_format.keep_with_next = False
    strip_wrap(para)
    r = para.add_run(caption)
    r.bold = True
    r.font.size = Pt(10.5)


def insert_figure_after_anchor(doc: Document, spec: FigureSpec, cw: float, ph: float):
    anchor = find_paragraph_contains(doc, spec.anchor_contains)
    if anchor is None:
        raise RuntimeError(f"Anchor not found: {spec.anchor_contains}")
    asset = ASSETS / spec.asset
    w_in, h_in = compute_picture_size(asset, cw, ph, spec.width_ratio, spec.max_height_ratio)
    insert_point = anchor
    if spec.page_break_before:
        insert_point = add_page_break_after(anchor)
    img_para = insert_after(insert_point)
    fill_image_paragraph(img_para, asset, w_in, h_in)
    cap_para = insert_after(img_para)
    fill_caption_paragraph(cap_para, spec.caption)


def insert_figure_before_anchor(doc: Document, spec: FigureSpec, cw: float, ph: float):
    anchor = find_paragraph_contains(doc, spec.anchor_contains)
    if anchor is None:
        raise RuntimeError(f"Anchor not found: {spec.anchor_contains}")
    asset = ASSETS / spec.asset
    w_in, h_in = compute_picture_size(asset, cw, ph, spec.width_ratio, spec.max_height_ratio)
    if spec.page_break_before:
        add_page_break_before(anchor)
    cap_para = insert_before(anchor)
    fill_caption_paragraph(cap_para, spec.caption)
    img_para = insert_before(cap_para)
    fill_image_paragraph(img_para, asset, w_in, h_in)


def insert_figure_before_caption(doc: Document, caption: str, asset_name: str, cw: float, ph: float, width_ratio: float, max_h_ratio: float):
    cap = find_caption(doc, caption)
    if cap is None:
        raise RuntimeError(f"Caption not found: {caption}")
    prev = cap._element.getprevious()
    if prev is not None and "pic:pic" in prev.xml:
        fill_caption_paragraph(cap, caption)
        return
    asset = ASSETS / asset_name
    w_in, h_in = compute_picture_size(asset, cw, ph, width_ratio, max_h_ratio)
    img_para = insert_before(cap)
    fill_image_paragraph(img_para, asset, w_in, h_in)
    fill_caption_paragraph(cap, caption)


def remove_all_image_paragraphs(doc: Document):
    for p in list(doc.paragraphs):
        if para_has_image(p):
            p._element.getparent().remove(p._element)


def remove_structural_captions(doc: Document):
    for p in list(doc.paragraphs):
        if p.text.strip() in STRUCTURAL_CAPTIONS:
            p._element.getparent().remove(p._element)


def cleanup_blanks(doc: Document):
    changed = True
    while changed:
        changed = False
        prev_blank = False
        for p in list(doc.paragraphs):
            if is_blank_paragraph(p):
                if prev_blank:
                    p._element.getparent().remove(p._element)
                    changed = True
                    break
                prev_blank = True
            else:
                prev_blank = False


def cleanup_consecutive_page_breaks(doc: Document):
    """合并连续分页符段，减少空白页。"""
    changed = True
    while changed:
        changed = False
        prev_break = False
        for p in list(doc.paragraphs):
            if is_page_break_paragraph(p):
                if prev_break:
                    p._element.getparent().remove(p._element)
                    changed = True
                    break
                prev_break = True
            else:
                prev_break = False


def prioritize_tables(doc: Document):
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                tr_pr.append(OxmlElement("w:cantSplit"))
        for ri, row in enumerate(table.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if run.font.size is None or run.font.size < Pt(9):
                            run.font.size = Pt(9)
                        if ri == 0:
                            run.bold = True


def page_break_before_table_captions(doc: Document):
    for text in TABLE_PAGE_BREAKS_BEFORE:
        cap = find_caption(doc, text)
        if cap is None:
            continue
        prev = cap._element.getprevious()
        if prev is not None:
            prev_p = Paragraph(prev, cap._parent)
            if is_page_break_paragraph(prev_p):
                continue
        add_page_break_before(cap)


def refactor():
    backup = BACKUP_DIR / f"Hybrid_TEE_Rollup_中文论文初稿_before_refactor_{datetime.now():%Y%m%d_%H%M%S}.docx"
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(DOCX_PATH, backup)

    doc = Document(str(DOCX_PATH))
    cw, ph = page_sizes(doc)

    remove_all_image_paragraphs(doc)
    remove_structural_captions(doc)
    cleanup_blanks(doc)

    # 图1：3.2 段后，分页后独立展示
    insert_figure_after_anchor(doc, STRUCTURAL_FIGURES[0], cw, ph)
    # 图2：4 系统概览 首段后
    insert_figure_after_anchor(doc, STRUCTURAL_FIGURES[1], cw, ph)
    # 图3：「正常路径」正文之前
    insert_figure_before_anchor(doc, FIG3_SPEC, cw, ph)
    # 图4：5 章引言段之后、5.1 之前
    insert_figure_after_anchor(doc, STRUCTURAL_FIGURES[2], cw, ph)

    # 图5-9：保留图题，仅重建图片
    for caption, asset in zip(EXPERIMENT_CAPTIONS, EXPERIMENT_ASSETS):
        insert_figure_before_caption(doc, caption, asset, cw, ph, 0.57, 0.40)

    cleanup_blanks(doc)
    cleanup_consecutive_page_breaks(doc)
    page_break_before_table_captions(doc)
    prioritize_tables(doc)

    out = DRAFT / "_Hybrid_TEE_Rollup_中文论文初稿_refactored.docx"
    doc.save(str(out))
    try:
        shutil.copy2(out, DOCX_PATH)
        print(f"Updated: {DOCX_PATH}")
    except PermissionError:
        print(f"WARN: could not overwrite {DOCX_PATH} (file may be open).")
        print(f"      Please close Word and use: {out}")
    return backup


def verify(doc_path: Path):
    doc = Document(str(doc_path))
    cw, ph = page_sizes(doc)
    ns_wp = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    ns_a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    all_caps = STRUCTURAL_CAPTIONS | set(EXPERIMENT_CAPTIONS)
    print("--- layout verify ---")
    for i, p in enumerate(doc.paragraphs):
        if para_has_image(p) or p.text.strip() in all_caps:
            t = p.text.strip()[:50]
            info = []
            if para_has_image(p):
                for d in p._element.findall(f".//{{{ns_wp}}}inline"):
                    ext = d.find(f".//{{{ns_a}}}ext")
                    if ext is not None:
                        wi = int(ext.get("cx")) / 914400
                        hi = int(ext.get("cy")) / 914400
                        info.append(f"{wi:.2f}x{hi:.2f}in w={wi/cw*100:.0f}% h={hi/ph*100:.0f}%page")
            anch = len(p._element.findall(f".//{{{ns_wp}}}anchor"))
            print(f"P{i:3d} anchor={anch} {info} | {t}")


def main():
    backup = refactor()
    verify(DRAFT / "_Hybrid_TEE_Rollup_中文论文初稿_refactored.docx")
    print(f"\nBackup: {backup}")


if __name__ == "__main__":
    main()
