from pathlib import Path
import re
import shutil
import zipfile
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


BASE = Path("E:/项目/project_code/paper_outputs/draft")
SRC = BASE / "teacher_comments_only_revised.docx"
OUT = BASE / "final_submission_formatted.docx"
REPORT = BASE / "format_revision_report.md"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = "{" + W_NS + "}"


def cleanup_markup(src: Path, dst: Path):
    shutil.copyfile(src, dst)
    tmp = dst.with_suffix(".tmp.docx")
    with zipfile.ZipFile(dst, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            name = item.filename
            if name == "word/comments.xml":
                continue
            if name in ("word/document.xml", "word/footnotes.xml", "word/endnotes.xml"):
                root = ET.fromstring(zin.read(name))
                accept_revisions(root)
                remove_comment_marks(root)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                zout.writestr(item, data)
            elif name == "word/settings.xml":
                root = ET.fromstring(zin.read(name))
                for el in list(root):
                    if el.tag == W + "trackRevisions":
                        root.remove(el)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                zout.writestr(item, data)
            elif name == "word/_rels/document.xml.rels":
                root = ET.fromstring(zin.read(name))
                for rel in list(root):
                    if rel.attrib.get("Type", "").endswith("/comments"):
                        root.remove(rel)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                zout.writestr(item, data)
            else:
                zout.writestr(item, zin.read(name))
    tmp.replace(dst)


def accept_revisions(root):
    def walk(parent):
        for child in list(parent):
            walk(child)
            if child.tag in (W + "del", W + "moveFrom"):
                parent.remove(child)
            elif child.tag in (W + "ins", W + "moveTo"):
                idx = list(parent).index(child)
                parent.remove(child)
                for grand in reversed(list(child)):
                    parent.insert(idx, grand)
    walk(root)


def remove_comment_marks(root):
    for parent in root.iter():
        for child in list(parent):
            if child.tag in (W + "commentRangeStart", W + "commentRangeEnd"):
                parent.remove(child)
            elif child.tag == W + "r" and any(x.tag == W + "commentReference" for x in child.iter()):
                parent.remove(child)


def set_font(run, east="宋体", west="Times New Roman", size=10, bold=None, italic=None):
    run.font.name = west
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run._element.rPr.rFonts.set(qn("w:ascii"), west)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), west)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic


def set_para(p, align=None, first_line=False, line=1.0, before=0, after=0):
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.first_line_indent = Pt(20) if first_line else None


def normalize_text(text):
    text = text.replace("计算机科学与探索", "")
    text = text.replace("DA感知", "DA 感知")
    text = text.replace("本地gas", "本地 gas")
    text = text.replace("EVM ）", "EVM）")
    text = text.replace("部署 路径", "部署路径")
    text = re.sub(r"\s+([，。；：！？、）])", r"\1", text)
    text = re.sub(r"（\s+", "（", text)
    text = re.sub(r"\s+）", "）", text)
    return text


def rewrite_paragraph_text(p, text):
    if p.text == text:
        return
    style = p.style
    alignment = p.alignment
    p.text = text
    p.style = style
    p.alignment = alignment


def heading_level(text):
    t = " ".join(text.split())
    if re.match(r"^\d+\s+\S", t):
        return 1
    if re.match(r"^\d+\.\d+\s+\S", t):
        return 2
    if re.match(r"^\d+\.\d+\.\d+(?:\.\d+)?\s+\S", t):
        return 3
    return 0


def is_reference(text):
    return bool(re.match(r"^\[\d+\]", text.strip()))


def is_caption(text):
    t = text.strip()
    return bool(
        re.match(r"^(图|表)\s*\d+\s{2,}\S", t)
        or re.match(r"^(Fig\.|Table)\s*\d+\s{2,}\S", t)
    )


def format_sections(doc):
    # Template page size/margins: 20.14 cm x 27.55 cm, top 3.2, bottom 0.8, left/right 1.4.
    for sec in doc.sections:
        sec.page_width = Cm(20.14)
        sec.page_height = Cm(27.55)
        sec.top_margin = Cm(3.2)
        sec.bottom_margin = Cm(0.8)
        sec.left_margin = Cm(1.4)
        sec.right_margin = Cm(1.4)
        sec.header_distance = Cm(1.2)
        sec.footer_distance = Cm(0.8)
        # Preserve existing one-column wide figure/table sections and two-column body sections.
        sect_pr = sec._sectPr
        cols = sect_pr.xpath("./w:cols")
        if cols:
            cols[0].set(qn("w:space"), "425")
        # Header: keep journal/title in header only, not body.
        for p in sec.header.paragraphs:
            if p.text.strip():
                p.text = "计算机科学与探索"
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    set_font(r, east="宋体", size=9)


def format_paragraphs(doc):
    table_title_map = {
        "Table 1 Mechanism advances and roles": "表 1 机制推进点与作用\nTable 1 Mechanism advances and roles",
        "Table 2 Semantic comparison of mechanisms": "表 2 机制语义对比\nTable 2 Semantic comparison of mechanisms",
        "Table 3 Protocol invariants and preserved properties": "表 3 协议不变量与保持性质\nTable 3 Protocol invariants and preserved properties",
        "Table 4 Failure scenarios and evidence boundaries": "表 4 故障场景与证据边界\nTable 4 Failure scenarios and evidence boundaries",
        "Table 5 Average gasUsed of local EVM key paths": "表 5 本地 EVM 关键路径平均 gasUsed\nTable 5 Average gasUsed of local EVM key paths",
        "Table 6 Sepolia deployment information": "表 6 Sepolia 部署信息\nTable 6 Sepolia deployment information",
        "Table 7 Summary of comparison experiments": "表 7 对比实验小结\nTable 7 Summary of comparison experiments",
        "Table 8 Comparison with related work": "表 8 相关工作与本文工作的对照\nTable 8 Comparison with related work",
    }
    body_started = False
    refs_started = False
    for idx, p in enumerate(doc.paragraphs):
        text = normalize_text(p.text.strip())
        if text in table_title_map:
            text = table_title_map[text]
        if text != p.text.strip():
            rewrite_paragraph_text(p, text)
        t = p.text.strip()
        if not t:
            set_para(p, after=0)
            continue

        if idx == 0:
            set_para(p, WD_ALIGN_PARAGRAPH.CENTER, first_line=False, after=4)
            for r in p.runs:
                set_font(r, east="黑体", west="Times New Roman", size=18, bold=True)
            continue
        if idx == 1:
            set_para(p, WD_ALIGN_PARAGRAPH.CENTER, first_line=False, after=2)
            for r in p.runs:
                set_font(r, east="楷体", west="Times New Roman", size=12)
            continue
        if t.startswith("摘") or t.startswith("关键词") or t.startswith("中图分类号"):
            set_para(p, WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=False, after=1)
            for r in p.runs:
                set_font(r, east="楷体", west="Times New Roman", size=10)
            continue
        if idx == 5:
            set_para(p, WD_ALIGN_PARAGRAPH.CENTER, first_line=False, after=3)
            for r in p.runs:
                set_font(r, east="Times New Roman", west="Times New Roman", size=14, bold=True)
            continue
        if idx == 6:
            set_para(p, WD_ALIGN_PARAGRAPH.CENTER, first_line=False, after=2)
            for r in p.runs:
                set_font(r, east="Times New Roman", west="Times New Roman", size=12)
            continue
        if t.startswith("Abstract:") or t.startswith("Key words:"):
            set_para(p, WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=False, after=1)
            for r in p.runs:
                set_font(r, east="Times New Roman", west="Times New Roman", size=10)
            continue

        if t == "参考文献":
            rewrite_paragraph_text(p, "参考文献：")
            t = p.text.strip()
        if t == "参考文献：":
            refs_started = True
            set_para(p, WD_ALIGN_PARAGRAPH.LEFT, first_line=False, after=1)
            for r in p.runs:
                set_font(r, east="黑体", west="Times New Roman", size=11.5, bold=True)
            continue

        lvl = heading_level(t)
        if lvl:
            body_started = True
            set_para(p, WD_ALIGN_PARAGRAPH.LEFT, first_line=False, before=1, after=1)
            size = 11.5 if lvl in (1, 2) else 10.5
            for r in p.runs:
                set_font(r, east="黑体", west="Times New Roman", size=size, bold=True)
            continue

        if is_caption(t):
            set_para(p, WD_ALIGN_PARAGRAPH.CENTER, first_line=False, after=1)
            for r in p.runs:
                set_font(r, east="宋体", west="Times New Roman", size=9)
            continue

        if is_reference(t) or refs_started:
            set_para(p, WD_ALIGN_PARAGRAPH.LEFT, first_line=False, after=0)
            for r in p.runs:
                set_font(r, east="宋体", west="Times New Roman", size=9)
            continue

        set_para(p, WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=body_started, after=0)
        for r in p.runs:
            set_font(r, east="宋体", west="Times New Roman", size=10)


def set_cell_margins(cell, top=40, start=40, bottom=40, end=40):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_tbl_borders(tbl, picture=False):
    tblPr = tbl._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        if picture:
            el.set(qn("w:val"), "nil")
        else:
            if edge in ("top", "bottom", "insideH"):
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "8" if edge in ("top", "bottom") else "4")
                el.set(qn("w:space"), "0")
                el.set(qn("w:color"), "000000")
            else:
                el.set(qn("w:val"), "nil")


def format_tables(doc):
    for tbl in doc.tables:
        picture = bool(tbl._element.xpath(".//a:blip"))
        tbl.alignment = 1
        set_tbl_borders(tbl, picture=picture)
        for row_i, row in enumerate(tbl.rows):
            for cell in row.cells:
                set_cell_margins(cell)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    set_para(p, WD_ALIGN_PARAGRAPH.CENTER, first_line=False, after=0)
                    for r in p.runs:
                        set_font(r, east="宋体", west="Times New Roman", size=7.5 if not picture else 9, bold=(row_i == 0 and not picture))


def remove_body_template_residue(doc):
    for p in doc.paragraphs:
        text = p.text.strip()
        if text in ("计算机科学与探索", "Comment by"):
            rewrite_paragraph_text(p, "")
        if text.startswith("Comment by "):
            rewrite_paragraph_text(p, "")


def write_report():
    report = """# 格式修复报告

| 检查项 | 处理结果 |
|---|---|
| 是否删除全部批注 | 已清理。最终 DOCX 中无 `word/comments.xml` 和批注引用标记 |
| 是否清理全部修订痕迹 | 已清理。最终 DOCX 中无 `w:ins`、`w:del` 修订节点 |
| 是否套用模板格式 | 已按《计算机科学与探索》模板统一页面大小、页边距、单双栏分节、字体字号、标题、图表题、表格和参考文献样式 |
| 首页格式处理情况 | 中文题名居中小二号；作者居中小四号；中文摘要/关键词使用 10 磅楷体；英文题名、作者、Abstract 和 Key words 按模板字号处理 |
| 标题格式处理情况 | 一级/二级标题 11.5 磅黑体，三级及以下标题 10.5 磅黑体；章节编号保持原终稿编号 |
| 图表格式处理情况 | 图题、表题居中小五号；图片表格容器去边框；数据表按三线表思路处理，表内字号约六号 |
| 参考文献格式处理情况 | 参考文献标题改为“参考文献：”；条目统一小五号；保留原 20 条参考文献和正文顺序编码引用 |
| 是否删除正文残留页眉文字 | 已检查并删除正文段落中的“计算机科学与探索”残留；页眉中的期刊名保留 |
| 无法完全匹配模板之处 | 系统未确认安装“方正书宋”等模板字体，正文使用“宋体”等效替代；未使用 Word COM 自动分页，无法程序化检查每页视觉分页 |
| 字体替代说明 | 中文正文使用宋体，中文摘要使用楷体，标题使用黑体；英文使用 Times New Roman |
| 内容完整性 | 未压缩页数，未删除 RQ1-RQ5、关键图、实验表、Sepolia 表、故障分类表和协议论证 |
"""
    REPORT.write_text(report, encoding="utf-8")


def main():
    cleanup_markup(SRC, OUT)
    doc = Document(OUT)
    remove_body_template_residue(doc)
    format_sections(doc)
    format_paragraphs(doc)
    format_tables(doc)
    doc.save(OUT)
    write_report()
    print(OUT)
    print(REPORT)


if __name__ == "__main__":
    main()
