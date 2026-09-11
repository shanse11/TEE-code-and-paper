from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


BASE = Path("E:/项目/project_code/paper_outputs/draft")
SRC = BASE / "Hybrid TEE-Rollup.docx"
OUT = BASE / "final_recovered_teacher_revised_8pages.docx"
TEACHER_REPORT = BASE / "teacher_comments_revision_report.md"
RECOVERY_REPORT = BASE / "content_recovery_report.md"
FIG_DIR = BASE / "teacher_recovered_assets"


def set_run_font(run, size=9.2, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def set_para_format(p, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=0):
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(after)
    if first_line:
        pf.first_line_indent = Pt(17)


def add_p(doc, text, size=9.2, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=True):
    p = doc.add_paragraph()
    set_para_format(p, first_line=first_line, align=align)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    set_para_format(p, first_line=False, align=WD_ALIGN_PARAGRAPH.LEFT, after=1)
    r = p.add_run(text)
    set_run_font(r, size=10.5 if level == 1 else 9.7, bold=True)
    return p


def add_caption(doc, cn, en):
    p = doc.add_paragraph()
    set_para_format(p, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    r = p.add_run(cn + "\n" + en)
    set_run_font(r, size=8.0)
    return p


def add_table_title(doc, cn, en):
    p = doc.add_paragraph()
    set_para_format(p, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    r = p.add_run(cn + "\n" + en)
    set_run_font(r, size=8.2)


def set_columns(section, num):
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    if cols:
        cols = cols[0]
    else:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(num))
    if num == 2:
        cols.set(qn("w:space"), "360")


def start_single_column(doc):
    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(sec, 1)
    return sec


def start_two_column(doc):
    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(sec, 2)
    return sec


def style_table(tbl, size=8.0):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    for row in tbl.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_run_font(run, size=size)


def set_cell_text(cell, text, size=8.0, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(str(text))
    set_run_font(r, size=size, bold=bold)


def table_from_rows(doc, rows, size=8.0):
    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    style_table(tbl, size=size)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            set_cell_text(tbl.cell(i, j), value, size=size, bold=(i == 0))
    return tbl


def extract_picture_from_table(src_doc, table_idx, out_name):
    FIG_DIR.mkdir(exist_ok=True)
    table = src_doc.tables[table_idx]
    for blip in table._element.xpath(".//a:blip"):
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue
        part = src_doc.part.related_parts[rid]
        ext = part.content_type.split("/")[-1]
        if ext == "jpeg":
            ext = "jpg"
        out = FIG_DIR / f"{out_name}.{ext}"
        out.write_bytes(part.blob)
        return out
    return None


def add_picture(doc, path, width=3.05):
    if path and path.exists():
        p = doc.add_paragraph()
        set_para_format(p, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run()
        r.add_picture(str(path), width=Inches(width))


def setup_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.25)
    sec.bottom_margin = Cm(1.15)
    sec.left_margin = Cm(1.25)
    sec.right_margin = Cm(1.25)
    set_columns(sec, 2)
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(9.2)
    return doc


def build():
    src = Document(SRC)
    fig1 = extract_picture_from_table(src, 2, "fig1_architecture")
    fig2 = extract_picture_from_table(src, 3, "fig2_challenge")
    fig3 = extract_picture_from_table(src, 5, "fig3_commit_size")
    fig4 = extract_picture_from_table(src, 6, "fig4_da_cost")
    fig5 = extract_picture_from_table(src, 8, "fig5_timeout_recovery")

    doc = setup_doc()
    p = doc.add_paragraph()
    set_para_format(p, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("面向高频 DApp 的 Hybrid TEE-Rollup 可恢复挑战协议与轻量 DA 成本评估")
    set_run_font(r, size=13.0, bold=True)
    add_p(doc, "杨帆", size=10.0, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)

    abstract = (
        "摘 要：高频去中心化应用（DApp）需要在低延迟、低链上开销和异常执行下的公开可验证仲裁之间取得平衡。"
        "全链上执行成本高，完全链下执行缺乏公开仲裁入口，单纯依赖可信执行环境（TEE）又难以覆盖超时、数据不可用和证明不一致等异常。"
        "针对 Hybrid TEE-Rollup 中轻量链上承诺、链外数据可用性（DA）负载和超时挑战之间的协同问题，本文提出三项核心设计：可恢复挑战协议、证据绑定型紧凑提交和 DA 感知成本评估模型。"
        "可恢复挑战协议在不改变争议事实的前提下，将因交互超时停滞的会话恢复为可重放、可结算状态，以保障仲裁连续性；证据绑定型紧凑提交将状态根、输出哈希、证明哈希、DA 指针与 DA 根绑定为异常路径的最小链上证据入口；DA 感知成本模型刻画负载规模、批处理策略和 DA 路径对摊销成本的影响。"
        "基于 Python 原型、本地 EVM gas 测量和 Sepolia 测试网部署记录，实验显示紧凑提交平均约 406 字节；在给定配置化成本模型和实验参数下，模块化 DA 采样路径的摊销成本估算较全链上调用数据下降 96.33%；在当前原型的合成超时场景中，恢复路径将挑战完成率由 0% 提升至 100%。"
        "本文结果基于模拟 TEE、可验证 DA 注册表和配置化成本模型，说明机制接口、链上对应物与测试网部署的可运行性，不代表真实硬件 TEE 安全、真实 DA 网络经济安全或主网生产性能。"
    )
    add_p(doc, abstract, size=8.8)
    add_p(doc, "关键词：Hybrid TEE-Rollup；可恢复挑战协议；证据绑定型紧凑提交；数据可用性；仲裁连续性", size=8.8)

    p = doc.add_paragraph()
    set_para_format(p, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Recoverable Challenge Protocol and Lightweight DA Cost Evaluation for Hybrid TEE-Rollup in High-Frequency DApps")
    set_run_font(r, size=10.7, bold=True)
    add_p(doc, "YANG Fan", size=9.3, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    en_abs = (
        "Abstract: High-frequency decentralized applications must balance low latency, low on-chain overhead, and publicly verifiable arbitration under abnormal execution. "
        "Full on-chain execution is costly, fully off-chain execution lacks a public arbitration entry, and a pure trusted execution environment (TEE) path does not by itself address timeout, data unavailability, or proof inconsistency. "
        "This paper studies the coordination among lightweight on-chain commitments, off-chain data availability (DA) payloads, and timeout challenges in Hybrid TEE-Rollup. "
        "It proposes a recoverable challenge protocol, an evidence-binding compact commit, and a DA-aware cost evaluation model. "
        "The recovery protocol restores timeout sessions to replayable and settleable states without changing disputed facts; the compact commit binds the state root, output hash, proof hash, DA pointer, and DA root as a minimal evidence entry for abnormal paths; and the cost model characterizes the effects of payload size, batching, and DA path selection. "
        "Experiments based on a Python prototype, local EVM gas measurement, and Sepolia deployment records show that compact commits average about 406 bytes. Under the configured model and parameters, the modular DA sampling path reduces estimated amortized cost by 96.33% compared with full on-chain calldata. In synthetic timeout scenarios of the current prototype, recovery improves challenge completion from 0% to 100%. "
        "The results are bounded by a simulated TEE, a verifiable DA registry, and a configurable cost model, and do not imply hardware TEE security, economic security of production DA networks, or mainnet performance."
    )
    add_p(doc, en_abs, size=8.3)
    add_p(doc, "Key words: Hybrid TEE-Rollup; recoverable challenge; evidence-binding compact commit; data availability; arbitration continuity", size=8.3)

    add_heading(doc, "1 引言")
    add_heading(doc, "1.1 研究背景与问题", 2)
    add_p(doc, "高频 DApp 通常具有短交互周期、批量状态更新和频繁结算需求。若所有执行细节都放在链上，调用数据和验证逻辑会迅速放大 gas 开销；若完全依赖链下执行，用户又难以在异常发生时取得公开可验证的仲裁入口。Rollup 通过“链下执行、链上验证”的方式缓解该矛盾，但不同 Rollup 路线在成本、最终性和异常处理之间仍存在取舍[6-7]。")
    add_p(doc, "ZK-Rollup 依赖有效性证明压缩链上验证，但证明生成成本和复杂执行支持仍是工程负担。Optimistic Rollup 通过挑战窗口和欺诈证明降低正常路径成本，Arbitrum、Optimism 等系统展示了交互式争议定位与 fault proof 的实践价值[10,16-17]。然而，高频应用不仅关心错误能否被最终发现，也关心挑战流程在超时、中断或数据缺失后是否仍能推进。")
    add_p(doc, "TEE 路线试图利用可信执行环境降低链下执行可信成本。Hybrid TEE-Rollup 结合 TEE 快速路径和 Rollup 异常验证入口，在正常情况下可以减少链上负担[1]。但纯 TEE 假设无法覆盖所有风险：远程证明可能无效，宿主可能隐藏输入输出，DA 负载可能不可得，挑战交互可能因超时停滞。因此，系统需要把 TEE 输出、DA 证据和链上挑战放到同一个可验证框架中。")
    add_p(doc, "本文关注的核心问题是：在不把完整负载搬到链上的前提下，如何保留异常路径所需的最小证据入口；在挑战会话超时后，如何不改变争议事实而恢复协议活性；在不同 DA 路径和批处理参数下，如何解释高频负载的摊销成本趋势。将轻量提交、DA 证据和恢复挑战联合研究，是因为三者分别对应正常路径成本、异常路径证据和协议活性，缺一都会削弱 Hybrid TEE-Rollup 的可用性。")
    add_heading(doc, "1.2 本文贡献", 2)
    for item in [
        "（1）提出可恢复挑战协议，将超时后停滞的挑战会话恢复为可重放、可结算状态，在不改变挑战事实的条件下保障仲裁连续性。",
        "（2）设计证据绑定型紧凑提交，将状态根、输出哈希、证明哈希、DA 指针与 DA 根绑定，形成异常路径的最小链上证据入口。",
        "（3）构建 DA 感知成本评估模型，分析负载规模、批处理和 DA 路径对摊销成本的影响，并通过原型实验给出趋势性证据。"
    ]:
        add_p(doc, item)

    add_heading(doc, "2 背景与威胁模型")
    add_heading(doc, "2.1 Rollup、TEE 与数据可用性", 2)
    add_p(doc, "Rollup 的基本思想是把大量执行移到链下完成，只把状态承诺、证明或争议入口保留在链上。ZK-Rollup 以有效性证明证明状态转移正确，链上验证成本稳定但链下证明生成复杂；Optimistic Rollup 默认链下结果正确，在挑战窗口内允许验证者发起欺诈证明，适合把成本集中到异常路径[6-7]。交互式挑战通常通过二分定位缩小争议步骤，最终在链上或可验证环境中检查单步执行。")
    add_p(doc, "数据可用性决定验证者能否取得重放和检查所需的原始负载。LazyLedger、轻客户端和模块化 DA 网络把数据发布与执行验证解耦，EIP-4844 进一步为以太坊 Rollup 提供 blob 数据通道[3-5,12-13]。本文不实现真实 DA 网络，而用可验证 DA 注册表抽象负载哈希、DA 根、字段级 Merkle 证明和可用性状态。")
    add_p(doc, "TEE 在 Hybrid TEE-Rollup 中承担快速执行和摘要生成角色。SGX、TDX 等硬件可提供隔离执行和远程证明接口[8-9]，但本文使用模拟 TEE 建模输入输出绑定关系，避免把硬件安全、侧信道防护和供应链信任混入机制评估。TEE 输出必须通过链上承诺、DA 证明和挑战流程形成可审计证据，而不能仅凭硬件声明获得最终信任。")
    add_heading(doc, "2.2 问题范围与对抗假设", 2)
    add_p(doc, "本文考虑的异常包括错误输出、无效远程证明、DA 负载缺失、字段级证明损坏、响应篡改、轨迹长度不匹配和挑战超时。对手可以在链下执行和交互过程中制造不一致，但不能破坏哈希抗碰撞性，不能伪造链上合约状态，也不能修改已经提交的状态根、DA 根和挑战证据。")
    add_p(doc, "研究边界包括三点。第一，模拟 TEE 只验证协议接口和摘要绑定，不证明真实硬件 TEE 安全。第二，可验证 DA 注册表只证明字段级验证流程可运行，不证明 Celestia、EigenDA 等真实网络的经济安全。第三，成本模型是配置化趋势模型，不预测主网价格、拥堵费率或 blob 基础费。形式化安全和交互式证明理论作为语义参照，而非完整证明系统[18-20]。")

    add_heading(doc, "3 系统架构与总体设计")
    add_heading(doc, "3.1 系统角色与正常路径", 2)
    add_p(doc, "系统包含四类角色：执行层负责接收请求并在模拟 TEE 中生成输出摘要；DA 层保存完整负载并提供 Merkle 证明；链上提交层接收紧凑提交并记录状态根、输出哈希、证明哈希、DA 指针和 DA 根；验证仲裁层负责挑战、恢复、重放和结算。")
    add_p(doc, "正常路径按以下顺序运行：用户请求进入执行层，模拟 TEE 产生执行输出和证明摘要；完整负载写入 DA 注册表并得到负载哈希与 DA 根；提交者调用链上合约提交紧凑承诺；若挑战窗口内无人提出有效异议，提交被最终确认。该路径把高频负载的大部分数据留在链外，同时为异常路径保留可追溯入口。")
    add_picture(doc, fig1, width=3.05)
    add_caption(doc, "图 1 Hybrid TEE-Rollup 原型系统架构", "Fig.1 Prototype architecture of Hybrid TEE-Rollup")
    add_heading(doc, "3.2 异常路径与合约映射", 2)
    add_p(doc, "异常路径从不一致检测开始。验证者可以根据链上 DA 指针加载字段级证明，检查负载是否可用、输出是否匹配、证明摘要是否与提交一致；若争议无法直接裁决，则进入交互式挑战和二分定位。挑战中断或超时后，恢复操作检查会话是否满足恢复条件，并把它推进到可重放或可结算状态。")
    add_p(doc, "Solidity 原型提供机制对应而非 Python 原型的逐行移植。MockTEEVerifier 抽象 TEE 证明验证，MockDARegistry 记录负载哈希和 DA 根，MerkleVerifier 验证字段级证明，HybridTEERollup 维护提交、挑战、恢复、重放和结算状态。该映射用于证明协议路径具有链上对应物，不代表生产合约已经审计或主网可用。")
    add_heading(doc, "3.3 设计原则", 2)
    add_p(doc, "系统设计遵循三条原则：正常路径最小化链上对象，异常路径保留足够证据，恢复路径只处理活性而不重写事实。第一条原则降低高频提交成本；第二条原则保证验证者在异常时可以加载 DA 证据并进入仲裁；第三条原则避免恢复操作成为改变裁决结果的入口。")

    add_heading(doc, "4 核心机制与成本评估")
    add_heading(doc, "4.1 可恢复挑战协议", 2)
    add_p(doc, "挑战会话可以表示为包含会话编号、提交承诺、挑战方、被挑战方、执行轨迹、当前区间、超时参数和证据集合的状态对象。会话从打开争议开始，经过响应、二分定位、单步重放和结算等阶段。若某一阶段超过交互期限，会话进入停滞状态；停滞状态不是裁决结果，而是协议活性中断。")
    add_p(doc, "可恢复挑战协议把恢复函数定义为受限状态转移。恢复函数可以改变会话的活性状态、下一步操作和超时计数，但不能改变状态根、输出哈希、证明哈希、DA 指针、DA 根、负载哈希、执行轨迹或已经提交的重放证据。该约束把安全性与活性分离：正确性仍由重放和证据检查裁决，恢复只保证争议能够继续推进。")
    add_p(doc, "协议包含三条典型路径。正常路径中提交未被挑战并最终确认；可仲裁异常路径中验证者发现不一致，打开挑战，通过二分定位和重放获得裁决；超时恢复路径中挑战已被打开但交互停滞，恢复操作将会话重新连接到重放或结算流程。三条路径共享同一组挑战事实，因此恢复不会制造新的争议事实。")
    add_picture(doc, fig2, width=3.05)
    add_caption(doc, "图 2 可恢复挑战协议流程", "Fig.2 Workflow of the recoverable challenge protocol")
    add_p(doc, "仲裁连续性是本文强调的协议性质：当异常已被检测且证据入口仍然存在时，协议应保留一条从当前会话状态通向重放和结算的有效路径。没有恢复机制时，超时可能使检测结果停留在未解决状态；引入恢复后，超时不再等价于争议失效，而是触发受约束的活性修复。")
    add_p(doc, "恢复机制不改变正确性裁决。若原始输出正确，重放仍应支持被挑战方；若输出、证明或 DA 字段存在不一致，重放和结算仍应支持挑战方。恢复的意义在于把“已检测但无法推进”的会话转化为“可继续验证”的会话，从而避免高频系统在异常路径上积累不可结算争议。")
    add_heading(doc, "4.2 证据绑定型紧凑提交与 DA 绑定", 2)
    add_p(doc, "证据绑定型紧凑提交包含五个核心字段：状态根约束执行前后状态，输出哈希约束链下执行结果，证明哈希约束 TEE 或证明摘要，DA 指针定位链外负载，DA 根约束完整负载的 Merkle 承诺。该结构不是简单压缩字段，而是把正常路径提交对象转化为异常路径的最小证据入口。")
    add_p(doc, "当响应被篡改时，输出哈希和状态根可用于发现声明结果与提交上下文不一致；当证明摘要不一致时，证明哈希提供检查入口；当 DA 证明损坏时，DA 根和字段级 Merkle 证明可定位具体字段；当负载缺失时，DA 指针和可用性状态说明验证者无法取得重放材料。不同异常都围绕同一紧凑提交展开，避免异常路径重新构造证据上下文。")
    add_p(doc, "紧凑提交与字段级 Merkle 证明的绑定关系是：链上只保存 DA 根和指针，链外保存完整负载和字段路径；挑战者在需要时提交字段值、路径和对应索引，合约根据 DA 根验证字段属于原始负载。这样既隔离完整负载规模，又使异常路径拥有可验证的字段级证据。")
    add_picture(doc, fig3, width=3.05)
    add_caption(doc, "图 3 完整负载与紧凑提交的字节规模对比", "Fig.3 Byte-size comparison between full payloads and compact commits")
    add_heading(doc, "4.3 DA 感知成本评估模型", 2)
    add_p(doc, "本文比较四类 DA 路径：全链上调用数据把完整负载作为 calldata 发布，成本随负载线性增长；紧凑外部 DA 只在链上记录指针和根，成本主要来自固定提交和外部可用性假设；类 EIP-4844 路径用 blob 思想降低数据发布成本[5]；模块化 DA 采样路径进一步把可用性证明外包给专门网络[12-13]。")
    add_p(doc, "成本模型写为 C_total = C_commit + C_DA + p_challenge × C_challenge，C_amortized = C_total / batch_size。其中，C_commit 表示链上固定提交成本，C_DA 表示给定路径下的数据发布或采样成本，C_challenge 表示异常挑战、恢复、重放和结算的期望成本，p_challenge 表示挑战发生概率，batch_size 表示批处理规模。")
    add_p(doc, "该模型用于解释结构性权衡。负载越大，全链上路径越容易被数据发布成本主导；批大小越大，固定提交成本越容易被摊薄；DA 路径越轻量，越需要明确异常路径中的证据可得性和边界。模型中的 96.33% 等结果是给定参数下的趋势估算，不能外推为主网价格结论。")
    add_picture(doc, fig4, width=3.05)
    add_caption(doc, "图 4 不同 DA 配置下的摊销 gas 趋势", "Fig.4 Amortized gas trends under different DA configurations")

    add_heading(doc, "5 实验评估")
    add_heading(doc, "5.1 实验设置与研究问题", 2)
    add_p(doc, "实验由三部分组成：Python 原型验证提交、DA 证明、挑战、恢复、重放和结算流程；配置化成本脚本生成不同负载、批大小和 DA 路径下的摊销趋势；Hardhat 本地 EVM 与 Sepolia 部署记录验证合约对应物。实验使用模拟 TEE 和合成负载，目的是验证机制可运行性和趋势解释能力。")
    add_p(doc, "研究问题包括：RQ1，紧凑提交能否隔离负载增长；RQ2，批处理与 DA 路径如何影响摊销成本；RQ3，恢复机制能否改善超时挑战活性；RQ4，不同故障场景能否被检测、恢复或惩罚；RQ5，Solidity、本地 EVM 与 Sepolia 证据能否支撑链上对应物。")
    add_heading(doc, "5.2 RQ1：紧凑提交能否隔离负载增长", 2)
    add_p(doc, "RQ1 比较完整负载与紧凑提交的字节规模。实验目的不是证明某个编码格式最优，而是观察链上提交对象是否随负载增长而线性放大。结果显示，紧凑提交平均约为 406 字节，在不同负载规模下保持近似稳定；完整负载则随 payload size 增长而增长。")
    add_p(doc, "该结果说明，状态根、输出哈希、证明哈希、DA 指针和 DA 根足以构成异常路径入口，而完整负载可以留在 DA 层。其意义在于把高频正常路径的链上开销从“随负载增长”转化为“固定承诺加外部 DA 证据”。边界在于，紧凑提交本身不保证 DA 网络真实可用，仍依赖后续字段级证明和可用性假设。")
    add_heading(doc, "5.3 RQ2：批处理与 DA 路径如何影响摊销成本", 2)
    add_p(doc, "RQ2 比较全链上调用数据、紧凑外部 DA、紧凑类 EIP-4844 和紧凑模块化 DA 采样四类路线。实验改变负载大小和批大小，观察 C_amortized 的趋势。结果表明，当负载发布成为主要开销时，批处理可以显著摊薄固定成本，而模块化 DA 采样在大负载、大批量条件下优势更明显。")
    add_p(doc, "以负载 8192、批大小 1000 为例，在给定配置化成本模型和实验参数下，模块化 DA 采样路径的摊销成本估算较全链上调用数据下降 96.33%。这一结果意味着 DA 路径选择会改变高频应用的成本可行性，但它不是主网交易费预测，也没有模拟 blob 基础费率、网络拥堵和真实 DA 服务定价。")
    add_heading(doc, "5.4 RQ3：恢复机制能否改善超时挑战活性", 2)
    add_p(doc, "RQ3 构造合成超时场景，比较不可恢复路径和可恢复路径。不可恢复路径中，挑战会话在超时后无法自然进入重放和结算，完成率为 0%；可恢复路径中，恢复操作检查会话证据并重新连接到后续状态，当前原型设置下完成率为 100%。")
    add_p(doc, "该结果说明，超时检测本身不等于仲裁完成。恢复机制的价值在于把停滞会话重新转化为可裁决会话，从而改善协议活性。实验边界也很明确：合成场景没有覆盖真实网络延迟、恶意拥塞、多并发挑战队列和恢复频率限制，因此不能直接推出生产环境的活性保证。")
    add_picture(doc, fig5, width=3.05)
    add_caption(doc, "图 5 超时恢复对挑战完成率的影响", "Fig.5 Impact of timeout recovery on challenge completion rate")
    add_heading(doc, "5.5 RQ4：不同故障场景能否被检测、恢复或惩罚", 2)
    add_p(doc, "RQ4 将故障分为正常路径、证明类异常、DA 类异常、执行响应异常和超时异常。实验目的在于区分“可检测”“可恢复推进”“可进入重放”和“可触发罚没”四种语义。远程证明无效、数据不可用和 DA 证明无效可以被检测或拒绝，但当前原型未把所有此类故障统一映射为罚没。响应篡改和轨迹长度不匹配可进入挑战闭环并触发惩罚。")
    start_single_column(doc)
    add_table_title(doc, "表 1 故障场景分类与证据边界", "Table 1 Failure scenario classification and evidence boundaries")
    fault_rows = [
        ["故障类型", "是否可检测", "是否可恢复推进", "是否可进入重放", "是否触发罚没", "当前证据边界"],
        ["正常路径", "否", "不适用", "否", "否", "说明正常提交不误报"],
        ["远程证明无效", "是", "否", "否", "否", "可拒绝无效证明，不证明真实 TEE 安全"],
        ["数据不可用", "是", "否", "否", "否", "可识别 DA 缺失，未接入真实 DA 网络"],
        ["DA 证明无效", "是", "否", "否", "否", "字段级 Merkle 路径校验失败"],
        ["响应篡改", "是", "是", "是", "是", "可通过输出哈希和重放证据裁决"],
        ["轨迹长度不匹配", "是", "是", "是", "是", "可进入二分定位和结算"],
        ["超时不可恢复", "是", "否", "否", "否", "检测到停滞但无法推进"],
        ["超时可恢复", "是", "是", "是", "按重放结果", "恢复只改变活性状态，不改变事实"],
    ]
    table_from_rows(doc, fault_rows, size=7.2)
    start_two_column(doc)
    add_p(doc, "表 1 表明，本文机制并不把所有异常都解释为同一种安全事件。DA 不可用更接近证据缺失，响应篡改更接近可惩罚执行错误，超时则是活性问题。区分这些语义有助于避免把检测率、挑战成功率和罚没率混为一个指标。")
    add_heading(doc, "5.6 RQ5：Solidity、本地 EVM 与 Sepolia 证据", 2)
    add_p(doc, "RQ5 评估链上对应物。Solidity 合约把 Python 原型中的关键路径映射为 registerDA、submitRollup、challengeOpen、challengeRespond、challengeStep、challengeRecover、challengeReplay、challengeResolve 和 challengeFinalize 等操作。本地 EVM gas 结果说明这些路径可以被合约执行，Sepolia 部署记录说明核心合约具备公开测试网可部署性。")
    start_single_column(doc)
    add_table_title(doc, "表 2 本地 EVM 关键路径平均 gasUsed", "Table 2 Average gasUsed of local EVM key paths")
    gas_rows = [
        ["操作", "平均 gasUsed", "解释"],
        ["registerDA", "113526.00", "MockDARegistry 记录负载哈希与 DA 根"],
        ["submitRollup", "299033.67", "紧凑提交路径"],
        ["challengeOpen", "283769.00", "打开争议并加载证据"],
        ["challengeRespond", "65886.00", "被挑战方响应"],
        ["challengeStep", "93807.00", "二分定位单轮推进"],
        ["challengeRecover", "156891.00", "恢复超时会话"],
        ["challengeReplay", "135632.00", "单步重放仲裁"],
        ["challengeResolve", "113402.00", "结算争议结果"],
        ["challengeFinalize", "33477.00", "挑战期后最终确认"],
    ]
    table_from_rows(doc, gas_rows, size=7.6)
    add_table_title(doc, "表 3 Sepolia 部署信息", "Table 3 Sepolia deployment information")
    sepolia_rows = [
        ["项目", "数值"],
        ["Network / Chain ID", "Sepolia / 11155111"],
        ["部署区块", "10825904"],
        ["MockTEEVerifier", "0xC42CB0Cf0D112Bd59E0f212F2DB2002ca50a8b6A"],
        ["MockDARegistry", "0xc42C64a7De05bf3ad4f1c75CbD8E0506F7De99Ff"],
        ["HybridTEERollup", "0xB9B72f10bB8aBC1ed090f1B0443Fe57189c497Fd"],
    ]
    table_from_rows(doc, sepolia_rows, size=7.5)
    start_two_column(doc)
    add_p(doc, "gas 数据应理解为链上路径存在性的证据，而不是主网性能评估。challengeRecover 的平均 gas 高于单轮 challengeStep，说明恢复是异常路径中的额外成本；但该成本的比较对象不是正常路径吞吐，而是未解决争议和可恢复仲裁之间的差异。")

    add_heading(doc, "6 讨论、局限性与相关工作")
    add_heading(doc, "6.1 证据边界与局限性", 2)
    add_p(doc, "本文的证据边界集中在执行环境、DA 网络和成本模型三个方面。模拟 TEE 说明接口能够产生可绑定摘要，但不覆盖真实硬件的侧信道攻击、远程证明兼容性和供应链信任。可验证 DA 注册表说明字段级证明能与链上根绑定，但不代表真实 DA 网络的经济安全和大规模采样保证。配置化成本模型说明路径趋势，不代表主网价格。")
    add_p(doc, "当前原型尚未覆盖多并发挑战、恢复队列增长、挑战者激励、罚没参数设计和跨批次依赖。恢复机制还需要频率限制和反滥用策略，否则恶意参与者可能通过反复触发恢复增加系统负担。这些问题不影响本文机制语义，但决定其能否进入完整生产协议栈。")
    add_heading(doc, "6.2 相关工作对比", 2)
    add_p(doc, "TEEROLLUP / Hybrid TEE-Rollup 研究利用异构 TEE 降低 Rollup 验证成本，并保留挑战和 DA 惩罚机制[1]。本文与其方向一致，但聚焦更窄的异常路径：当挑战会话超时后，如何在不改变争议事实的前提下恢复仲裁连续性。")
    add_p(doc, "乐观 TEE Rollup 和 OTR 类工作通常将 TEE 快速执行与乐观验证结合，目标是降低正常路径成本。opML 和 optimistic ML fraud proof 则把交互式欺诈证明用于机器学习或大模型计算，强调二分定位和单步仲裁。本文借鉴交互式争议定位思想，但不实现完整 FPVM 或真实 ML 证明，而是服务于 Hybrid TEE-Rollup 中的异常仲裁。")
    add_p(doc, "Dynamic Fraud Proof 强调无争议场景下的最终性压缩，并在争议出现时动态延迟结算[2]。该方向关注挑战窗口和验证者参与机制的动态调整；本文关注的是已经进入挑战后的停滞会话恢复问题，恢复目标不是缩短最终性，而是防止异常路径卡死。")
    add_p(doc, "LazyLedger、模块化 DA、Celestia、EigenDA 和 EIP-4844 展示了数据可用性与执行验证解耦的主要路线[3-5,12-13]。本文不提出新的 DA 网络，而是研究紧凑链上承诺如何与 DA 指针、DA 根和字段级证明绑定，使链上对象在低成本的同时保留异常路径证据。")
    add_p(doc, "Arbitrum 和 Optimism fault proof 展示了乐观 Rollup 的争议推进实践[16-17]。本文与这些系统的共同点是把成本集中到异常路径；区别在于本文额外引入模拟 TEE 快速路径、DA 绑定提交和超时恢复语义，目标是解释高频 DApp 下轻量提交与可恢复仲裁的组合机制。")
    add_heading(doc, "6.3 未来工作", 2)
    add_p(doc, "后续工作可沿三条路线推进：接入真实 TEE 与真实 DA 网络，重新测量证明验证、数据可用性和恢复开销；完善链上重放证明和罚没策略，使不同故障映射到更精细的经济后果；研究并发挑战、恢复频率限制和队列调度，避免恢复机制被滥用。")

    add_heading(doc, "7 结论")
    add_p(doc, "本文在 Hybrid TEE-Rollup 背景下研究高频 DApp 的轻量提交、DA 证据和超时挑战问题，提出可恢复挑战协议、证据绑定型紧凑提交和 DA 感知成本评估模型。可恢复挑战协议把恢复定义为受约束的活性状态转移，保持挑战事实不变并保障仲裁连续性；证据绑定型紧凑提交为异常路径提供最小链上证据入口；成本模型解释负载规模、批处理和 DA 路径之间的趋势权衡。")
    add_p(doc, "实验基于 Python 原型、本地 EVM gas 和 Sepolia 部署记录，说明机制接口、链上对应物和测试网部署均可运行。结果同时受限于模拟 TEE、可验证 DA 注册表和配置化成本模型。总体而言，本文的价值在于恢复并连接 Hybrid TEE-Rollup 中的三个关键环节：正常路径低成本、异常路径证据可得和超时后仲裁可继续推进。")

    add_heading(doc, "参考文献")
    refs = [
        "[1] WEN X, FENG Q, LYU H, et al. TEEROLLUP: Efficient Rollup Design Using Heterogeneous TEE[EB/OL]. arXiv:2409.14647, 2024.",
        "[2] PICCO G, FORTUGNO A. Dynamic Fraud Proof[EB/OL]. arXiv:2502.10321, 2025.",
        "[3] AL-BASSAM M. LazyLedger: A Distributed Data Availability Ledger with Client-Side Smart Contracts[EB/OL]. arXiv:1905.09274, 2019.",
        "[4] TAS E N, TSE D, YANG L, et al. Light Clients for Lazy Blockchains[EB/OL]. arXiv:2203.15968, 2022.",
        "[5] Ethereum Foundation. EIP-4844: Shard Blob Transactions[EB/OL]. 2024.",
        "[6] Ethereum Foundation. Optimistic Rollups[EB/OL]. Ethereum Documentation, 2024.",
        "[7] BUTERIN V. An Incomplete Guide to Rollups[EB/OL]. 2021.",
        "[8] Intel Corporation. Intel Software Guard Extensions Developer Guide[EB/OL]. 2024.",
        "[9] Intel Corporation. Intel Trust Domain Extensions Architecture Specification[EB/OL]. 2024.",
        "[10] Offchain Labs. Arbitrum: A Next-Generation Layer 2 for Ethereum[EB/OL]. 2021.",
        "[11] Cartesi. Cartesi Rollups Documentation[EB/OL]. 2024.",
        "[12] EigenLayer. EigenDA Documentation[EB/OL]. 2024.",
        "[13] Celestia. Celestia: Modular Blockchain Network[EB/OL]. 2024.",
        "[14] RISC Zero. RISC Zero zkVM Documentation[EB/OL]. 2024.",
        "[15] Succinct Labs. SP1: A Performant, 100% Open-Source zkVM[EB/OL]. 2024.",
        "[16] Arbitrum Foundation. Arbitrum Nitro Technical Documentation[EB/OL]. 2024.",
        "[17] OP Labs. Optimism Documentation: Fault Proofs and Dispute Games[EB/OL]. 2024.",
        "[18] CANETTI R, et al. Universally Composable Security: A New Paradigm for Cryptographic Protocols[EB/OL]. 2000.",
        "[19] GOLDWASSER S, MICALI S, RACKOFF R. The Knowledge Complexity of Interactive Proof Systems[J]. SIAM Journal on Computing, 1989, 18(1): 186-208.",
        "[20] NIST. Recommendation for Key Management: Part 1 - General (Rev. 5)[EB/OL]. 2020.",
    ]
    for ref in refs:
        add_p(doc, ref, size=7.4, first_line=False)

    doc.save(OUT)
    return OUT


def write_reports():
    teacher = """# 导师批注处理说明

| 导师批注原意 | 修改位置 | 具体修改动作 | 是否完成 | 说明 |
|---|---|---|---|---|
| 英文摘要请对照中文摘要修改 | 中英文摘要 | 重写中文摘要为 350-450 字左右的完整摘要，并同步重写英文摘要，保持背景、方法、结果和边界一致 | 完成 | 删除病句、错字、异常符号和不规范空格 |
| 参考文献要在正文中合适的位置引用 | 第 1、2、4、6 章 | 在 Rollup、Optimistic Rollup、TEE、DA、TEEROLLUP、Dynamic Fraud Proof、Arbitrum/Optimism、交互式证明等位置加入顺序编码引用 | 完成 | 保留 20 条参考文献，正文引用不再只集中于文末 |
| 删除“详见第 X 章”等提前预告 | 全文 | 重写衔接句，删除章节预告式表达 | 完成 | 已检索清理“详见第”“将在第”“后文将”等表达 |
| 原第 4、5、6 章合并为新第 4 章 | 第 4 章 | 合并为“4 核心机制与成本评估”，下设 4.1 可恢复挑战协议、4.2 证据绑定型紧凑提交与 DA 绑定、4.3 DA 感知成本评估模型 | 完成 | 后续实验、讨论和结论重新编号 |
| 参考文献格式不规范 | 参考文献 | 统一为顺序编码条目，并在正文中按语义引用 | 基本完成 | 建议最终按学校模板补充访问日期等细节 |
"""
    TEACHER_REPORT.write_text(teacher, encoding="utf-8")

    recovery = """# 内容恢复说明

| 项目 | 说明 |
|---|---|
| 上一版中过度删除的部分 | 引言动机、Rollup/TEE/DA 背景、系统正常与异常路径、第 4 章理论解释、RQ1-RQ5 实验解释、故障分类完整语义、相关工作定位被压得过短 |
| 本次从导师批注原稿中恢复的内容 | 恢复高频 DApp 动机、全链上/链下/TEE 路线不足、Rollup 与交互式挑战背景、系统角色与 Solidity 合约映射、挑战会话与恢复约束、紧凑提交字段和异常定位、DA 成本公式、RQ1-RQ5 分节、完整故障分类、gas 表、Sepolia 表、相关工作自然段 |
| 恢复后仍然压缩的内容 | 删除机制推进点表、协议不变量长表、二分轮次图、失败检测图、实验小结表、相关工作大表；将重复边界说明合并到摘要、实验和讨论中的必要位置 |
| 最终页数 | 本机 Word COM 自动分页不可用，无法程序化确认；文档已按 7.5-8 页目标采用双栏、窄页边距和宽表单栏切换设计 |
| 是否保持 8 页以内 | 已按 8 页以内进行版式和内容控制，建议在 Word 中打开后用状态栏页数做最终确认 |
| 是否保留论文完整性 | 是。新版不再是 3 页摘要稿，恢复了完整学术论文所需的问题动机、技术背景、机制论证、实验 RQ、图表证据、局限性和相关工作 |
"""
    RECOVERY_REPORT.write_text(recovery, encoding="utf-8")


if __name__ == "__main__":
    print(build())
    write_reports()
    print(TEACHER_REPORT)
    print(RECOVERY_REPORT)
