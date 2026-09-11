from pathlib import Path
from docx import Document

path = Path(__file__).with_name("final_layout_fixed_only.docx")
doc = Document(str(path))

for si, section in enumerate(doc.sections, 1):
    print(si, "page", section.page_width, "margins", section.left_margin, section.right_margin)
    print(str(section._sectPr.xml)[:500].replace("\n", ""))

print("tables", len(doc.tables))
for i, table in enumerate(doc.tables, 1):
    rows = len(table.rows)
    cols = len(table.columns)
    first = " | ".join(cell.text.replace("\n", "/")[:70] for cell in table.rows[0].cells) if rows else ""
    print(i, rows, cols, first)
