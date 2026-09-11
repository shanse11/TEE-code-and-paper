from __future__ import annotations

from pathlib import Path
import sys

from docx import Document


TEMPLATE = Path(r"D:\Microsoft downloads\20260123103641 (1).docx")
PAPER = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(r"E:\项目\project_code\paper_outputs\draft\Hybrid_TEE_Rollup_参考文献格式修订版.docx")


def pt(value):
    return None if value is None else round(value.pt, 2)


def inspect(path: Path, label: str) -> None:
    doc = Document(str(path))
    print(f"\n== {label} ==")
    refs = []
    start = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().replace(" ", "") in {"参考文献", "参考文献："}:
            start = i
            refs.append((i, p))
            continue
        if start is not None and p.text.strip().startswith("["):
            refs.append((i, p))
            if len(refs) >= 8:
                break
    if start is None:
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip().startswith("[1]"):
                refs.append((i, p))
                break
    for i, p in refs:
        fmt = p.paragraph_format
        run = p.runs[0] if p.runs else None
        print("P", i, repr(p.text[:130]))
        print(" style", p.style.name if p.style else None, "align", p.alignment, "left", pt(fmt.left_indent), "first", pt(fmt.first_line_indent), "space_before", pt(fmt.space_before), "space_after", pt(fmt.space_after), "line", fmt.line_spacing)
        if run:
            print(" run", "font", run.font.name, "size", pt(run.font.size), "bold", run.bold, "sup", run.font.superscript)


inspect(TEMPLATE, "template")
inspect(PAPER, "paper")
