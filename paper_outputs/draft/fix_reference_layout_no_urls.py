from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


BASE = Path(__file__).resolve().parent
SRC = BASE / "Hybrid_TEE_Rollup_参考文献格式模板修订版.docx"
FALLBACK = BASE / "Hybrid_TEE_Rollup_参考文献格式修订版.docx"
OUT = BASE / "Hybrid_TEE_Rollup_参考文献排版修订版.docx"

REFS = [
    "[1] WEN X, FENG Q, LYU H, et al. TEEROLLUP: efficient Rollup design using heterogeneous TEE[EB/OL]. arXiv:2409.14647, 2024 [2026-06-15].",
    "[2] BUTERIN V. An incomplete guide to Rollups[EB/OL]. 2021 [2026-06-15].",
    "[3] ETHEREUM FOUNDATION. Optimistic Rollups[EB/OL]. 2024 [2026-06-15].",
    "[4] ARBITRUM FOUNDATION. Arbitrum Nitro technical documentation[EB/OL]. 2024 [2026-06-15].",
    "[5] OP LABS. Optimism documentation: fault proofs and dispute games[EB/OL]. 2024 [2026-06-15].",
    "[6] PICCO G, FORTUGNO A. Dynamic fraud proof[EB/OL]. arXiv:2502.10321, 2025 [2026-06-15].",
    "[7] AL-BASSAM M. LazyLedger: a distributed data availability ledger with client-side smart contracts[EB/OL]. arXiv:1905.09274, 2019 [2026-06-15].",
    "[8] TAS E N, TSE D, YANG L, et al. Light clients for lazy blockchains[EB/OL]. arXiv:2203.15968, 2022 [2026-06-15].",
    "[9] EIGENLABS. EigenDA documentation[EB/OL]. 2024 [2026-06-15].",
    "[10] CELESTIA. Celestia: modular blockchain network[EB/OL]. 2024 [2026-06-15].",
    "[11] INTEL CORPORATION. Intel Software Guard Extensions developer guide[EB/OL]. 2024 [2026-06-15].",
    "[12] ETHEREUM FOUNDATION. EIP-4844: shard blob transactions[EB/OL]. 2024 [2026-06-15].",
]


def set_east_asian_font(run, name: str) -> None:
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def clean_reference_text(text: str) -> str:
    text = text.replace("\u00ad", "").replace("\ufffe", "")
    text = re.sub(r"https?://\S+", "", text)
    text = re.sub(r"\[\[www\]\(http://www/\)\]\(http://www\)\.?", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"\s+\.", ".", text)
    if not text.endswith("."):
        text += "."
    return text


def style_reference_paragraph(paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.left_indent = Pt(18.4)
    fmt.first_line_indent = Pt(-18.4)
    fmt.line_spacing = 1.2
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def replace_reference(paragraph, text: str) -> None:
    paragraph.clear()
    run = paragraph.add_run(clean_reference_text(text))
    run.font.name = "Times New Roman"
    set_east_asian_font(run, "宋体")
    run.font.size = Pt(9)
    run.bold = False
    run.italic = False
    run.font.superscript = False
    style_reference_paragraph(paragraph)


def style_heading(paragraph) -> None:
    paragraph.clear()
    run = paragraph.add_run("参考文献：")
    run.font.size = Pt(11.5)
    run.bold = True
    set_east_asian_font(run, "黑体")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = None
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.line_spacing = None


def main() -> None:
    src = SRC if SRC.exists() else FALLBACK
    doc = Document(str(src))
    paragraphs = doc.paragraphs

    ref_idx = None
    for i, p in enumerate(paragraphs):
        if p.text.strip().replace(" ", "") in {"参考文献", "参考文献："}:
            ref_idx = i
            break
    if ref_idx is None:
        raise RuntimeError("Reference heading not found")

    style_heading(paragraphs[ref_idx])

    ref_paras = [p for p in paragraphs[ref_idx + 1 :] if re.match(r"^\[\d+\]", p.text.strip())]
    if len(ref_paras) != len(REFS):
        raise RuntimeError(f"Expected {len(REFS)} reference entries, found {len(ref_paras)}")

    for p, text in zip(ref_paras, REFS):
        replace_reference(p, text)

    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
