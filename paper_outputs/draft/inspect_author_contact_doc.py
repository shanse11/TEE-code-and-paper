from __future__ import annotations

from pathlib import Path
from docx import Document

P = Path(__file__).resolve().parent / "论文终版" / "Hybrid_TEE_Rollup.docx"

doc = Document(str(P))
print("paragraphs", len(doc.paragraphs))
print("FIRST")
for i, p in enumerate(doc.paragraphs[:30], 1):
    print(f"{i:03d}: {p.text!r}")
print("TAIL")
for i, p in list(enumerate(doc.paragraphs, 1))[-45:]:
    print(f"{i:03d}: {p.text!r}")

print("TABLES", len(doc.tables))
for ti, table in enumerate(doc.tables[:3], 1):
    print("TABLE", ti)
    for ri, row in enumerate(table.rows[:5], 1):
        print(" | ".join(cell.text.replace("\n", "\\n") for cell in row.cells))
