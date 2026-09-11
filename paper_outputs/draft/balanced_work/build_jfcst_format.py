from __future__ import annotations

import copy
import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from PIL import Image


SRC = Path("E:/项目/project_code/paper_outputs/draft/final_balanced_low_aigc_academic.docx")
TEMPLATE = Path("D:/Microsoft downloads/20260123103641 (1).docx")
OUT = Path("E:/项目/project_code/paper_outputs/draft/final_jfcst_submission_format_optimized.docx")


FIG_EN = {
    1: "Prototype architecture of Hybrid TEE-Rollup",
    2: "Workflow of the recoverable challenge protocol",
    3: "Byte-size comparison between full payloads and compact commits",
    4: "Amortized gas trends under different DA configurations",
    5: "Relationship between bisection rounds and trace steps",
    6: "Impact of timeout recovery on challenge completion rate",
    7: "Detection and arbitration outcomes under failure scenarios",
}

TAB_EN = {
    1: "Mechanism advances and roles",
    2: "Semantic comparison of mechanisms",
    3: "Protocol invariants and preserved properties",
    4: "Failure scenarios and evidence boundaries",
    5: "Average gasUsed of local EVM key paths",
    6: "Sepolia deployment information",
    7: "Summary of comparison experiments",
    8: "Comparison with related work",
}

EN_ABSTRACT = (
    "High-frequency decentralized applications (DApps) require low latency, low on-chain overhead, "
    "and publicly verifiable arbitration under abnormal executions. This paper studies the interaction "
    "between lightweight on-chain commitments, off-chain data availability (DA) payloads, and timeout "
    "challenges in Hybrid TEE-Rollup. It proposes a Recoverable Challenge protocol, an Evidence-Carrying "
    "Compact Commit, and a DA-Aware Cost Model. The recoverable challenge path restores timeout sessions "
    "to replayable and settleable states; the compact commit binds the state root, output hash, proof hash, "
    "DA pointer, and DA root; the cost model characterizes the effects of payload size, batching, and DA "
    "path selection. Experiments based on a Python prototype, local EVM gas measurements, and Sepolia "
    "deployment records show that compact commits average about 406 bytes. Under payload size 8192 and "
    "batch size 1000, the modular DA sampling path reduces the estimated amortized cost by 96.33% compared "
    "with full on-chain calldata. In synthetic timeout scenarios, recovery improves the challenge success "
    "rate from 0% on the unrecoverable path to 100%. The results are bounded by a simulated TEE, a verifiable "
    "DA registry, and a configurable cost model."
)


def clear_document(doc: Document) -> None:
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_east_asian_font(run, east_asia="SimSun", ascii_font="Times New Roman", size=None, bold=None):
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_paragraph_format(paragraph, *, align=None, first_line=None, before=0, after=0, line=12.5):
    if align is not None:
        paragraph.alignment = align
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if first_line is not None:
        pf.first_line_indent = Pt(first_line)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line)


def add_text_para(doc, text, *, style=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=10, east="SimSun",
                  ascii_font="Times New Roman", bold=False, first_line=20, before=0, after=0, line=13):
    p = doc.add_paragraph(style=style)
    set_paragraph_format(p, align=align, first_line=first_line, before=before, after=after, line=line)
    r = p.add_run(text)
    set_east_asian_font(r, east, ascii_font, size=size, bold=bold)
    return p


def add_title(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6, line=20)
    r = p.add_run(text)
    set_east_asian_font(r, "SimHei", "Times New Roman", size=18, bold=True)
    return p


def add_author(doc, text, english=False):
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=3, line=15)
    r = p.add_run(text)
    set_east_asian_font(r, "KaiTi", "Times New Roman", size=12, bold=False)
    return p


def add_abstract_para(doc, label, text, english=False):
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=0, before=2, after=2, line=13)
    r0 = p.add_run(label)
    set_east_asian_font(r0, "KaiTi" if not english else "Times New Roman", "Times New Roman", size=10, bold=True)
    r1 = p.add_run(text)
    set_east_asian_font(r1, "KaiTi" if not english else "Times New Roman", "Times New Roman", size=10, bold=False)
    return p


def add_heading(doc, number, title, level):
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=0, before=6, after=3, line=14)
    r = p.add_run(f"{number}  {title}")
    size = 11.5 if level in (1, 2) else 10.5
    set_east_asian_font(r, "SimHei", "Times New Roman", size=size, bold=True)
    return p


def add_small_heading(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=0, before=4, after=2, line=13)
    r = p.add_run(text)
    set_east_asian_font(r, "SimHei", "Times New Roman", size=10, bold=True)
    return p


def add_caption(doc, cn, en, kind="fig"):
    for text in (cn, en):
        p = doc.add_paragraph()
        set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0, before=0, after=0, line=10.5)
        r = p.add_run(text)
        set_east_asian_font(r, "SimSun", "Times New Roman", size=9, bold=False)


def get_blip_rids(paragraph):
    return paragraph._p.xpath(".//a:blip/@r:embed")


def image_blob_from_paragraph(src_doc, paragraph):
    rids = get_blip_rids(paragraph)
    if not rids:
        return None
    part = src_doc.part.related_parts[rids[0]]
    return part.blob


def crop_image_whitespace(blob):
    """Crop near-white margins so charts remain readable after Word scaling."""
    try:
        im = Image.open(BytesIO(blob)).convert("RGBA")
        bg = Image.new("RGBA", im.size, "WHITE")
        bg.alpha_composite(im)
        rgb = bg.convert("RGB")
        mask = rgb.point(lambda p: 0 if p > 248 else 255).convert("L")
        bbox = mask.getbbox()
        if not bbox:
            return blob
        margin = 12
        left = max(bbox[0] - margin, 0)
        top = max(bbox[1] - margin, 0)
        right = min(bbox[2] + margin, rgb.width)
        bottom = min(bbox[3] + margin, rgb.height)
        cropped = rgb.crop((left, top, right, bottom))
        out = BytesIO()
        cropped.save(out, format="PNG")
        return out.getvalue()
    except Exception:
        return blob


def set_columns(section, num=1, space_twips=425):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(space_twips))


def set_page(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(0.8)


def configure_sections(doc):
    for section in doc.sections:
        set_page(section)
        set_columns(section, 1)


def add_header_footer(doc):
    for section in doc.sections:
        header = section.header
        header.is_linked_to_previous = False
        if not header.paragraphs:
            header.add_paragraph()
        hp = header.paragraphs[0]
        hp.text = ""
        set_paragraph_format(hp, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0, before=0, after=0, line=10)
        r = hp.add_run("计算机科学与探索")
        set_east_asian_font(r, "SimSun", "Times New Roman", size=9)
        footer = section.footer
        footer.is_linked_to_previous = False
        if not footer.paragraphs:
            footer.add_paragraph()
        fp = footer.paragraphs[0]
        fp.text = ""
        set_paragraph_format(fp, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0, before=0, after=0, line=10)
        r = fp.add_run("")
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        r._r.append(fld_begin)
        r._r.append(instr)
        r._r.append(fld_end)


def set_cell_text_style(cell, header=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0, before=0, after=0, line=10)
        for run in p.runs:
            set_east_asian_font(run, "SimSun", "Times New Roman", size=7.5, bold=header)


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "nil")
    for edge in ("top", "bottom"):
        el = borders.find(qn(f"w:{edge}"))
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:color"), "000000")
    # header bottom rule
    if table.rows:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tr_borders = tr_pr.find(qn("w:tblHeader"))
        if tr_borders is None:
            hdr = OxmlElement("w:tblHeader")
            hdr.set(qn("w:val"), "true")
            tr_pr.append(hdr)
        for cell in table.rows[0].cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_borders = tc_pr.find(qn("w:tcBorders"))
            if tc_borders is None:
                tc_borders = OxmlElement("w:tcBorders")
                tc_pr.append(tc_borders)
            bottom = tc_borders.find(qn("w:bottom"))
            if bottom is None:
                bottom = OxmlElement("w:bottom")
                tc_borders.append(bottom)
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "000000")


def set_table_no_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "nil")


def add_caption_to_paragraph(paragraph, text):
    set_paragraph_format(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0, before=0, after=0, line=10.5)
    run = paragraph.add_run(text)
    set_east_asian_font(run, "SimSun", "Times New Roman", size=9, bold=False)


def add_table_from_data(doc, rows, table_no, caption):
    add_caption(doc, f"表 {table_no}  {caption}", f"Table {table_no}  {TAB_EN.get(table_no, caption)}", "table")
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            table.cell(i, j).text = value
            set_cell_text_style(table.cell(i, j), header=(i == 0))
    set_table_borders(table)
    p = doc.add_paragraph()
    set_paragraph_format(p, after=3)
    return table


def rows_from_table(tbl):
    return [[cell.text.replace("\n", " ").strip() for cell in row.cells] for row in tbl.rows]


def add_image(doc, src_doc, paragraph, fig_no, caption, wide=False):
    blob = image_blob_from_paragraph(src_doc, paragraph)
    if not blob:
        return
    blob = crop_image_whitespace(blob)
    if wide:
        sec = doc.add_section(WD_SECTION.CONTINUOUS)
        set_page(sec)
        set_columns(sec, 1)
    block = doc.add_table(rows=1, cols=1)
    block.alignment = WD_TABLE_ALIGNMENT.CENTER
    block.autofit = True
    set_table_no_borders(block)
    cell = block.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.keep_with_next = True
    width = Inches(6.5 if wide else 3.25)
    p.add_run().add_picture(BytesIO(blob), width=width)
    add_caption_to_paragraph(cell.add_paragraph(), f"图 {fig_no}  {caption}")
    add_caption_to_paragraph(cell.add_paragraph(), f"Fig.{fig_no}  {FIG_EN.get(fig_no, caption)}")
    if wide:
        sec = doc.add_section(WD_SECTION.CONTINUOUS)
        set_page(sec)
        set_columns(sec, 2)


def translate_fig_caption(text):
    m = re.match(r"图\s*(\d+)\s*(.*)", text)
    if not m:
        return None
    return int(m.group(1)), m.group(2).strip()


def translate_table_caption(text):
    # Only treat short title-like paragraphs as table captions; explanatory
    # paragraphs such as "表 2 用于..." must remain body text.
    if len(text.strip()) > 35:
        return None
    m = re.match(r"表\s*(\d+)\s*(.*)", text)
    if not m:
        return None
    return int(m.group(1)), m.group(2).strip()


def normalize_heading_text(number, title):
    replacements = {
        "协议性质分析": "协议性质",
        "对比实验小结": "实验小结",
    }
    return replacements.get(title, title)


def add_body_paragraph(doc, text):
    if not text.strip():
        return
    m = re.match(r"^(\d+(?:\.\d+)*)\s+(.+)$", text.strip())
    if m:
        num = m.group(1)
        title = normalize_heading_text(num, m.group(2))
        level = num.count(".") + 1
        add_heading(doc, num, title, level)
        return
    cn_sub = {
        "（一）挑战事实不变性": ("4.4.1", "挑战事实不变性"),
        "（二）安全性与活性的分离": ("4.4.2", "安全性与活性分离"),
        "（三）仲裁连续性": ("4.4.3", "仲裁连续性"),
        "（四）恢复操作正确性与抗滥用边界": ("4.4.4", "恢复操作正确性与抗滥用边界"),
    }
    if text.strip() in cn_sub:
        num, title = cn_sub[text.strip()]
        add_heading(doc, num, title, 3)
        return
    if text.strip() in {"协议性质", "协议不变量与保持性质"}:
        add_small_heading(doc, text.strip())
        return
    if text.strip().startswith("["):
        add_reference_para(doc, text.strip())
        return
    if "\n" in text:
        for part in text.split("\n"):
            add_text_para(doc, part.strip(), size=10, first_line=20, line=13)
        return
    add_text_para(doc, text.strip(), size=10, first_line=20, line=13)


def add_reference_para(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=0, before=0, after=0, line=10.5)
    r = p.add_run(text)
    set_east_asian_font(r, "SimSun", "Times New Roman", size=8.5)


def add_front_matter(doc, src_doc):
    title = src_doc.paragraphs[0].text.strip()
    author = "杨帆"
    abstract = src_doc.paragraphs[3].text.strip()
    keywords = src_doc.paragraphs[4].text.strip().replace("关键词：", "")

    add_title(doc, title)
    add_author(doc, author)
    add_abstract_para(doc, "摘  要：", abstract, english=False)
    add_abstract_para(doc, "关键词：", keywords, english=False)
    add_text_para(doc, "中图分类号：TP311    文献标志码：A", align=WD_ALIGN_PARAGRAPH.LEFT,
                  first_line=0, size=10, east="KaiTi", line=13)

    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=3, line=16)
    r = p.add_run("Recoverable Challenge Protocol and Lightweight DA Cost Evaluation for\nHybrid TEE-Rollup in High-Frequency DApps")
    set_east_asian_font(r, "Times New Roman", "Times New Roman", size=13, bold=True)
    add_author(doc, "YANG Fan", english=True)
    add_abstract_para(doc, "Abstract: ", EN_ABSTRACT, english=True)
    add_abstract_para(doc, "Key words: ", "Hybrid TEE-Rollup; interactive challenge; data availability; high-frequency DApp; cost amortization; fault recovery", english=True)


def build():
    src = Document(str(SRC))
    doc = Document(str(TEMPLATE))
    clear_document(doc)
    configure_sections(doc)
    add_front_matter(doc, src)

    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    set_page(sec)
    set_columns(sec, 2)

    pending_fig = None
    pending_table = None
    fig_no = 0
    table_no = 0
    para_idx = 0
    table_idx = 0
    started = False

    body = src.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            paragraph = src.paragraphs[para_idx]
            text = paragraph.text.strip()
            has_img = bool(child.xpath(".//w:drawing"))
            para_idx += 1
            if not started:
                if text == "1 引言":
                    started = True
                else:
                    continue
            if has_img:
                # Caption arrives in the next paragraph; defer insertion.
                pending_fig = paragraph
                continue
            fig_cap = translate_fig_caption(text)
            if fig_cap and pending_fig is not None:
                fig_no, cap = fig_cap
                add_image(doc, src, pending_fig, fig_no, cap, wide=(fig_no in {1, 2}))
                pending_fig = None
                continue
            table_cap = translate_table_caption(text)
            if table_cap:
                pending_table = table_cap
                continue
            if text.startswith("参考文献"):
                continue
            add_body_paragraph(doc, text)
        elif tag == "tbl":
            tbl = src.tables[table_idx]
            table_idx += 1
            if pending_table:
                table_no, cap = pending_table
                # Wide tables are easier to read in a one-column continuous section.
                sec = doc.add_section(WD_SECTION.CONTINUOUS)
                set_page(sec)
                set_columns(sec, 1)
                add_table_from_data(doc, rows_from_table(tbl), table_no, cap)
                sec = doc.add_section(WD_SECTION.CONTINUOUS)
                set_page(sec)
                set_columns(sec, 2)
                pending_table = None
            else:
                add_table_from_data(doc, rows_from_table(tbl), table_idx, f"表格 {table_idx}")

    # Insert reference heading before first reference if not already inserted.
    # The source loop adds references as normal paragraphs; ensure heading precedes them.
    # Simpler pass: if no heading was added, append a clear one before references by rebuilding tail is not needed
    # because source references follow 8.6. Add explicit heading before them by moving? Instead add if absent near end.
    # The existing body loop includes references, so detect and insert heading before first reference via XML.
    insert_reference_heading(doc)
    add_header_footer(doc)
    doc.save(str(OUT))
    print(OUT)


def insert_reference_heading(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:p"):
            txt = "".join(t.text for t in child.iter(qn("w:t")) if t.text).strip()
            if txt.startswith("[1]"):
                p = OxmlElement("w:p")
                child.addprevious(p)
                # Attach a temporary paragraph object by reloading is cumbersome; write OOXML directly.
                ppr = OxmlElement("w:pPr")
                p.append(ppr)
                jc = OxmlElement("w:jc")
                jc.set(qn("w:val"), "left")
                ppr.append(jc)
                r = OxmlElement("w:r")
                p.append(r)
                rpr = OxmlElement("w:rPr")
                r.append(rpr)
                b = OxmlElement("w:b")
                rpr.append(b)
                rfonts = OxmlElement("w:rFonts")
                rfonts.set(qn("w:eastAsia"), "SimHei")
                rfonts.set(qn("w:ascii"), "Times New Roman")
                rfonts.set(qn("w:hAnsi"), "Times New Roman")
                rpr.append(rfonts)
                sz = OxmlElement("w:sz")
                sz.set(qn("w:val"), "23")
                rpr.append(sz)
                t = OxmlElement("w:t")
                t.text = "参考文献"
                r.append(t)
                return


if __name__ == "__main__":
    build()
